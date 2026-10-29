import json
import gc
import io
import re
import tempfile
import time
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
import zipfile
import urllib.error

import app as app_module
import grounding
import ollama_client
import semantic_matching
import seniority
import tailoring_pipeline
import tailoring
import resume_layout
import sources.jobspy_source as jobspy_source
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


class FakeEmbeddingBackend:
    model_name = "fake-local-test-model"

    def encode(self, texts):
        vectors = []
        groups = {
            "automation": {"automation", "scripts", "workflow", "reporting"},
            "api": {"api", "apis", "flask", "backend", "integrations"},
            "cloud": {"cloud", "aws", "azure", "deployment"},
            "frontend": {"react", "javascript", "interface", "ui"},
            "unrelated": {"sales", "kitchen", "biology"},
        }
        for text in texts:
            lowered = text.lower()
            vector = []
            for terms in groups.values():
                vector.append(1.0 if any(term in lowered for term in terms) else 0.0)
            vectors.append(vector)
        return vectors


class FakeUrlopenResponse:
    def __init__(self, body: dict[str, object]):
        self.body = body

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def read(self):
        return json.dumps(self.body).encode("utf-8")


class FakeGeocodingProvider:
    name = "fake"

    def __init__(self, responses: dict[str, dict[str, object]]):
        self.responses = responses
        self.calls: list[tuple[str, str]] = []

    def geocode(self, query: str, *, country: str = "Canada") -> dict[str, object]:
        self.calls.append((query, country))
        if query not in self.responses:
            raise ValueError(f"No geocoding match for {query}")
        response = self.responses[query]
        if isinstance(response, Exception):
            raise response
        return response


class JobAssistantTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        app_module.DATA_DIR = Path(self.tmpdir.name)
        app_module.DB_PATH = app_module.DATA_DIR / "test.sqlite"
        app_module.app.config.update(TESTING=True)
        self.client = app_module.app.test_client()

    def tearDown(self):
        self.tmpdir.cleanup()

    def wait_for_tailoring_ai(self, response):
        self.assertEqual(response.status_code, 202, response.get_json())
        job_id = response.get_json()["job_id"]
        for _ in range(200):
            result = self.client.get(f"/api/tailoring/improve-jobs/{job_id}")
            if result.status_code != 202:
                return result
            time.sleep(0.01)
        self.fail("Tailoring AI job did not finish within the test timeout")

    def test_profile_save_and_txt_upload(self):
        response = self.client.post(
            "/api/profile",
            json={
                "name": "Terry",
                "key_skills": "Python, Flask, React, SQL, dashboards",
                "work_authorization": "Legally authorized to work in Canada.",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["name"], "Terry")

        upload = self.client.post(
            "/api/profile/upload",
            data={"file": (io.BytesIO(b"Built Flask dashboards with SQL."), "resume.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(upload.status_code, 200)
        self.assertIn("Flask dashboards", upload.get_json()["resume_text"])

    def test_profile_save_populates_evidence_library(self):
        self.client.post(
            "/api/profile",
            json={
                "name": "Terry",
                "work_experience": "Harmony 88 Co. | Backend Developer | 2023-2025. Built Python tools backed by PostgreSQL and automated monitoring.",
                "resume_text": "Harmony 88 Co. Built Python tools backed by PostgreSQL and automated monitoring. Implemented validation and logging.",
            },
        )
        evidence = self.client.get("/api/evidence").get_json()
        self.assertGreaterEqual(len(evidence), 2)
        self.assertTrue(any(item["source_kind"] == "profile" for item in evidence))
        self.assertTrue(any(item["source_kind"] == "resume_import" for item in evidence))

    def test_job_analysis_and_resume_variants_endpoints(self):
        self.client.post(
            "/api/profile",
            json={
                "name": "Terry",
                "work_experience": "Harmony 88 Co. | Backend Developer | 2023-2025. Built Python tools backed by PostgreSQL and automated monitoring.",
                "resume_text": "Harmony 88 Co. Built Python tools backed by PostgreSQL and automated monitoring. Implemented validation and logging.",
            },
        )
        job = self.client.post(
            "/api/jobs",
            json={
                "company": "Scotiabank",
                "title": "Software Developer",
                "location": "Toronto",
                "job_description": "Python PostgreSQL monitoring validation logging automation",
            },
        ).get_json()
        analysis = self.client.get(f"/api/jobs/{job['id']}/analysis").get_json()
        self.assertEqual(analysis["job_id"], job["id"])
        self.assertGreaterEqual(analysis["approved_evidence_count"], 2)
        self.assertTrue(analysis["matched_evidence"])
        variants = self.client.get(f"/api/jobs/{job['id']}/resume-variants").get_json()
        self.assertIsInstance(variants, list)

    def test_lakeview_qualification_fit_surfaces_stretch_reasons_without_changing_priority(self):
        self.client.post(
            "/api/profile",
            json={
                "name": "QA Applicant",
                "resume_text": "Built Python Flask APIs and SQL reports for internal tools.",
                "work_experience": "Software developer with three years of application experience.",
                "maximum_required_years": "3",
                "target_seniority_levels": "junior, intermediate",
            },
        )
        job = self.client.post(
            "/api/jobs",
            json={
                "company": "Lakeview Loan Servicing",
                "title": "Backend Engineer",
                "location": "Toronto",
                "overall_score": 89,
                "job_description": "Requirements: 3 to 5+ years of production backend experience, advanced distributed systems, Kafka and Kubernetes.",
                "required_skills": "3 to 5+ years of production backend experience, advanced distributed systems, Kafka, Kubernetes",
                "status": "Saved",
            },
        ).get_json()
        self.assertEqual(job["overall_score"], 89)
        self.assertLess(job["qualification_fit_score"], 75)
        self.assertEqual(job["qualification_fit_label"], "Stretch opportunity")
        reasons = " ".join(job["qualification_fit_reasons"])
        self.assertIn("Not demonstrated", reasons)
        self.assertIn("experience", reasons)

    def test_qualification_fit_penalizes_missing_explicit_years_without_profile_maximum(self):
        profile = {
            "target_seniority_levels": "junior, intermediate",
            "maximum_required_years": "",
            "work_experience": "Built internal Python tools and APIs.",
        }
        analysis = seniority.analyze_job_seniority({
            "title": "Backend Engineer",
            "job_description": "Required: 3 to 5+ years of production backend experience and advanced distributed systems.",
        })
        fit = seniority.score_seniority_fit(profile, analysis)
        self.assertTrue(fit["stretch_role"])
        self.assertLess(fit["score"], 75)
        self.assertIn("does not explicitly demonstrate", fit["explanation"])

    def test_ollama_profile_defaults_and_overrides(self):
        with patch.dict(
            "os.environ",
            {
                "OLLAMA_ENABLED": "true",
                "OLLAMA_BASE_URL": "http://ollama.local:11434",
                "OLLAMA_MODEL": "llama3.1:8b",
                "OLLAMA_TIMEOUT_SECONDS": "45",
                "OLLAMA_MAX_OUTPUT_TOKENS": "777",
                "OLLAMA_MAX_PROMPT_CHARS": "11111",
            },
            clear=False,
        ):
            config = ollama_client.OllamaConfig.from_mapping({})
        self.assertTrue(config.enabled)
        self.assertEqual(config.base_url, "http://ollama.local:11434")
        self.assertEqual(config.model, "llama3.1:8b")
        self.assertEqual(config.timeout_seconds, 45)
        self.assertEqual(config.max_output_tokens, 777)
        self.assertEqual(config.max_prompt_chars, 11111)

    def test_ollama_health_reports_available_model(self):
        with patch.object(ollama_client.urllib.request, "urlopen", return_value=FakeUrlopenResponse({"models": [{"name": "qwen3:4b"}]})):
            report = ollama_client.build_ollama_health_report(
                ollama_client.OllamaConfig(
                    enabled=True,
                    base_url="http://127.0.0.1:11434",
                    model="qwen3:4b",
                    timeout_seconds=1,
                    max_output_tokens=1200,
                    max_prompt_chars=24000,
                )
            )
        self.assertTrue(report["reachable"])
        self.assertTrue(report["model_available"])
        self.assertEqual(report["status"], "ok")

    def test_ollama_health_reports_missing_model(self):
        with patch.object(ollama_client.urllib.request, "urlopen", return_value=FakeUrlopenResponse({"models": [{"name": "another-model"}]})):
            report = ollama_client.build_ollama_health_report(
                ollama_client.OllamaConfig(
                    enabled=True,
                    base_url="http://127.0.0.1:11434",
                    model="qwen3:4b",
                    timeout_seconds=1,
                    max_output_tokens=1200,
                    max_prompt_chars=24000,
                )
            )
        self.assertTrue(report["reachable"])
        self.assertFalse(report["model_available"])
        self.assertEqual(report["status"], "missing_model")
        self.assertIn("ollama pull qwen3:4b", report["summary"])

    def test_ollama_health_reports_timeout_and_disabled_state(self):
        with patch.object(ollama_client.urllib.request, "urlopen", side_effect=TimeoutError):
            report = ollama_client.build_ollama_health_report(
                ollama_client.OllamaConfig(
                    enabled=True,
                    base_url="http://127.0.0.1:11434",
                    model="qwen3:4b",
                    timeout_seconds=1,
                    max_output_tokens=1200,
                    max_prompt_chars=24000,
                )
            )
        self.assertEqual(report["status"], "timeout")
        self.assertFalse(report["reachable"])
        self.assertFalse(report["model_available"])

        disabled = ollama_client.build_ollama_health_report(
            ollama_client.OllamaConfig(
                enabled=False,
                base_url="http://127.0.0.1:11434",
                model="qwen3:4b",
                timeout_seconds=1,
                max_output_tokens=1200,
                max_prompt_chars=24000,
            )
        )
        self.assertEqual(disabled["status"], "disabled")
        self.assertFalse(disabled["reachable"])

    def test_ollama_json_completion_disables_thinking_and_parses_final_response(self):
        captured = {}

        def fake_urlopen(request, timeout):
            captured["payload"] = json.loads(request.data.decode("utf-8"))
            return FakeUrlopenResponse({"response": '{"ok": true, "message": "hello"}'})

        config = ollama_client.OllamaConfig(
            enabled=True,
            base_url="http://127.0.0.1:11434",
            model="qwen3:4b",
            timeout_seconds=1,
            max_output_tokens=120,
            max_prompt_chars=24000,
        )
        with patch.object(ollama_client.urllib.request, "urlopen", side_effect=fake_urlopen):
            result = ollama_client.request_json_completion(config, "Return a tiny JSON object.")

        self.assertTrue(result["ok"])
        self.assertEqual(result["data"], {"ok": True, "message": "hello"})
        self.assertIs(captured["payload"]["think"], False)
        self.assertEqual(captured["payload"]["options"]["num_predict"], 120)

    def test_ollama_json_completion_rejects_empty_final_response(self):
        config = ollama_client.OllamaConfig(
            enabled=True,
            base_url="http://127.0.0.1:11434",
            model="qwen3:4b",
            timeout_seconds=1,
            max_output_tokens=120,
            max_prompt_chars=24000,
        )
        with patch.object(
            ollama_client.urllib.request,
            "urlopen",
            return_value=FakeUrlopenResponse({"response": "", "thinking": "private reasoning"}),
        ):
            result = ollama_client.request_json_completion(config, "Return JSON.")

        self.assertFalse(result["ok"])
        self.assertEqual(result["error"], "Ollama returned an empty response.")

    def test_ollama_json_completion_retries_bounded_malformed_output(self):
        responses = iter([
            FakeUrlopenResponse({"response": "Here is the answer: not JSON"}),
            FakeUrlopenResponse({"response": "```json\n{\"proposed_bullet\": \"Built Python services.\"}\n```"}),
        ])
        config = ollama_client.OllamaConfig(True, "http://127.0.0.1:11434", "qwen3:4b", 1, 120, 24000)
        with patch.object(ollama_client.urllib.request, "urlopen", side_effect=lambda request, timeout: next(responses)):
            result = ollama_client.request_json_completion(config, "Return JSON.")
        self.assertTrue(result["ok"])
        self.assertEqual(result["attempts"], 2)
        self.assertEqual(result["data"]["proposed_bullet"], "Built Python services.")

    def test_job_create_scores_and_dashboard(self):
        self.seed_profile_for_documents(with_template=True)
        response = self.client.post(
            "/api/jobs",
            json={
                "company": "Northstar",
                "title": "Junior Full Stack Developer",
                "location": "Toronto hybrid",
                "salary": "$70k-$85k",
                "job_description": "Python JavaScript Flask React SQL AWS dashboards APIs internal tools",
            },
        )
        self.assertEqual(response.status_code, 201)
        job = response.get_json()
        self.assertGreaterEqual(job["fit_score"], 70)
        self.assertGreaterEqual(job["skills_match_score"], 70)

        dashboard = self.client.get("/api/dashboard").get_json()
        self.assertEqual(dashboard["jobs_saved"], 1)
        self.assertEqual(dashboard["high_fit"], 1)

    def test_generation_saves_versions_and_exports(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/jobs",
            json={
                "company": "OpsTools",
                "title": "Implementation Specialist",
                "location": "Remote Canada",
                "job_description": "APIs automation reporting business systems SQL Python",
            },
        ).get_json()
        generated = self.client.post("/api/generate", json={**created, "id": created["id"]})
        self.assertEqual(generated.status_code, 200)
        body = generated.get_json()
        self.assertIn("cover_letter", body)
        self.assertIn("resume_text", body)

        versions = self.client.get(f"/api/jobs/{created['id']}/versions").get_json()
        self.assertGreaterEqual(len(versions), 3)

        txt_export = self.client.get(f"/export/{created['id']}/cover_letter.txt")
        self.assertEqual(txt_export.status_code, 200)
        self.assertIn(b"Implementation Specialist", txt_export.data)

        docx_export = self.client.get(f"/export/{created['id']}/resume.docx")
        self.assertEqual(docx_export.status_code, 200)
        self.assertTrue(docx_export.data.startswith(b"PK"))

    def test_csv_import(self):
        csv_data = b"company,title,location,url,source,salary,description\nAcme,Cloud Support,Toronto,https://example.com,CSV,$60k,Azure SQL APIs\n"
        response = self.client.post(
            "/api/jobs/import-csv",
            data={"file": (io.BytesIO(csv_data), "jobs.csv")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["imported"], 1)
        jobs = self.client.get("/api/jobs").get_json()
        self.assertEqual(jobs[0]["company"], "Acme")

    def test_discovery_page_uses_compact_toolbar_and_full_width_layout(self):
        html = self.client.get("/").data.decode("utf-8")
        self.assertIn('id="discovery-submit">Search jobs</button>', html)
        self.assertIn('id="discovery-sort-mode"', html)
        self.assertIn('id="discovery-filter-toggle"', html)
        self.assertIn('id="discovery-filter-panel"', html)
        self.assertNotIn('id="discovery-sort-controls"', html)
        self.assertNotIn('workspace narrow', html)
        self.assertIn('workspace discovery-layout', html)
        self.assertLess(html.index('id="discovery-sort-mode"'), html.index('id="discovery-results"'))
        self.assertLess(html.index('id="discovery-results"'), html.index('id="csv-import-form"'))
        self.assertIn('id="csv-import-form"', html)
        self.assertIn('id="pipeline-job-details"', html)
        self.assertIn('Analyze job and generate diagnostic suggestions', html)
        self.assertIn('Sort: distance', html)
        self.assertIn('Sort: posting date', html)
        self.assertIn('id="pipeline-sort"', html)
        self.assertIn('Strong qualification matches', html)
        self.assertIn('data-view-application-package', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('focusPreparationPackage', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertNotIn('data-start-pipeline-preparation', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertNotIn('data-retry-pipeline-preparation', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('pipeline-drawer-in', Path('static/styles.css').read_text(encoding='utf-8'))
        self.assertNotIn('window.state.jobs', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('preparation-preview-columns', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('pipeline-matches-label', html)
        self.assertIn('Create candidate profile', html)
        self.assertIn('candidate-profile-form-error', html)
        self.assertIn('.candidate-profile-manager [hidden]', Path('static/styles.css').read_text(encoding='utf-8'))
        self.assertIn('refreshCandidateProfiles', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('openTailoringStudio', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('tailoring-studio', html)
        self.assertIn('keyword-memory-list', html)
        self.assertIn('tailoring-requirement-search', html)
        self.assertIn('tailoring-requirement-detail', html)
        self.assertIn('tailoring-editor', html)
        self.assertIn('tailoring-diff', html)
        self.assertIn('tailoring-open-preview', html)
        self.assertIn('tailoring-suggested-bullet', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('tailoring-requirement-filters', html)
        self.assertIn('buildTailoringRequirementGroups', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('Related résumé evidence for review', Path('static/app.js').read_text(encoding='utf-8'))

    def test_candidate_profile_creation_is_committed_and_returned_by_list_endpoint(self):
        response = self.client.post(
            "/api/candidate-profiles",
            json={"name": "Software Engineering", "identity_mode": "shared"},
        )
        self.assertEqual(response.status_code, 201)
        profile_id = response.get_json()["id"]
        self.assertIsInstance(profile_id, int)
        profiles = self.client.get("/api/candidate-profiles")
        self.assertEqual(profiles.status_code, 200)
        saved = next(profile for profile in profiles.get_json() if profile["id"] == profile_id)
        self.assertEqual(saved["name"], "Software Engineering")
        self.assertEqual(saved["identity_mode"], "shared")

    def test_tailoring_studio_analyzes_selected_resume_and_applies_copy_edit(self):
        self.client.post("/api/profile", json={"name": "Terry Mak", "email": "terry@example.com"})
        profile = self.client.post("/api/candidate-profiles", json={"name": "Software Engineering", "identity_mode": "shared"}).get_json()
        document = Document()
        document.add_heading("Terry Mak", level=1)
        document.add_paragraph("Experience")
        bullet = document.add_paragraph("Built Python automation for internal reporting.", style="List Bullet")
        bullet.runs[0].font.bold = True
        document.add_paragraph("Skills")
        document.add_paragraph("Python, SQL, Flask")
        buffer = io.BytesIO()
        document.save(buffer)
        upload = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "base.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        )
        self.assertEqual(upload.status_code, 201)
        resume_id = upload.get_json()["id"]
        job = self.client.post("/api/jobs", json={"company": "Example Co", "title": "Backend Developer", "job_description": "Python and REST APIs required. Kafka preferred."}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume_id}")
        self.assertEqual(workspace.status_code, 200, workspace.get_json())
        body = workspace.get_json()
        self.assertGreaterEqual(body["keywords"]["matched_count"], 1)
        editable = next(item for item in body["bullets"] if item["editable"])
        result = self.client.post(f"/api/tailoring/drafts/{body['draft']['id']}/apply", json={"paragraph_index": editable["paragraph_index"], "original_text": editable["text"], "proposed_text": "Built Python and REST API automation for internal reporting.", "operation": "edit"})
        self.assertEqual(result.status_code, 200, result.get_json())
        self.assertIn("REST API", result.get_json()["text"])
        draft_download = self.client.get(f"/api/tailoring/drafts/{body['draft']['id']}/download")
        self.assertEqual(draft_download.status_code, 200)
        self.assertEqual(draft_download.mimetype, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        second = self.client.post(f"/api/tailoring/drafts/{body['draft']['id']}/apply", json={"paragraph_index": editable["paragraph_index"], "original_text": "Built Python and REST API automation for internal reporting.", "proposed_text": "Built Python REST API automation for internal reporting.", "operation": "edit"})
        self.assertEqual(second.status_code, 200, second.get_json())
        self.assertIn("Built Python REST API automation", second.get_json()["text"])
        refreshed = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume_id}").get_json()
        self.assertGreaterEqual(refreshed["keywords"]["matched_count"], body["keywords"]["matched_count"])
        self.assertNotIn("REST API", app_module.extract_docx_text(Path(upload.get_json()["path"]).read_bytes()))

    def test_tailoring_apply_and_undo_rebuild_only_selected_draft_edit(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Tailoring Undo QA"}).get_json()
        document = Document()
        document.add_paragraph("Taylor Example")
        document.add_paragraph("Experience")
        first = document.add_paragraph("Built Python services for internal reporting.", style="List Bullet")
        second = document.add_paragraph("Created SQL validation workflows for operations.", style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "undo-qa.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Undo QA", "title": "Backend Engineer", "job_description": "Python, SQL and REST APIs required."}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        bullets = [item for item in workspace["bullets"] if item["editable"]]
        self.assertGreaterEqual(len(bullets), 2)
        draft_id = workspace["draft"]["id"]
        for bullet, replacement in zip(bullets[:2], ["Built Python REST services for internal reporting.", "Created SQL validation and audit workflows for operations."]):
            applied = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={
                "paragraph_index": bullet["paragraph_index"], "original_text": bullet["text"],
                "proposed_text": replacement, "operation": "edit",
            })
            self.assertEqual(applied.status_code, 200, applied.get_json())
        before_undo = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertEqual(len(before_undo["draft"]["changes"]), 2)
        undone = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={
            "paragraph_index": bullets[0]["paragraph_index"], "operation": "leave_unchanged",
        })
        self.assertEqual(undone.status_code, 200, undone.get_json())
        after_undo = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertEqual([change["paragraph_index"] for change in after_undo["draft"]["changes"]], [bullets[1]["paragraph_index"]])
        self.assertIn("Built Python services for internal reporting.", after_undo["preview"]["text"])
        self.assertIn("Created SQL validation and audit workflows for operations.", after_undo["preview"]["text"])
        self.assertIn("Applied to draft", Path("static/app.js").read_text(encoding="utf-8"))
        self.assertIn("Latest edit undone. The previous saved revision is restored.", Path("static/app.js").read_text(encoding="utf-8"))
        self.assertIn('toggleAttribute("disabled"', Path("static/app.js").read_text(encoding="utf-8"))

    def test_tailoring_applies_skills_and_experience_atomically_and_blocks_phantom_approval_changes(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Section Edit QA"}).get_json()
        document = Document()
        document.add_paragraph("Taylor Example")
        document.add_paragraph("Experience")
        experience = document.add_paragraph("Built Python services for internal reporting.", style="List Bullet")
        document.add_paragraph("Skills")
        skills = document.add_paragraph("Technologies: React, Flask, Node.js, Express, REST APIs, PostgreSQL")
        buffer = io.BytesIO(); document.save(buffer)
        source_bytes = buffer.getvalue()
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(source_bytes), "section-edit-qa.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Section Edit Co", "title": "Backend Engineer", "job_description": "Python, REST APIs and PostgreSQL required."}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        experience_item = next(item for item in workspace["bullets"] if item["text"] == experience.text)
        skills_item = next(item for item in workspace["bullets"] if item["text"] == skills.text)
        self.assertTrue(skills_item["editable"])

        first = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/apply", json={
            "paragraph_index": experience_item["paragraph_index"], "original_text": experience.text,
            "proposed_text": "Built Python REST services for internal reporting.", "operation": "edit",
        })
        self.assertEqual(first.status_code, 200, first.get_json())
        second = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/apply", json={
            "paragraph_index": skills_item["paragraph_index"], "original_text": skills.text,
            "proposed_text": skills.text + ", etc", "operation": "edit",
        })
        self.assertEqual(second.status_code, 200, second.get_json())
        draft_id = workspace["draft"]["id"]
        draft = self.client.get(f"/api/tailoring/drafts/{draft_id}/download")
        generated = Document(io.BytesIO(draft.data))
        generated_text = "\n".join(paragraph.text for paragraph in generated.paragraphs)
        self.assertIn("Built Python REST services for internal reporting.", generated_text)
        self.assertIn("Technologies: React, Flask, Node.js, Express, REST APIs, PostgreSQL, etc", generated_text)
        self.assertEqual(len(self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()["draft"]["changes"]), 2)
        undone_skill = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": skills_item["paragraph_index"], "operation": "leave_unchanged"})
        self.assertEqual(undone_skill.status_code, 200, undone_skill.get_json())
        self.assertIn(skills.text, undone_skill.get_json()["text"])
        reapplied_skill = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={
            "paragraph_index": skills_item["paragraph_index"], "original_text": skills.text,
            "proposed_text": skills.text + ", etc", "operation": "edit",
        })
        self.assertEqual(reapplied_skill.status_code, 200, reapplied_skill.get_json())
        reverted_skill = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": skills_item["paragraph_index"], "operation": "revert_original"})
        self.assertEqual(reverted_skill.status_code, 200, reverted_skill.get_json())
        self.assertIn(skills.text, reverted_skill.get_json()["text"])
        reapplied_for_review = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={
            "paragraph_index": skills_item["paragraph_index"], "original_text": skills.text,
            "proposed_text": skills.text + ", etc", "operation": "edit",
        })
        self.assertEqual(reapplied_for_review.status_code, 200, reapplied_for_review.get_json())

        protected = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={
            "paragraph_index": 1, "original_text": "Experience", "proposed_text": "Education", "operation": "edit",
        })
        self.assertEqual(protected.status_code, 409)
        self.assertIn("not saved", protected.get_json()["error"])
        after_rejection = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertEqual(len(after_rejection["draft"]["changes"]), 2)
        self.assertIn("etc", after_rejection["preview"]["text"])

        with app_module.app.app_context():
            app_module.get_db().execute("INSERT INTO application_preparation_queue(job_id, status, candidate_profile_id, resume_id) VALUES (?, 'queued', ?, ?)", [job["id"], profile["id"], resume["id"]])
            app_module.get_db().execute("UPDATE tailoring_drafts SET changes_json = ? WHERE id = ?", [json.dumps(after_rejection["draft"]["changes"] + [{"paragraph_index": skills_item["paragraph_index"], "original_text": skills.text, "proposed_text": "Technologies: fabricated", "keyword": "skills"}]), draft_id])
            app_module.get_db().commit()
        with patch.object(app_module, "validate_tailoring_draft_layout", return_value={"available": True, "verified": True, "original_pages": 1, "candidate_pages": 1, "line_mismatches": []}):
            approval = self.client.post(f"/api/tailoring/drafts/{draft_id}/approve")
        self.assertEqual(approval.status_code, 409)
        self.assertIn("does not contain its saved proposed wording", approval.get_json()["error"])
        frontend = Path("static/app.js").read_text(encoding="utf-8")
        self.assertIn('"Failed to undo" : "Failed to apply"', frontend)

    def test_tailoring_protected_metadata_is_not_offered_as_editable(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Protected Metadata QA"}).get_json()
        document = Document(); document.add_heading("Taylor Example", level=1); document.add_paragraph("Experience")
        document.add_paragraph("Built Python services for reporting.", style="List Bullet"); document.add_paragraph("Skills")
        document.add_paragraph("Python, SQL")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(f"/api/candidate-profiles/{profile['id']}/resumes/upload", data={"file": (io.BytesIO(buffer.getvalue()), "protected.docx"), "is_default": "true"}, content_type="multipart/form-data").get_json()
        job = self.client.post("/api/jobs", json={"company": "Protected Co", "title": "Engineer", "job_description": "Python required."}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        headings = [item for item in workspace["bullets"] if item["text"] in {"Taylor Example", "Experience", "Skills"}]
        self.assertTrue(headings)
        self.assertTrue(all(not item["editable"] for item in headings))

    def test_tailoring_revision_history_supports_stepwise_undo_and_revert(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Revision History QA"}).get_json()
        document = Document(); document.add_paragraph("Taylor Example"); document.add_paragraph("Experience")
        bullet = document.add_paragraph("Built Python services for reporting.", style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "revision-qa.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Revision QA", "title": "Backend Engineer", "job_description": "Python required."}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        draft_id = workspace["draft"]["id"]
        index = next(item["paragraph_index"] for item in workspace["bullets"] if item["editable"])
        original = "Built Python services for reporting."
        edit_a = "Built Python backend services for reporting."
        edit_b = "Built Python backend services and APIs for reporting."
        self.assertEqual(self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": index, "original_text": original, "proposed_text": edit_a, "operation": "edit"}).status_code, 200)
        self.assertEqual(self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": index, "original_text": edit_a, "proposed_text": edit_b, "operation": "edit"}).status_code, 200)
        undone = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": index, "operation": "leave_unchanged"})
        self.assertEqual(undone.status_code, 200, undone.get_json())
        self.assertIn(edit_a, undone.get_json()["text"])
        refreshed = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertEqual(refreshed["draft"]["changes"][0]["proposed_text"], edit_a)
        undone_again = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": index, "operation": "leave_unchanged"})
        self.assertEqual(undone_again.status_code, 200, undone_again.get_json())
        self.assertIn(original, undone_again.get_json()["text"])
        reapplied = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": index, "original_text": original, "proposed_text": edit_a, "operation": "edit"})
        self.assertEqual(reapplied.status_code, 200)
        reverted = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": index, "operation": "revert_original"})
        self.assertEqual(reverted.status_code, 200, reverted.get_json())
        self.assertIn(original, reverted.get_json()["text"])
        download = self.client.get(f"/api/tailoring/drafts/{draft_id}/download")
        self.assertIn("Revision_History_QA_Resume_Backend_Engineer_Draft.docx", download.headers.get("Content-Disposition", ""))

    def test_preparation_download_names_are_profile_and_job_specific(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Terry Mak", "applicant_name": "Terry Mak"}).get_json()
        job = self.client.post("/api/jobs", json={"company": "Name QA", "title": "Senior Backend / API Engineer"}).get_json()
        with app_module.app.app_context():
            row = {"candidate_profile_id": profile["id"], "job_id": job["id"]}
            self.assertEqual(app_module.tailoring_resume_filename(row, draft=True), "Terry_Mak_Resume_Senior_Backend_API_Engineer_Draft.docx")

    def test_tailoring_apply_defers_word_validation_and_preview_cache_tracks_draft_version(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Tailoring Performance QA"}).get_json()
        document = Document(); document.add_paragraph("Taylor Example"); document.add_paragraph("Experience")
        document.add_paragraph("Built Python services for internal reporting.", style="List Bullet")
        document.add_paragraph("Created SQL validation workflows for operations.", style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "performance-qa.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Performance QA", "title": "Backend Engineer", "job_description": "Python and SQL required."}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        bullet = next(item for item in workspace["bullets"] if item["editable"])
        draft_id = workspace["draft"]["id"]
        with patch.object(app_module, "compare_docx_layout") as layout_mock:
            applied = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={
                "paragraph_index": bullet["paragraph_index"], "original_text": bullet["text"],
                "proposed_text": "Built Python REST services for internal reporting.", "operation": "edit",
            })
        self.assertEqual(applied.status_code, 200, applied.get_json())
        self.assertFalse(layout_mock.called)
        self.assertTrue(applied.get_json()["layout"]["pending"])
        self.assertIn("word_layout_validation", applied.get_json()["timings_ms"])

        layout = {"available": True, "verified": True, "original_pages": 1, "candidate_pages": 1, "line_mismatches": []}
        def fake_export(_source, output):
            Path(output).write_bytes(b"%PDF-1.7 performance QA")
            return {"available": True, "path": str(output)}

        def fake_combined(_source, _candidate, output):
            Path(output).write_bytes(b"%PDF-1.7 performance QA")
            return {**layout, "available": True, "verified": True, "pdf_path": str(output)}

        with patch.object(app_module, "compare_docx_layout_and_export", side_effect=fake_combined) as combined_mock:
                first_preview = self.client.get(f"/api/tailoring/drafts/{draft_id}/preview").get_json()
                second_preview = self.client.get(f"/api/tailoring/drafts/{draft_id}/preview").get_json()
        self.assertTrue(first_preview["pdf_available"])
        self.assertFalse(first_preview["cached"])
        self.assertTrue(second_preview["cached"])
        self.assertEqual(combined_mock.call_count, 1)
        self.assertGreaterEqual(first_preview["timings_ms"]["preview_total"], 0)

        with patch.object(app_module, "compare_docx_layout", return_value=layout) as compare_after_change:
            changed = self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={
                "paragraph_index": bullet["paragraph_index"], "original_text": "Built Python REST services for internal reporting.",
                "proposed_text": "Built Python and SQL REST services for internal reporting.", "operation": "edit",
            })
        self.assertEqual(changed.status_code, 200, changed.get_json())
        self.assertFalse(compare_after_change.called)
        with patch.object(app_module, "compare_docx_layout_and_export", side_effect=fake_combined) as combined_after_version:
                refreshed_preview = self.client.get(f"/api/tailoring/drafts/{draft_id}/preview").get_json()
        self.assertTrue(refreshed_preview["pdf_available"])
        self.assertFalse(refreshed_preview["cached"])
        self.assertEqual(combined_after_version.call_count, 1)

    def test_tailoring_workspace_reports_missing_description_then_uses_persisted_description(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Lakeview QA"}).get_json()
        document = Document(); document.add_paragraph("Alex Rivera"); document.add_paragraph("Built Python services.", style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "lakeview.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Lakeview Loan Servicing", "title": "Backend Engineer"}).get_json()
        empty = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertFalse(empty["keywords"]["available"])
        self.assertEqual(empty["keywords"]["total"], 0)
        self.assertIn("complete job description", empty["keywords"]["unavailable_reason"])
        updated = self.client.put(f"/api/jobs/{job['id']}", json={
            "paste_description": "Requirements:\n3 to 5+ years of production backend experience. Advanced distributed systems. Python and PostgreSQL. Preferred:\nKafka and Kubernetes.",
            "required_skills": "3 to 5+ years of production backend experience, advanced distributed systems, Python, PostgreSQL",
            "preferred_skills": "Kafka, Kubernetes",
        })
        self.assertEqual(updated.status_code, 200)
        self.assertIn("production backend experience", updated.get_json()["job_description"])
        complete = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertTrue(complete["keywords"]["available"])
        self.assertGreater(complete["keywords"]["total"], 1)
        self.assertTrue(any(item["normalized"] == "messaging and queue platforms" for item in complete["keywords"]["missing"]))
        override = self.client.post(f"/api/tailoring/{job['id']}/requirement-overrides", json={"normalized_requirement": "messaging and queue platforms", "override": "not_a_requirement"})
        self.assertEqual(override.status_code, 200)
        adjusted = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()["keywords"]
        self.assertEqual(adjusted["excluded_count"], 1)
        self.assertEqual(adjusted["adjusted_total"], adjusted["total"] - 1)
        self.assertTrue(any(item.get("excluded_from_adjusted_coverage") for item in adjusted["missing"]))
        restored = self.client.delete(f"/api/tailoring/{job['id']}/requirement-overrides/messaging%20and%20queue%20platforms")
        self.assertEqual(restored.status_code, 200)

    def test_tailoring_ai_structured_proposal_is_extracted_without_dictionary_text(self):
        self.client.post("/api/profile", json={"name": "Terry Mak", "email": "terry@example.com", "resume_text": "Built Python automation."})
        profile = self.client.post("/api/candidate-profiles", json={"name": "Software Engineering", "identity_mode": "shared"}).get_json()
        document = Document()
        document.add_paragraph("Built Python automation for reporting.", style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(f"/api/candidate-profiles/{profile['id']}/resumes/upload", data={"file": (io.BytesIO(buffer.getvalue()), "base.docx"), "is_default": "true"}, content_type="multipart/form-data").get_json()
        job = self.client.post("/api/jobs", json={"company": "Example", "title": "Developer", "job_description": "Python automation"}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        with patch.object(app_module, "request_json_completion", return_value={"ok": True, "data": {"proposed_bullet": "Built Python automation for reporting workflows."}}):
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=app_module.OllamaConfig(True, "http://ollama", "qwen3:4b", 10, 100, 1000)):
                response = self.wait_for_tailoring_ai(self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": "Built Python automation for reporting."}))
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["proposed_text"], "Built Python automation for reporting workflows.")
        self.assertNotIn("proposed_bullet", response.get_json()["proposed_text"])

    def test_tailoring_ai_improve_numeric_safety_retry_and_successful_diff(self):
        self.client.post("/api/profile", json={"name": "Terry Mak", "email": "terry@example.com", "resume_text": "Built Flowtally services."})
        profile = self.client.post("/api/candidate-profiles", json={"name": "Software Engineering", "identity_mode": "shared"}).get_json()
        document = Document()
        document.add_paragraph("Built Flowtally services.", style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "base.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Example", "title": "Developer", "job_description": "Python and REST APIs"}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        config = app_module.OllamaConfig(True, "http://ollama", "qwen3:4b", 10, 100, 1000)
        with patch.object(app_module, "compare_docx_layout_and_export", return_value={"available": False, "verified": False, "reason": "Microsoft Word unavailable"}):
                preview_response = self.client.get(f"/api/tailoring/drafts/{workspace['draft']['id']}/preview")
        self.assertEqual(preview_response.status_code, 200, preview_response.get_json())
        self.assertFalse(preview_response.get_json()["pdf_available"])
        self.assertIn("Microsoft Word unavailable", preview_response.get_json()["pdf_reason"])

        def improve(original, response_values):
            with patch.object(app_module, "request_json_completion", side_effect=response_values) as request_mock:
                with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                    response = self.wait_for_tailoring_ai(self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original, "keyword": "REST APIs"}))
            return response, request_mock

        response, request_mock = improve(
            "Architected REST APIs for the Flowtally platform.",
            [{"ok": True, "data": {"proposed_bullet": "Designed REST APIs for the Flowtally platform."}}],
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["proposed_text"], "Designed REST APIs for the Flowtally platform.")
        self.assertEqual(request_mock.call_count, 1)

        response, _ = improve(
            "Built React 18 services with 75% test coverage.",
            [{"ok": True, "data": {"proposed_bullet": "Developed React 18 services with 75% test coverage."}}],
        )
        self.assertEqual(response.status_code, 200, response.get_json())

        response, _ = improve(
            "Built 25+ tests and reduced errors by 75% across 3 services.",
            [{"ok": True, "data": {"proposed_bullet": "Added 25+ tests and reduced errors by 75% across 3 services."}}],
        )
        self.assertEqual(response.status_code, 200, response.get_json())

        response, request_mock = improve(
            "Built 25+ tests and reduced errors by 75% across 3 services.",
            [
                {"ok": True, "data": {"proposed_bullet": "Built 20+ tests and reduced errors by 75% across 3 services."}},
                {"ok": True, "data": {"proposed_bullet": "Added 25+ tests and reduced errors by 75% across 3 services."}},
            ],
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertTrue(response.get_json()["retry_attempted"])
        self.assertEqual(request_mock.call_count, 2)

        response, request_mock = improve(
            "Built 25+ tests and reduced errors by 75% across 3 services.",
            [
                {"ok": True, "data": {"proposed_bullet": "Built 20+ tests and reduced errors by 75% across 3 services."}},
                {"ok": True, "data": {"proposed_bullet": "Added 25+ tests and reduced errors by 60% across 3 services."}},
            ],
        )
        self.assertEqual(response.status_code, 422, response.get_json())
        self.assertTrue(response.get_json()["retry_attempted"])
        self.assertEqual(response.get_json()["original_text"], "Built 25+ tests and reduced errors by 75% across 3 services.")
        self.assertEqual(response.get_json()["proposed_text"], response.get_json()["original_text"])
        self.assertIn("rejected after one safety retry", response.get_json()["error"])
        self.assertEqual(request_mock.call_count, 2)

    def test_tailoring_ai_rejects_cross_project_claims_omissions_and_long_rewrites(self):
        original = "Architected a React and TypeScript frontend with a Flask and PostgreSQL backend, implementing REST APIs, Google OAuth, role-based access, tenant isolation, audit trails, and database migrations."
        cross_project = "Designed and implemented REST APIs for product pricing and inventory systems, using Flask and PostgreSQL to support Google Docs, Amazon, and idempotent operations."
        omission = "Designed Flask and PostgreSQL APIs for product pricing."
        long_rewrite = original + " This additional wording is unnecessarily long and should remain review only."
        self.assertIn("unsupported résumé claims", app_module.validate_tailoring_bullet_suggestion(original, cross_project))
        self.assertIn("omitted supported résumé facts", app_module.validate_tailoring_bullet_suggestion(original, omission))
        self.assertIn("target length budget", app_module.validate_tailoring_bullet_suggestion(original, long_rewrite))

    def test_tailoring_ai_improve_flowtally_api_bullet_rejects_trivial_article_deletion(self):
        self.client.post("/api/profile", json={"name": "Terry Mak", "email": "terry@example.com", "resume_text": "Flowtally"})
        profile = self.client.post("/api/candidate-profiles", json={"name": "Software Engineering", "identity_mode": "shared"}).get_json()
        document = Document()
        original = "Architected a React and TypeScript frontend with a Flask and PostgreSQL backend, implementing REST APIs, Google OAuth, role-based access, tenant isolation, audit trails, and database migrations."
        document.add_paragraph(original, style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "flowtally.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Lakeview", "title": "Backend Engineer", "job_description": "Design APIs and services"}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        config = app_module.OllamaConfig(True, "http://ollama", "qwen3:4b", 10, 100, 1000)
        trivial = "Architected React and TypeScript frontend with Flask and PostgreSQL backend, implementing REST APIs, Google OAuth, role-based access, tenant isolation, audit trails, and database migrations."
        meaningful = "Architected and implemented REST APIs with a React and TypeScript frontend and a Flask and PostgreSQL backend, preserving Google OAuth, role-based access, tenant isolation, audit trails, and database migrations."
        with patch.object(app_module, "request_json_completion", side_effect=[{"ok": True, "data": {"proposed_bullet": trivial}}, {"ok": True, "data": {"proposed_bullet": original}}]):
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                rejected_trivial = self.wait_for_tailoring_ai(self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original, "keyword": "API and service design"}))
        self.assertEqual(rejected_trivial.status_code, 200, rejected_trivial.get_json())
        self.assertEqual(rejected_trivial.get_json()["message"], "The existing bullet is already suitable.")
        self.assertFalse(rejected_trivial.get_json()["meaningful_improvement"])
        self.assertEqual(rejected_trivial.get_json()["proposed_text"], original)

        with patch.object(app_module, "request_json_completion", side_effect=[{"ok": True, "data": {"proposed_bullet": trivial}}, {"ok": True, "data": {"proposed_bullet": meaningful}}]):
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                improved = self.wait_for_tailoring_ai(self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original, "keyword": "API and service design"}))
        self.assertEqual(improved.status_code, 200, improved.get_json())
        self.assertFalse(improved.get_json()["meaningful_improvement"])
        self.assertEqual(improved.get_json()["proposed_text"], original)

    def test_approved_tailoring_draft_is_immutable_and_new_revision_copies_artifact(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Backend", "identity_mode": "distinct", "applicant_name": "Alex Rivera"}).get_json()
        resume_doc = Document(); resume_doc.add_paragraph("Alex Rivera"); resume_doc.add_paragraph("Experience"); resume_doc.add_paragraph("Built Python automation for reports.", style="List Bullet")
        resume_buffer = io.BytesIO(); resume_doc.save(resume_buffer)
        resume = self.client.post(f"/api/candidate-profiles/{profile['id']}/resumes/upload", data={"file": (io.BytesIO(resume_buffer.getvalue()), "base.docx"), "is_default": "true"}, content_type="multipart/form-data").get_json()
        cover_doc = Document(); cover_doc.add_paragraph("Alex Rivera"); cover_buffer = io.BytesIO(); cover_doc.save(cover_buffer)
        cover = self.client.post(f"/api/candidate-profiles/{profile['id']}/cover-letter-templates/upload", data={"file": (io.BytesIO(cover_buffer.getvalue()), "cover.docx"), "is_default": "true"}, content_type="multipart/form-data").get_json()
        job = self.client.post("/api/jobs", json={"company": "Package Co", "title": "Developer", "required_skills": "Python", "job_description": "Python"}).get_json()
        package = self.client.post("/api/preparation/queue", json={"job_id": job["id"], "candidate_profile_id": profile["id"], "resume_id": resume["id"], "cover_letter_template_id": cover["id"]}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        draft_id = workspace["draft"]["id"]
        bullet = next(item for item in workspace["bullets"] if item["editable"])
        self.assertEqual(self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": bullet["paragraph_index"], "original_text": bullet["text"], "proposed_text": "Built Python and REST API automation for reports.", "operation": "edit"}).status_code, 200)
        with patch.object(app_module, "validate_tailoring_draft_layout", return_value={"available": True, "verified": True, "original_pages": 1, "candidate_pages": 1, "line_mismatches": []}):
            approved = self.client.post(f"/api/tailoring/drafts/{draft_id}/approve", json={"preparation_id": package["id"]})
        self.assertEqual(approved.status_code, 200, approved.get_json())
        approved_path = Path(approved.get_json()["draft"]["approved_tailored_resume_path"])
        approved_bytes = approved_path.read_bytes()
        locked_workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertEqual(locked_workspace["draft"]["status"], "approved")
        self.assertTrue(locked_workspace["draft"]["locked"])
        self.assertEqual(self.client.post(f"/api/tailoring/drafts/{draft_id}/apply", json={"paragraph_index": bullet["paragraph_index"], "original_text": bullet["text"], "proposed_text": "Unsafe change", "operation": "edit"}).status_code, 409)
        created = self.client.post(f"/api/tailoring/drafts/{draft_id}/new-revision", json={})
        self.assertEqual(created.status_code, 201, created.get_json())
        revision = created.get_json()["draft"]
        self.assertNotEqual(revision["id"], draft_id)
        self.assertEqual(revision["revision_number"], 2)
        self.assertEqual(Path(revision["draft_path"]).read_bytes(), approved_bytes)
        repeated = self.client.post(f"/api/tailoring/drafts/{draft_id}/new-revision", json={})
        self.assertEqual(repeated.status_code, 200, repeated.get_json())
        self.assertEqual(repeated.get_json()["draft"]["id"], revision["id"])
        current = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        self.assertEqual(current["draft"]["id"], revision["id"])
        revision_bullet = next(item for item in current["bullets"] if item["editable"])
        changed = self.client.post(f"/api/tailoring/drafts/{revision['id']}/apply", json={"paragraph_index": revision_bullet["paragraph_index"], "original_text": revision_bullet["text"], "proposed_text": "Built Python and REST API automation for reviewed reports.", "operation": "edit"})
        self.assertEqual(changed.status_code, 200, changed.get_json())
        package_after_edit = self.client.get(f"/api/preparation/queue").get_json()
        package_after_edit = next(item for item in package_after_edit if int(item["id"]) == int(package["id"]))
        self.assertEqual(package_after_edit["approved_tailoring_draft_id"], draft_id)
        self.assertEqual(Path(package_after_edit["approved_tailored_resume_path"]).read_bytes(), approved_bytes)
        self.assertEqual(approved_path.read_bytes(), approved_bytes)

    def test_approved_tailoring_becomes_package_resume_and_survives_cover_regeneration(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Backend", "identity_mode": "distinct", "applicant_name": "Alex Rivera"}).get_json()
        resume_doc = Document(); resume_doc.add_paragraph("Alex Rivera"); resume_doc.add_paragraph("Experience"); resume_doc.add_paragraph("Built Python automation for reports.", style="List Bullet")
        resume_buffer = io.BytesIO(); resume_doc.save(resume_buffer)
        resume = self.client.post(f"/api/candidate-profiles/{profile['id']}/resumes/upload", data={"file": (io.BytesIO(resume_buffer.getvalue()), "base.docx"), "is_default": "true"}, content_type="multipart/form-data").get_json()
        cover_doc = Document(); cover_doc.add_paragraph("Alex Rivera"); cover_doc.add_paragraph("I am applying to Package Co.")
        cover_buffer = io.BytesIO(); cover_doc.save(cover_buffer)
        cover = self.client.post(f"/api/candidate-profiles/{profile['id']}/cover-letter-templates/upload", data={"file": (io.BytesIO(cover_buffer.getvalue()), "cover.docx"), "is_default": "true"}, content_type="multipart/form-data").get_json()
        job = self.client.post("/api/jobs", json={"company": "Package Co", "title": "Developer", "overall_score": 90, "remote_eligibility": "canadian_eligible", "required_skills": "Python, REST APIs", "job_description": "Python and REST APIs"}).get_json()
        item = self.client.post("/api/preparation/queue", json={"job_id": job["id"], "candidate_profile_id": profile["id"], "resume_id": resume["id"], "cover_letter_template_id": cover["id"]}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        bullet = next(entry for entry in workspace["bullets"] if entry["editable"])
        applied = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/apply", json={"paragraph_index": bullet["paragraph_index"], "keyword": "REST APIs", "original_text": bullet["text"], "proposed_text": "Built Python automation and REST APIs for reports.", "operation": "edit"})
        self.assertEqual(applied.status_code, 200, applied.get_json())
        self.assertEqual(self.client.get(f"/api/candidate-profiles/{profile['id']}/keyword-memory").get_json(), [])
        with patch.object(app_module, "compare_docx_layout", return_value={"available": True, "verified": True, "original_pages": 1, "candidate_pages": 1, "line_mismatches": []}):
            approved = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/approve", json={"preparation_id": item["id"]})
        self.assertEqual(approved.status_code, 200, approved.get_json())
        memory = self.client.get(f"/api/candidate-profiles/{profile['id']}/keyword-memory").get_json()
        self.assertEqual(memory[0]["approved_replacement"], "Built Python automation and REST APIs for reports.")
        tailored_path = Path(approved.get_json()["preparation"]["approved_tailored_resume_path"])
        self.assertTrue(tailored_path.is_file())
        cover_path = app_module.generated_applications_dir() / str(job["id"]) / "cover-output.docx"; cover_path.parent.mkdir(parents=True, exist_ok=True); cover_path.write_bytes(cover_buffer.getvalue())
        def fake_prepare(job_id, kinds, profile_override):
            with app_module.app.app_context():
                app_module.get_db().execute("UPDATE jobs SET resume_generation_status = 'generated', cover_letter_generation_status = 'generated', cover_letter_docx_path = ? WHERE id = ?", [str(cover_path), job_id]); app_module.get_db().commit()
            return {"resume": "generated", "cover_letter": "generated", "error": ""}
        with patch.object(app_module, "generate_and_store_shortlist_documents", side_effect=fake_prepare):
            prepared = self.client.post(f"/api/preparation/{item['id']}/prepare")
        self.assertEqual(prepared.status_code, 200, prepared.get_json())
        package = prepared.get_json()
        self.assertEqual(package["status"], "ready_for_review", package)
        self.assertEqual(Path(package["prepared_resume_path"]), tailored_path)
        self.assertIn("rest apis", [item["normalized"] for item in package["keywords_score"].get("matched", [])], {"score": package["keywords_score"], "text": app_module.extract_docx_text(tailored_path.read_bytes()), "job": package["job"].get("required_skills")})
        self.assertIn("REST APIs", app_module.extract_docx_text(tailored_path.read_bytes()))
        tailored_path.unlink()
        unavailable = self.client.post(f"/api/preparation/{item['id']}/regenerate")
        self.assertEqual(unavailable.status_code, 200)
        self.assertEqual(unavailable.get_json()["status"], "failed")
        self.assertIn("approved tailored résumé", unavailable.get_json()["error"])
        second_job = self.client.post("/api/jobs", json={"company": "Second Co", "title": "Backend Developer", "required_skills": "REST APIs", "job_description": "REST APIs"}).get_json()
        second_workspace = self.client.get(f"/api/tailoring/{second_job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        reusable = next(item for group in ("matched", "related", "missing") for item in second_workspace["keywords"][group] if item["normalized"] == "api and service design")
        self.assertTrue(reusable["memory"]["source_matches"])
        reused = self.client.post(f"/api/tailoring/drafts/{second_workspace['draft']['id']}/apply-memory", json={"keyword": "REST APIs", "paragraph_index": reusable["memory"]["matching_paragraph_index"]})
        self.assertEqual(reused.status_code, 200, reused.get_json())
        self.assertIn("REST APIs", reused.get_json()["text"])

    def test_tailoring_keyword_memory_is_profile_scoped_and_editable(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Backend"}).get_json()
        saved = self.client.post(f"/api/candidate-profiles/{profile['id']}/keyword-memory", json={"keyword": "Kafka", "status": "yes", "resume_id": 7, "project_employer": "Flowtally", "supporting_bullet": "Built Kafka integrations", "source_fingerprint": "abc", "approved_replacement": "Built Kafka integrations for reporting"})
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.get_json()["status"], "yes")
        updated = self.client.post(f"/api/candidate-profiles/{profile['id']}/keyword-memory", json={"keyword": "Kafka", "status": "no"})
        self.assertEqual(updated.get_json()["project_employer"], "Flowtally")
        self.assertEqual(updated.get_json()["resume_id"], 7)
        listed = self.client.get(f"/api/candidate-profiles/{profile['id']}/keyword-memory").get_json()
        self.assertEqual(len(listed), 1)
        self.assertEqual(self.client.delete(f"/api/candidate-profiles/{profile['id']}/keyword-memory/kafka").status_code, 200)
        self.assertEqual(self.client.get(f"/api/candidate-profiles/{profile['id']}/keyword-memory").get_json(), [])
        self.assertIn('Checking requirements', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('Configure candidate profile documents', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('Candidate profile documents are separate from the legacy Profile documents', Path('static/app.js').read_text(encoding='utf-8'))
        self.assertIn('z-index: 60', Path('static/styles.css').read_text(encoding='utf-8'))

    def test_skill_confirmation_memory_is_individual_profile_scoped_and_does_not_change_coverage(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Backend"}).get_json()
        document = Document(); document.add_paragraph("Built Python services with REST APIs.", style="List Bullet")
        buffer = io.BytesIO(); document.save(buffer)
        resume = self.client.post(f"/api/candidate-profiles/{profile['id']}/resumes/upload", data={"file": (io.BytesIO(buffer.getvalue()), "base.docx"), "is_default": "true"}, content_type="multipart/form-data").get_json()
        job = self.client.post("/api/jobs", json={"company": "Example", "title": "Backend", "job_description": "Required\nPython or Java\nPreferred\nKafka or RabbitMQ\nResponsibilities\nBuild REST APIs"}).get_json()
        before = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()["keywords"]
        group = next(item for item in before["missing"] if item["normalized"] == "messaging and queue platforms")
        self.assertEqual(group["skill_memory"], {})
        inventory_names = {item["keyword"] for item in before["skill_inventory"]}
        self.assertTrue({"kafka", "rabbitmq"}.issubset({name.lower() for name in inventory_names}))
        self.assertTrue(all(item["parent_keyword"] == "messaging and queue platforms" for item in before["skill_inventory"] if item["keyword"].lower() in {"kafka", "rabbitmq"}))
        saved = self.client.post(f"/api/candidate-profiles/{profile['id']}/skill-memory", json={"keyword": "Kafka", "status": "no", "notes": "Not used professionally"})
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.get_json()["confirmation_source"], "skill_confirmation")
        listed = self.client.get(f"/api/candidate-profiles/{profile['id']}/skill-memory").get_json()
        self.assertEqual([(item["normalized_keyword"], item["status"]) for item in listed], [("kafka", "no")])
        after = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()["keywords"]
        group = next(item for item in after["missing"] if item["normalized"] == "messaging and queue platforms")
        self.assertEqual(group["skill_memory"]["kafka"]["status"], "no")
        self.assertEqual(after["matched_count"], before["matched_count"])
        self.client.post(f"/api/candidate-profiles/{profile['id']}/skill-memory", json={"keyword": "RabbitMQ", "status": "yes"})
        group = next(item for item in self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()["keywords"]["missing"] if item["normalized"] == "messaging and queue platforms")
        self.assertEqual(group["skill_memory"]["rabbitmq"]["status"], "yes")
        self.assertEqual(self.client.delete(f"/api/candidate-profiles/{profile['id']}/skill-memory/kafka").status_code, 200)
        self.assertEqual(self.client.get(f"/api/candidate-profiles/{profile['id']}/skill-memory").get_json()[0]["normalized_keyword"], "rabbitmq")
    def test_jobspy_discovery_requires_keywords(self):
        response = self.client.post("/api/discovery/jobspy", json={"location": "Toronto"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("Keywords are required", response.get_json()["error"])

    def test_discovery_location_resolution_normalizes_city_and_postal_inputs(self):
        provider = FakeGeocodingProvider(
            {
                "Toronto, Ontario, Canada": {
                    "latitude": 43.6532,
                    "longitude": -79.3832,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 80,
                    "precision": "city",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "Toronto, Ontario, Canada",
                    "error": "",
                },
                "M2J 0B6, Canada": {
                    "latitude": 43.787,
                    "longitude": -79.353,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 95,
                    "precision": "postal code",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "M2J 0B6, Canada",
                    "error": "",
                },
            }
        )
        city = app_module.resolve_discovery_location_context("Toronto", {}, provider)
        toronto_on = app_module.resolve_discovery_location_context("Toronto, ON", {}, provider)
        postal = app_module.resolve_discovery_location_context("m2j0b6", {}, provider)
        self.assertEqual(city["normalized_input"], "Toronto")
        self.assertEqual(city["distance_origin"], "Toronto")
        self.assertEqual(city["discovery_location"], "Toronto, ON")
        self.assertEqual(toronto_on["normalized_input"], "Toronto, ON")
        self.assertEqual(toronto_on["discovery_location"], "Toronto, ON")
        self.assertEqual(postal["normalized_input"], "M2J 0B6")
        self.assertEqual(postal["distance_origin"], "M2J 0B6")
        self.assertEqual(postal["discovery_location"], "Toronto, ON")
        self.assertEqual(postal["resolution_method"], "postal_code_prefix_fallback")

    def test_canadian_postal_code_normalization_and_cached_profile_geocoding(self):
        self.assertEqual(app_module.normalize_canadian_postal_code("m5v2t6"), "M5V 2T6")
        self.assertEqual(app_module.extract_canadian_postal_code("Suite 400, M5V 2T6"), "M5V 2T6")
        provider = FakeGeocodingProvider(
            {
                "M5V 2T6, Canada": {
                    "ok": True,
                    "status": "saved",
                    "latitude": 43.6426,
                    "longitude": -79.3871,
                    "provider": "fake",
                    "confidence": 95,
                    "precision": "postal code",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "M5V 2T6, Canada",
                    "error": "",
                }
            }
        )
        with app_module.app.app_context():
            app_module.init_db()
            profile, first = app_module.geocode_profile_home_location({"location": "m5v2t6"}, provider)
            self.assertTrue(first["ok"])
            self.assertEqual(profile["home_location_normalized"], "M5V 2T6, Canada")
            self.assertEqual(profile["home_location_precision"], "postal code")
            self.assertEqual(len(provider.calls), 1)

            cached_profile, second = app_module.geocode_profile_home_location(profile, provider)
            self.assertTrue(second["ok"])
            self.assertEqual(cached_profile["home_location_normalized"], "M5V 2T6, Canada")
            self.assertEqual(len(provider.calls), 1)

    def test_invalid_canadian_postal_code_is_rejected(self):
        response = self.client.post("/api/profile", json={"location": "M5V2T"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("Invalid Canadian postal code", response.get_json()["error"])

    def test_api_health_reports_environment_diagnostics(self):
        health = self.client.get("/api/health")
        self.assertEqual(health.status_code, 200)
        body = health.get_json()
        self.assertEqual(body["sqlite_db_path"], str(app_module.DB_PATH))
        self.assertEqual(body["discovery_cache_path"], str(app_module.DB_PATH))
        self.assertEqual(body["schema_version"], 1)
        self.assertIn("python_executable", body)
        self.assertIn("python_version", body)
        self.assertIn("os", body)
        self.assertIn("jobspy_version", body)
        self.assertIn("pandas_version", body)
        self.assertIn("numpy_version", body)
        self.assertIn("cwd", body)

    def test_discovery_returns_jobs_without_profile_origin(self):
        provider = FakeGeocodingProvider(
            {
                "Toronto, Ontario, Canada": {
                    "latitude": 43.6532,
                    "longitude": -79.3832,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 80,
                    "precision": "city",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "Toronto, ON",
                    "error": "",
                },
            }
        )
        fake = {
            "jobs": [
                {
                    "source_job_id": "1",
                    "title": "Software Developer",
                    "company": "Northstar",
                    "location": "Toronto, ON",
                    "source": "linkedin",
                    "job_url": "https://example.com/1",
                    "description": "Python SQL APIs",
                    "posted_date": "2026-06-19",
                }
            ],
            "errors": {},
            "source_counts": {"linkedin": 1},
            "source_diagnostics": {
                "linkedin": {"source": "linkedin", "attempted": True, "success": True, "result_count": 1, "error": "", "skipped_reason": "", "duration_ms": 5, "raw_records_row_count": 1, "normalized_row_count": 1},
                "indeed": {"source": "indeed", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
                "google": {"source": "google", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
                "zip_recruiter": {"source": "zip_recruiter", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
            },
            "source_snapshots": {"linkedin": [{"source": "linkedin", "title": "Software Developer", "company": "Northstar", "location": "Toronto, ON", "date_posted": "2026-06-19", "url": "https://example.com/1", "raw_field_types": {"title": "str"}}]},
        }
        location = {"original_input": "Toronto, Ontario", "normalized_input": "Toronto, Ontario", "distance_origin": "Toronto, Ontario", "discovery_location": "Toronto, ON", "resolution_method": "test", "lookup": {}}
        with patch.object(app_module, "get_geocoding_provider", return_value=provider), patch.object(app_module, "resolve_discovery_location_context", return_value=location), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            response = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "software developer", "location": "Toronto, ON", "enabled_sources": ["linkedin"], "bypass_cache": True},
            )
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["total"], 1)
        self.assertEqual(body["visible_total"], 1)
        self.assertFalse(body["diagnostics"]["origin_available"])
        self.assertEqual(body["bypass_cache"], True)

    def test_discovery_source_smoke_test_reports_each_enabled_source(self):
        provider = FakeGeocodingProvider(
            {
                "Toronto, Ontario, Canada": {
                    "latitude": 43.6532,
                    "longitude": -79.3832,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 80,
                    "precision": "city",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "Toronto, ON",
                    "error": "",
                }
            }
        )
        fake = {
            "jobs": [],
            "errors": {"indeed": "RuntimeError: parser error"},
            "source_counts": {"linkedin": 0, "indeed": 0},
            "source_diagnostics": {
                "linkedin": {"source": "linkedin", "attempted": True, "success": True, "result_count": 0, "error": "", "skipped_reason": "", "duration_ms": 5, "raw_records_row_count": 0, "normalized_row_count": 0},
                "indeed": {"source": "indeed", "attempted": True, "success": False, "result_count": 0, "error": "RuntimeError: parser error", "skipped_reason": "", "duration_ms": 7, "raw_records_row_count": 0, "normalized_row_count": 0},
                "google": {"source": "google", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
                "zip_recruiter": {"source": "zip_recruiter", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
            },
            "source_snapshots": {},
        }
        with patch.object(app_module, "get_geocoding_provider", return_value=provider), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake) as mocked_fetch:
            response = self.client.post(
                "/api/discovery/jobspy/test-sources",
                json={"enabled_sources": ["linkedin", "indeed"], "bypass_cache": True},
            )
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["location_context"]["discovery_location"], "Toronto, ON")
        self.assertEqual(body["source_diagnostics"]["linkedin"]["attempted"], True)
        self.assertEqual(body["source_diagnostics"]["indeed"]["attempted"], True)
        self.assertEqual(body["source_diagnostics"]["indeed"]["error"], "RuntimeError: parser error")
        self.assertEqual(mocked_fetch.call_args.kwargs["location"], "Toronto, ON")
        self.assertEqual(body["source_counts"]["linkedin"], 0)

    def test_discovery_debug_latest_endpoint_returns_latest_run(self):
        provider = FakeGeocodingProvider(
            {
                "Toronto, Ontario, Canada": {
                    "latitude": 43.6532,
                    "longitude": -79.3832,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 80,
                    "precision": "city",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "Toronto, ON",
                    "error": "",
                },
            }
        )
        fake = {
            "jobs": [],
            "errors": {},
            "source_counts": {"linkedin": 0},
            "source_diagnostics": {
                "linkedin": {"source": "linkedin", "attempted": True, "success": True, "result_count": 0, "error": "", "skipped_reason": "", "duration_ms": 4, "raw_records_row_count": 0, "normalized_row_count": 0},
                "indeed": {"source": "indeed", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
                "google": {"source": "google", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
                "zip_recruiter": {"source": "zip_recruiter", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
            },
            "source_snapshots": {},
        }
        with patch.object(app_module, "get_geocoding_provider", return_value=provider), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "software developer", "location": "Toronto, ON", "enabled_sources": ["linkedin"], "bypass_cache": True},
            )
        latest = self.client.get("/api/discovery/debug/latest")
        self.assertEqual(latest.status_code, 200)
        body = latest.get_json()
        self.assertIn("source_snapshots", body)
        self.assertIn("metadata", body)
        self.assertEqual(body["metadata"]["jobs_returned_to_frontend"], 0)

    def test_toronto_city_and_job_distance_metadata_are_geocoded(self):
        provider = FakeGeocodingProvider(
            {
                "Toronto, Ontario, Canada": {
                    "latitude": 43.6532,
                    "longitude": -79.3832,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 80,
                    "precision": "city",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "Toronto, Ontario, Canada",
                    "error": "",
                },
                "M5V 2T6, Canada": {
                    "latitude": 43.6426,
                    "longitude": -79.3871,
                    "display_name": "M5V 2T6, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 95,
                    "precision": "postal code",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "M5V 2T6, Canada",
                    "error": "",
                },
            }
        )
        with patch.object(app_module, "get_geocoding_provider", return_value=provider):
            profile_response = self.client.post("/api/profile", json={"location": "Toronto"})
        self.assertEqual(profile_response.status_code, 200)
        profile = profile_response.get_json()
        self.assertEqual(profile["home_location_normalized"], "Toronto, ON")
        self.assertEqual(profile["home_location_precision"], "city")
        self.assertEqual(profile["location_result"]["status"], "saved")

        fake = {
            "jobs": [
                {
                    "source_job_id": "job-1",
                    "title": "Software Developer",
                    "company": "Northstar",
                    "location": "M5V2T6",
                    "source": "linkedin",
                    "job_url": "https://example.com/job/1",
                    "description": "Python SQL APIs",
                    "posted_date": date(2026, 6, 19),
                }
            ],
            "errors": {},
            "source_counts": {"linkedin": 1},
        }
        with patch.object(app_module, "get_geocoding_provider", return_value=provider), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            response = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "software developer", "location": "Toronto", "enabled_sources": ["linkedin"]},
            )

        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["total"], 1)
        job = body["jobs"][0]
        self.assertEqual(job["normalized_location"], "M5V 2T6, Canada")
        self.assertEqual(job["distance_label"], "approximate")
        self.assertTrue(job["distance_note"].startswith("Approx. distance from Toronto, ON"))
        self.assertGreaterEqual(job["distance_km"], 0)

    def test_discovery_jobspy_accepts_date_objects_in_raw_payloads(self):
        provider = FakeGeocodingProvider(
            {
                "Toronto, Ontario, Canada": {
                    "latitude": 43.6532,
                    "longitude": -79.3832,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 80,
                    "precision": "city",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "Toronto, Ontario, Canada",
                    "error": "",
                }
            }
        )
        fake = {
            "jobs": [
                {
                    "source_job_id": "job-serial",
                    "title": "Backend Developer",
                    "company": "SerialCo",
                    "location": "Toronto",
                    "source": "linkedin",
                    "job_url": "https://example.com/job/serial",
                    "description": "Python Flask SQL",
                    "posted_date": date(2026, 6, 19),
                }
            ],
            "errors": {},
            "source_counts": {"linkedin": 1},
        }
        with patch.object(app_module, "get_geocoding_provider", return_value=provider), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            response = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "backend developer", "location": "Toronto", "enabled_sources": ["linkedin"]},
            )

        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["jobs"][0]["posted_date"], "2026-06-19")
        self.assertEqual(body["visible_total"], 1)

    def test_discovery_jobspy_uses_resolved_search_location_and_reports_source_diagnostics(self):
        provider = FakeGeocodingProvider(
            {
                "Toronto, Ontario, Canada": {
                    "latitude": 43.6532,
                    "longitude": -79.3832,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 80,
                    "precision": "city",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "Toronto, Ontario, Canada",
                    "error": "",
                },
                "M2J 0B6, Canada": {
                    "latitude": 43.787,
                    "longitude": -79.353,
                    "display_name": "Toronto, Ontario, Canada",
                    "ok": True,
                    "status": "saved",
                    "provider": "fake",
                    "confidence": 95,
                    "precision": "postal code",
                    "geocoded_at": "2026-06-20T10:00:00",
                    "normalized_query": "M2J 0B6, Canada",
                    "error": "",
                },
            }
        )
        calls: list[dict[str, object]] = []

        def fake_fetch_jobspy_jobs(**kwargs):
            calls.append(kwargs)
            return {
                "jobs": [],
                "errors": {},
                "source_counts": {"linkedin": 0, "indeed": 0},
                "source_diagnostics": {
                    "linkedin": {"source": "linkedin", "attempted": True, "success": True, "result_count": 0, "error": "", "skipped_reason": "", "duration_ms": 5, "effective_country_indeed": "Canada"},
                    "indeed": {"source": "indeed", "attempted": True, "success": True, "result_count": 0, "error": "", "skipped_reason": "", "duration_ms": 7, "effective_country_indeed": "Canada"},
                    "google": {"source": "google", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
                    "zip_recruiter": {"source": "zip_recruiter", "attempted": False, "success": False, "result_count": 0, "error": "", "skipped_reason": "disabled", "duration_ms": 0},
                },
            }

        with patch.object(app_module, "get_geocoding_provider", return_value=provider), patch.object(app_module, "fetch_jobspy_jobs", side_effect=fake_fetch_jobspy_jobs):
            response = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "Registered Early Childhood Educator", "location": "M2J0B6", "distance": 25, "enabled_sources": ["linkedin", "indeed"], "date_window": "30d"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(calls)
        self.assertEqual(calls[0]["location"], "Toronto, ON")
        self.assertEqual(calls[0]["country_indeed"], "Canada")
        body = response.get_json()
        self.assertEqual(body["search"]["original_location"], "M2J0B6")
        self.assertEqual(body["search"]["normalized_location"], "M2J 0B6")
        self.assertEqual(body["search"]["distance_origin"], "M2J 0B6")
        self.assertEqual(body["search"]["discovery_location"], "Toronto, ON")
        self.assertIn("source_diagnostics", body["diagnostics"])
        self.assertIn("indeed", body["diagnostics"]["source_diagnostics"])
        self.assertTrue(body["diagnostics"]["source_diagnostics"]["indeed"]["attempted"])

    def test_jobspy_source_diagnostics_capture_success_failure_and_disabled_sources(self):
        def fake_scrape_jobs(**kwargs):
            if kwargs["site_name"] == "linkedin":
                return [
                    {"title": "Role", "company": "A", "location": "Toronto, ON", "date_posted": "2026-06-19"}
                ]
            raise RuntimeError("parser error")

        linkedin_only = jobspy_source.fetch_jobspy_jobs(
            keywords="developer",
            location="Toronto, ON",
            distance=10,
            results_wanted=5,
            hours_old=72,
            enabled_sources=["linkedin"],
            country_indeed="Canada",
            scrape_jobs_func=fake_scrape_jobs,
        )
        self.assertEqual(linkedin_only["source_counts"]["linkedin"], 1)
        self.assertTrue(linkedin_only["source_diagnostics"]["linkedin"]["attempted"])
        self.assertFalse(linkedin_only["source_diagnostics"]["indeed"]["attempted"])
        self.assertEqual(linkedin_only["source_diagnostics"]["indeed"]["skipped_reason"], "disabled")

        indeed_only = jobspy_source.fetch_jobspy_jobs(
            keywords="developer",
            location="Toronto, ON",
            distance=10,
            results_wanted=5,
            hours_old=72,
            enabled_sources=["indeed"],
            country_indeed="Canada",
            scrape_jobs_func=fake_scrape_jobs,
        )
        self.assertTrue(indeed_only["source_diagnostics"]["linkedin"]["skipped_reason"])
        self.assertTrue(indeed_only["source_diagnostics"]["indeed"]["attempted"])
        self.assertFalse(indeed_only["source_diagnostics"]["indeed"]["success"])
        self.assertIn("parser error", indeed_only["source_diagnostics"]["indeed"]["error"])

        result = jobspy_source.fetch_jobspy_jobs(
            keywords="developer",
            location="Toronto, ON",
            distance=10,
            results_wanted=5,
            hours_old=72,
            enabled_sources=["linkedin", "indeed"],
            country_indeed="Canada",
            scrape_jobs_func=fake_scrape_jobs,
        )
        self.assertEqual(result["source_counts"]["linkedin"], 1)
        self.assertIn("indeed", result["errors"])
        self.assertTrue(result["source_diagnostics"]["linkedin"]["attempted"])
        self.assertTrue(result["source_diagnostics"]["linkedin"]["success"])
        self.assertEqual(result["source_diagnostics"]["linkedin"]["result_count"], 1)
        self.assertTrue(result["source_diagnostics"]["indeed"]["attempted"])
        self.assertFalse(result["source_diagnostics"]["indeed"]["success"])
        self.assertIn("parser error", result["source_diagnostics"]["indeed"]["error"])
        self.assertFalse(result["source_diagnostics"]["google"]["attempted"])
        self.assertEqual(result["source_diagnostics"]["google"]["skipped_reason"], "disabled")

    def test_jobspy_discovery_scores_saved_profile_and_handles_missing_fields(self):
        self.client.post(
            "/api/profile",
            json={
                "preferred_roles": "software developer",
                "preferred_locations": "Toronto, remote",
                "key_skills": "Python, Flask, React, SQL, APIs",
                "resume_text": "Built Flask APIs and React dashboards with SQL.",
            },
        )
        fake = {
            "jobs": [
                {
                    "source_job_id": "abc-1",
                    "title": "Junior Software Developer",
                    "company": "",
                    "location": "Toronto, Ontario",
                    "source": "linkedin",
                    "job_url": "https://example.com/job/1?utm_source=x",
                    "application_url": "https://apply.example.com/job/1",
                    "description": "Python Flask React SQL APIs dashboards",
                    "salary_text": "$70,000-$85,000",
                    "remote_type": "Hybrid",
                    "posted_date": "2026-06-10",
                }
            ],
            "errors": {},
            "source_counts": {"linkedin": 1, "indeed": 0},
        }
        with patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            response = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "software developer", "location": "Toronto, Ontario", "enabled_sources": ["linkedin"]},
            )

        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body["total"], 1)
        self.assertEqual(body["visible_total"], 1)
        job = body["jobs"][0]
        self.assertEqual(job["company"], "Unknown company")
        self.assertGreaterEqual(job["fit_score"], 70)
        self.assertIn("components", job["score_breakdown"])
        self.assertIn("overall_score", job["score_breakdown"])
        self.assertTrue(job["description_available"])

    def test_jobspy_discovery_keeps_source_errors_and_counts(self):
        fake = {
            "jobs": [
                {
                    "source_job_id": "indeed-1",
                    "title": "Software Developer",
                    "company": "Northstar",
                    "location": "Toronto",
                    "source": "indeed",
                    "job_url": "https://indeed.example/job/1",
                    "description": "Python SQL APIs",
                }
            ],
            "errors": {"linkedin": "TimeoutError: source timed out"},
            "source_counts": {"linkedin": 0, "indeed": 1},
        }
        with patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            response = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "software developer", "enabled_sources": ["linkedin", "indeed"]},
            )

        body = response.get_json()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(body["source_counts"]["indeed"], 1)
        self.assertEqual(body["source_counts"]["linkedin"], 0)
        self.assertIn("linkedin", body["errors"])
        self.assertEqual(body["total"], 1)
        self.assertEqual(body["visible_total"], 1)

    def test_jobspy_discovery_deduplicates_existing_jobs(self):
        self.client.post(
            "/api/jobs",
            json={
                "company": "Acme",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/jobs/123",
                "job_description": "Python SQL",
            },
        )
        fake = {
            "jobs": [
                {
                    "source_job_id": "1",
                    "title": "Software Developer",
                    "company": "Acme",
                    "location": "Toronto",
                    "source": "linkedin",
                    "job_url": "https://example.com/jobs/123?utm_source=linkedin",
                    "description": "Python SQL",
                },
                {
                    "source_job_id": "2",
                    "title": "Software Developer",
                    "company": "Acme",
                    "location": "Toronto",
                    "source": "indeed",
                    "job_url": "https://different.example/jobs/456",
                    "description": "Python SQL",
                },
            ],
            "errors": {},
            "source_counts": {"linkedin": 1, "indeed": 1},
        }
        with patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            response = self.client.post("/api/discovery/jobspy", json={"keywords": "software developer"})

        body = response.get_json()
        self.assertEqual(body["total"], 2)
        self.assertEqual(body["visible_total"], 1)
        self.assertEqual(body["duplicates"], 1)
        self.assertEqual(body["diagnostics"]["already_saved_matches"], 1)
        self.assertTrue(body["jobs"][0]["already_saved"])

    def test_discovered_result_can_save_to_tracker_and_prevent_future_duplicate_button(self):
        fake = {
            "jobs": [
                {
                    "source_job_id": "zip-1",
                    "title": "Cloud Support Analyst",
                    "company": "OpsTools",
                    "location": "Remote Canada",
                    "source": "zip_recruiter",
                    "job_url": "https://zip.example/job/1",
                    "application_url": "https://zip.example/apply/1",
                    "description": "Azure SQL APIs technical support automation",
                    "remote_type": "Remote",
                }
            ],
            "errors": {},
            "source_counts": {"zip_recruiter": 1},
        }
        with patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            discovery = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "cloud support", "enabled_sources": ["zip_recruiter"]},
            ).get_json()
            saved = self.client.post("/api/jobs", json=discovery["jobs"][0])
            second = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "cloud support", "enabled_sources": ["zip_recruiter"]},
            ).get_json()

        self.assertEqual(saved.status_code, 201)
        self.assertEqual(saved.get_json()["title"], "Cloud Support Analyst")
        self.assertTrue(second["jobs"][0]["already_saved"])
        self.assertEqual(second["duplicates"], 0)
        self.assertEqual(second["diagnostics"]["already_saved_matches"], 1)

    def test_saved_discovery_searches_are_named_disabled_by_default_and_editable(self):
        created = self.client.post(
            "/api/discovery/searches",
            json={"name": "Backend Developer", "keywords": "backend developer", "location": "Toronto", "enabled_sources": ["linkedin"], "interval_minutes": 60},
        )
        self.assertEqual(created.status_code, 201)
        search = created.get_json()
        self.assertFalse(search["enabled"])
        self.assertEqual(search["enabled_sources"], ["linkedin"])
        updated = self.client.put(f"/api/discovery/searches/{search['id']}", json={**search, "name": "Cloud Developer", "enabled": True})
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.get_json()["name"], "Cloud Developer")
        self.assertTrue(updated.get_json()["enabled"])

    def test_saved_discovery_run_persists_new_jobs_and_preserves_existing_status(self):
        provider = FakeGeocodingProvider({"Toronto, Ontario, Canada": {"latitude": 43.65, "longitude": -79.38, "display_name": "Toronto, Ontario, Canada", "precision": "city"}})
        fake = {
            "jobs": [{"source_job_id": "saved-1", "title": "Backend Developer", "company": "Acme", "location": "Toronto, Ontario", "source": "linkedin", "job_url": "https://example.com/acme/1", "description": "Python SQL backend", "remote_type": "Hybrid"}],
            "errors": {}, "source_counts": {"linkedin": 1}, "source_diagnostics": {"linkedin": {"success": True}}, "source_snapshots": {},
        }
        search = self.client.post("/api/discovery/searches", json={"name": "Backend", "keywords": "backend developer", "location": "Toronto, Ontario", "enabled_sources": ["linkedin"]}).get_json()
        location = {"original_input": "Toronto, Ontario", "normalized_input": "Toronto, Ontario", "distance_origin": "Toronto, Ontario", "discovery_location": "Toronto, ON", "resolution_method": "test", "lookup": {}}
        with patch.object(app_module, "get_geocoding_provider", return_value=provider), patch.object(app_module, "resolve_discovery_location_context", return_value=location), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            first = self.client.post(f"/api/discovery/searches/{search['id']}/run")
            self.assertEqual(first.status_code, 200, first.get_json())
            self.assertEqual(first.get_json()["persisted"]["new"], 1)
            job = self.client.get("/api/discovery/inbox").get_json()[0]
            self.client.post(f"/api/discovery/inbox/{job['id']}/review")
            second = self.client.post(f"/api/discovery/searches/{search['id']}/run")
        self.assertEqual(second.status_code, 200)
        self.assertEqual(second.get_json()["persisted"]["duplicate"], 1)
        refreshed = self.client.get("/api/discovery/inbox").get_json()[0]
        self.assertEqual(refreshed["discovery_state"], "reviewed")
        self.assertEqual(len(self.client.get("/api/discovery/runs").get_json()), 2)

    def test_discovery_scheduler_is_disabled_until_explicitly_enabled(self):
        from scripts.discovery_worker import run_due_searches

        with app_module.app.app_context():
            self.assertEqual(run_due_searches()["disabled"], 1)
        response = self.client.put("/api/discovery/scheduler", json={"enabled": True})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.get_json()["enabled"])

    def test_geocoding_rate_limit_is_cached_temporarily_and_does_not_retry_candidates(self):
        class RateLimitedProvider:
            name = "fake-nominatim"
            def __init__(self):
                self.calls = 0
            def geocode(self, query, *, country="Canada"):
                self.calls += 1
                raise urllib.error.HTTPError("https://example.test", 429, "Too Many Requests", {}, None)

        provider = RateLimitedProvider()
        with app_module.app.app_context():
            app_module.init_db()
            first = app_module.geocode_location_value("Toronto, Ontario", lookup_type="job", provider=provider)
            second = app_module.geocode_location_value("Toronto, Ontario", lookup_type="job", provider=provider)
        self.assertEqual(first["status"], "rate_limited")
        self.assertEqual(second["status"], "rate_limited")
        self.assertEqual(provider.calls, 1)
        self.assertIn("temporarily rate limited", second["error"])

    def test_discovery_task_status_is_read_only_and_scripts_are_present(self):
        response = self.client.get("/api/discovery/task-status")
        self.assertEqual(response.status_code, 200)
        self.assertIn("task_name", response.get_json())
        install = Path(app_module.BASE_DIR) / "scripts" / "install_discovery_task.ps1"
        uninstall = Path(app_module.BASE_DIR) / "scripts" / "uninstall_discovery_task.ps1"
        runner = Path(app_module.BASE_DIR) / "scripts" / "run_discovery_task.ps1"
        self.assertTrue(install.exists() and uninstall.exists() and runner.exists())
        install_text = install.read_text(encoding="utf-8")
        self.assertIn("Register-ScheduledTask", install_text)
        self.assertIn("[string]$TaskName", install_text)
        self.assertNotIn("schtasks.exe /Create", install_text)
        self.assertIn("Unregister-ScheduledTask", uninstall.read_text(encoding="utf-8"))
        self.assertIn("--once", runner.read_text(encoding="utf-8"))

    def test_remote_eligibility_requires_safe_location_or_explicit_canadian_statement(self):
        cases = [
            ({"remote_type": "Remote", "location": "London, UK"}, "unknown"),
            ({"remote_type": "Remote", "location": "Remote US only"}, "restricted"),
            ({"remote_type": "Remote", "location": "Remote Canada"}, "canadian_eligible"),
            ({"remote_type": "Remote", "location": "Toronto, Ontario"}, "canadian_eligible"),
            ({"remote_type": "Remote", "location": "Global"}, "unknown"),
            ({"remote_type": "Remote", "location": "Toronto, Ontario", "job_description": "Canada excluded from this remote role."}, "restricted"),
            ({"remote_type": "Remote", "location": "London, UK", "job_description": "Remote in Canada eligible."}, "canadian_eligible"),
        ]
        for job, expected in cases:
            self.assertEqual(app_module.discovery_remote_eligibility(job)[0], expected, job)

    def test_saved_search_salary_minimum_keeps_unknown_and_estimated_compensation(self):
        location = {"original_input": "Toronto", "normalized_input": "Toronto", "distance_origin": "Toronto", "discovery_location": "Toronto, ON", "resolution_method": "test", "lookup": {}}
        fake = {
            "jobs": [
                {"source_job_id": "below", "title": "Software Developer", "company": "Below", "location": "Toronto", "source": "linkedin", "job_url": "https://example.com/below", "salary_min": 50000, "salary_max": 60000, "description": "Python"},
                {"source_job_id": "above", "title": "Software Developer", "company": "Above", "location": "Toronto", "source": "linkedin", "job_url": "https://example.com/above", "salary_min": 90000, "salary_max": 100000, "description": "Python"},
                {"source_job_id": "missing", "title": "Software Developer", "company": "Missing", "location": "Toronto", "source": "linkedin", "job_url": "https://example.com/missing", "description": "Python"},
                {"source_job_id": "estimated", "title": "Software Developer", "company": "Estimated", "location": "Toronto", "source": "linkedin", "job_url": "https://example.com/estimated", "salary_text": "Estimated market range", "description": "Python"},
            ],
            "errors": {}, "source_counts": {"linkedin": 4}, "source_diagnostics": {}, "source_snapshots": {},
        }
        search = self.client.post("/api/discovery/searches", json={"name": "Salary test", "keywords": "software developer", "location": "Toronto", "minimum_salary": 80000, "enabled_sources": ["linkedin"]}).get_json()
        with patch.object(app_module, "get_geocoding_provider", return_value=FakeGeocodingProvider({})), patch.object(app_module, "resolve_discovery_location_context", return_value=location), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            response = self.client.post(f"/api/discovery/searches/{search['id']}/run")
        self.assertEqual(response.status_code, 200)
        jobs = response.get_json()["jobs"]
        self.assertEqual({job["company"] for job in jobs}, {"Above", "Missing", "Estimated"})
        self.assertTrue(all(job["discovery_salary_filter_status"] in {"known", "unknown"} for job in jobs))

    def test_manual_and_worker_runs_share_search_lease_and_heartbeat(self):
        from scripts.discovery_worker import run_due_searches

        search = self.client.post("/api/discovery/searches", json={"name": "Lease test", "keywords": "software developer", "enabled": True}).get_json()
        with app_module.app.app_context():
            app_module.init_db()
            self.assertTrue(app_module.acquire_discovery_search_lease(search["id"], "manual-test"))
            self.assertFalse(app_module.acquire_discovery_search_lease(search["id"], "scheduled-test"))
            original_interval = app_module.DISCOVERY_LEASE_RENEW_SECONDS
            app_module.DISCOVERY_LEASE_RENEW_SECONDS = 0.01
            stop, thread = app_module.start_discovery_search_lease_heartbeat(search["id"], "manual-test")
            try:
                import time
                time.sleep(0.04)
                row = app_module.query_db("SELECT expires_at FROM discovery_search_leases WHERE search_id = ?", [search["id"]], one=True)
                self.assertGreater(row["expires_at"], app_module.datetime.now().isoformat(timespec="seconds"))
            finally:
                app_module.DISCOVERY_LEASE_RENEW_SECONDS = original_interval
                stop.set()
                thread.join(timeout=1)
                app_module.release_discovery_search_lease(search["id"], "manual-test")
        with app_module.app.app_context(), patch.object(app_module, "run_discovery_pipeline", return_value={"total": 0, "persisted": {"new": 0, "duplicate": 0}, "errors": {}, "source_diagnostics": {}}):
            app_module.get_db().execute("UPDATE saved_search_profiles SET enabled = 1, next_run_at = '' WHERE id = ?", [search["id"]])
            app_module.get_db().execute("INSERT INTO app_settings (key, value) VALUES ('discovery_scheduler_enabled', '1') ON CONFLICT(key) DO UPDATE SET value = '1'")
            app_module.get_db().commit()
            first = run_due_searches()
            second = run_due_searches()
        self.assertEqual(first["runs"], 1)
        self.assertEqual(second["runs"], 0)

    def test_discovery_salary_normalization_and_hourly_conversion(self):
        salary = app_module.determine_salary(
            {
                "salary_text": "$30 an hour",
                "job_description": "Hourly support work",
            }
        )
        self.assertEqual(salary["salary_interval"], "hourly")
        self.assertEqual(salary["salary_min"], 62400)
        self.assertEqual(salary["salary_max"], 62400)

    def test_discovery_salary_sorting_keeps_unknowns_after_known_when_requested(self):
        jobs = [
            {
                "company": "KnownCo",
                "title": "Role A",
                "salary_min": 90000,
                "salary_max": 100000,
                "salary_midpoint": 95000,
                "salary_listed": 1,
                "distance_km": 10,
                "distance_confidence": 80,
                "overall_score": 82,
            },
            {
                "company": "UnknownCo",
                "title": "Role B",
                "salary_min": 0,
                "salary_max": 0,
                "salary_midpoint": 0,
                "salary_listed": 0,
                "distance_km": 10,
                "distance_confidence": 80,
                "overall_score": 91,
            },
        ]
        ordered = app_module.sort_discovery_jobs(jobs, "salary", "asc", unknown_salary_position="last")
        self.assertEqual(ordered[0]["company"], "KnownCo")
        self.assertEqual(ordered[1]["company"], "UnknownCo")
        ordered_first = app_module.sort_discovery_jobs(jobs, "salary", "asc", unknown_salary_position="first")
        self.assertEqual(ordered_first[0]["company"], "UnknownCo")

    def test_discovery_distance_remote_and_unknown_location_ordering(self):
        jobs = [
            {
                "company": "LocalCo",
                "title": "Role A",
                "distance_km": 5,
                "distance_confidence": 90,
                "remote_type": "",
                "overall_score": 70,
            },
            {
                "company": "RemoteCo",
                "title": "Role B",
                "distance_km": 9999,
                "distance_confidence": 95,
                "remote_type": "Remote",
                "overall_score": 70,
            },
            {
                "company": "UnknownCo",
                "title": "Role C",
                "distance_km": 9999,
                "distance_confidence": 10,
                "remote_type": "",
                "overall_score": 70,
            },
        ]
        remote_first = app_module.sort_discovery_jobs(jobs, "distance", "asc", remote_rule="remote_first")
        self.assertEqual(remote_first[0]["company"], "RemoteCo")
        remote_after_local = app_module.sort_discovery_jobs(jobs, "distance", "asc", remote_rule="remote_after_local")
        self.assertEqual(remote_after_local[-1]["company"], "UnknownCo")

    def test_discovery_newest_first_and_multi_key_sorting(self):
        jobs = [
            {"company": "B", "title": "Role", "posted_date": "2026-06-18", "overall_score": 80},
            {"company": "A", "title": "Role", "posted_date": "2026-06-19", "overall_score": 80},
            {"company": "C", "title": "Role", "posted_date": "2026-06-17", "overall_score": 90},
        ]
        ordered = app_module.sort_discovery_jobs(jobs, "posted_date", "desc", secondary_field="company", secondary_direction="asc")
        self.assertEqual([job["company"] for job in ordered], ["A", "B", "C"])

    def test_discovery_overall_score_and_explanation_totals(self):
        profile = {
            "location": "Toronto",
            "preferred_locations": "Toronto",
            "preferred_roles": "software developer",
            "key_skills": "Python, Flask, React, SQL",
            "resume_text": "Built Python Flask React SQL tools.",
        }
        job = {
            "company": "GroundedCo",
            "title": "Software Developer",
            "location": "Toronto",
            "remote_type": "Hybrid",
            "salary_min": 90000,
            "salary_max": 100000,
            "salary_midpoint": 95000,
            "salary_source": "structured",
            "salary_confidence": 95,
            "salary": "CAD $90,000-$100,000",
            "salary_listed": 1,
            "job_description": "Python Flask React SQL automation dashboards",
            "required_skills": "Python, SQL",
            "preferred_skills": "React",
            "responsibilities": "Build and maintain internal tools.",
            "posted_date": "2026-06-19",
        }
        breakdown = app_module.discovery_score_breakdown(job, profile)
        self.assertGreaterEqual(breakdown["overall_score"], 0)
        self.assertLessEqual(breakdown["overall_score"], 100)
        component_total = sum(item["score"] for item in breakdown["components"])
        penalty_total = sum(item["points"] for item in breakdown["penalties"])
        self.assertEqual(component_total - penalty_total, breakdown["overall_score"])

    def test_discovery_filters_before_sorting(self):
        jobs = [
            {"company": "SkipCo", "title": "Role A", "salary_max": 60000, "distance_km": 5, "overall_score": 90, "remote_type": "Remote", "posted_date": "2026-06-19"},
            {"company": "KeepCo", "title": "Role B", "salary_max": 90000, "distance_km": 10, "overall_score": 85, "remote_type": "Hybrid", "posted_date": "2026-06-19"},
            {"company": "LowerCo", "title": "Role C", "salary_max": 95000, "distance_km": 50, "overall_score": 60, "remote_type": "Hybrid", "posted_date": "2026-06-19"},
        ]
        filters = {"salary_min": 80000, "overall_score_min": 70, "remote_type": "Hybrid"}
        filtered = [job for job in jobs if app_module.apply_discovery_filters(job, filters)]
        ordered = app_module.sort_discovery_jobs(filtered, "salary", "desc")
        self.assertEqual([job["company"] for job in ordered], ["KeepCo"])

    def test_discovery_saved_search_sort_preferences_persist(self):
        search = {
            "keywords": "software developer",
            "location": "Toronto",
            "enabled_sources": ["linkedin"],
            "search_depth": "standard",
            "date_window": "14d",
            "per_source_cap": 100,
            "total_cap": 200,
            "sort": {"field": "salary", "direction": "asc", "secondary_field": "posted_date", "secondary_direction": "desc", "unknown_salary_position": "last", "remote_rule": "remote_after_local"},
            "filters": {"salary_min": 0, "max_distance": 0, "overall_score_min": 0, "posted_within": 0},
        }
        search_key = app_module.discovery_search_signature(search)
        self.client.post(
            "/api/discovery/preferences",
            json={"search_key": search_key, "preferences": app_module.discovery_preferences_payload(search)},
        )
        response = self.client.get(f"/api/discovery/preferences?search_key={search_key}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["preferences"]["sort"]["field"], "salary")

        fake = {
            "jobs": [
                {
                    "source_job_id": "1",
                    "title": "Role A",
                    "company": "A",
                    "location": "Toronto",
                    "source": "linkedin",
                    "job_url": "https://example.com/a",
                    "salary_text": "$120,000",
                    "salary_min": 120000,
                    "salary_max": 120000,
                    "salary_interval": "yearly",
                    "salary_currency": "CAD",
                    "posted_date": "2026-06-18",
                },
                {
                    "source_job_id": "2",
                    "title": "Role B",
                    "company": "B",
                    "location": "Toronto",
                    "source": "linkedin",
                    "job_url": "https://example.com/b",
                    "salary_text": "$80,000",
                    "salary_min": 80000,
                    "salary_max": 80000,
                    "salary_interval": "yearly",
                    "salary_currency": "CAD",
                    "posted_date": "2026-06-19",
                },
            ],
            "errors": {},
            "source_counts": {"linkedin": 2},
        }
        with patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            result = self.client.post(
                "/api/discovery/jobspy",
                json={"keywords": "software developer", "location": "Toronto", "enabled_sources": ["linkedin"]},
            ).get_json()
        self.assertEqual(result["preferences"]["sort"]["field"], "salary")
        self.assertEqual(result["jobs"][0]["company"], "B")

    def test_discovery_search_results_are_stored_and_filtering_is_local(self):
        fake = {
            "jobs": [
                {
                    "source_job_id": "1",
                    "title": "Software Developer",
                    "company": "Scotiabank",
                    "location": "Toronto",
                    "source": "linkedin",
                    "job_url": "https://example.com/1",
                    "salary_text": "$90,000",
                    "salary_min": 90000,
                    "salary_max": 90000,
                    "salary_interval": "yearly",
                    "salary_currency": "CAD",
                    "posted_date": "2026-06-19",
                },
                {
                    "source_job_id": "2",
                    "title": "Support Analyst",
                    "company": "Scotiabank",
                    "location": "Toronto",
                    "source": "linkedin",
                    "job_url": "https://example.com/2",
                    "salary_text": "$80,000",
                    "salary_min": 80000,
                    "salary_max": 80000,
                    "salary_interval": "yearly",
                    "salary_currency": "CAD",
                    "posted_date": "2026-06-18",
                },
            ],
            "errors": {},
            "source_counts": {"linkedin": 2},
        }
        with patch.object(app_module, "fetch_jobspy_jobs", return_value=fake):
            result = self.client.post(
                "/api/discovery/jobspy",
                json={
                    "keywords": "software developer",
                    "location": "Toronto, Ontario",
                    "enabled_sources": ["linkedin"],
                    "search_depth": "thorough",
                    "date_window": "30d",
                    "per_source_cap": 250,
                    "total_cap": 500,
                    "salary_min": 70000,
                    "overall_score_min": 70,
                    "company": "Scotiabank",
                },
            ).get_json()

        self.assertEqual(result["total"], 2)
        self.assertEqual(result["visible_total"], 2)
        self.assertIn("search", result)
        stored = self.client.get(f"/api/discovery/results?search_key={result['search_key']}").get_json()
        self.assertEqual(len(stored["jobs"]), 2)
        self.assertEqual(stored["search"]["search_depth"], "thorough")
        self.assertEqual(stored["metadata"]["raw_jobs_returned"], 2)
        self.assertTrue(stored["cache_hit"])
        self.assertEqual(stored["cached_raw_result_count"], 2)

    def test_salary_extraction_from_structured_fields(self):
        salary = app_module.determine_salary(
            {
                "salary_min": 65000,
                "salary_max": 80000,
                "salary_interval": "yearly",
                "salary_currency": "CAD",
                "job_description": "",
            }
        )
        self.assertEqual(salary["salary_source"], "structured")
        self.assertEqual(salary["salary_min"], 65000)
        self.assertEqual(salary["salary_max"], 80000)

    def test_salary_extraction_from_description_and_hourly_annualization(self):
        yearly = app_module.parse_salary_text("Compensation range: $65,000-$80,000 per year plus benefits.")
        hourly = app_module.parse_salary_text("From $30 an hour for full-time support work.")
        self.assertEqual(yearly["salary_source"], "description")
        self.assertEqual(yearly["salary_max"], 80000)
        self.assertEqual(hourly["salary_interval"], "hourly")
        self.assertEqual(hourly["salary_min"], 62400)

    def test_salary_parser_avoids_false_bonus_matches_and_estimates_missing(self):
        false_match = app_module.parse_salary_text("Includes a $5,000 signing bonus and wellness benefits.")
        estimate = app_module.determine_salary(
            {
                "title": "Software Developer",
                "location": "Toronto",
                "role_type": "software developer",
                "seniority_level": "Junior",
                "job_description": "Build APIs and dashboards.",
            }
        )
        self.assertEqual(false_match["salary_source"], "missing")
        self.assertEqual(estimate["salary_source"], "estimated")
        self.assertIn("Estimated market range", estimate["salary_text"])

    def test_revised_scoring_salary_location_dominate_and_missing_not_zero(self):
        profile = {
            "preferred_roles": "software developer",
            "preferred_locations": "Toronto, remote",
            "key_skills": "Python, SQL",
            "resume_text": "Target salary $75k. Built Python SQL tools.",
        }
        high = app_module.score_job(
            {
                "title": "Software Developer",
                "location": "Toronto hybrid",
                "remote_type": "Hybrid",
                "salary": "CAD $85,000 per year",
                "salary_min": 85000,
                "salary_max": 85000,
                "salary_source": "structured",
                "salary_confidence": 95,
                "job_description": "Python SQL APIs",
            },
            profile,
        )
        low = app_module.score_job(
            {
                "title": "Software Developer",
                "location": "Far Away",
                "remote_type": "On-site",
                "salary": "CAD $45,000 per year",
                "salary_min": 45000,
                "salary_max": 45000,
                "salary_source": "structured",
                "salary_confidence": 95,
                "job_description": "Python SQL APIs",
            },
            profile,
        )
        missing = app_module.score_job(
            {
                "title": "Software Developer",
                "location": "Toronto",
                "salary_source": "estimated",
                "salary_confidence": 38,
                "salary_min": 65000,
                "salary_max": 90000,
                "job_description": "Python SQL APIs",
            },
            profile,
        )
        self.assertGreater(high["fit_score"], low["fit_score"])
        self.assertGreater(missing["salary_score"], 0)
        self.assertEqual(high["breakdown"]["previous_formula"].startswith("salary 30%"), True)
        self.assertEqual(high["breakdown"]["weights"]["skills"], 30)

    def test_junior_profile_matched_with_zero_to_two_year_job(self):
        profile = {
            "target_seniority_levels": "junior",
            "maximum_required_years": "2",
            "willing_to_apply_to_stretch_roles": "true",
            "management_roles_allowed": "false",
        }
        job = {"title": "Junior Developer", "job_description": "Required: 0-2 years of experience building Python tools."}
        analysis = seniority.analyze_job_seniority(job)
        fit = seniority.score_seniority_fit(profile, analysis)
        self.assertEqual(analysis["detected_level"], "junior")
        self.assertGreaterEqual(fit["score"], 85)

    def test_junior_profile_penalized_for_five_plus_year_role(self):
        profile = {
            "target_seniority_levels": "junior",
            "maximum_required_years": "2",
            "willing_to_apply_to_stretch_roles": "false",
            "management_roles_allowed": "false",
        }
        job = {"title": "Software Developer", "job_description": "Minimum required 5+ years of experience with APIs."}
        fit = seniority.score_seniority_fit(profile, seniority.analyze_job_seniority(job))
        self.assertLess(fit["score"], 55)

    def test_junior_profile_strongly_penalized_for_ten_year_manager_role(self):
        profile = {
            "target_seniority_levels": "junior",
            "maximum_required_years": "2",
            "willing_to_apply_to_stretch_roles": "true",
            "management_roles_allowed": "false",
        }
        job = {
            "title": "Engineering Manager",
            "job_description": "Requires 10+ years of experience, people management, direct reports, hiring, and budget responsibility.",
        }
        fit = seniority.score_seniority_fit(profile, seniority.analyze_job_seniority(job))
        self.assertLessEqual(fit["score"], 15)
        self.assertTrue(fit["management_role"])

    def test_stretch_role_behaviour_and_management_disabled(self):
        stretch_profile = {
            "target_seniority_levels": "junior",
            "maximum_required_years": "2",
            "willing_to_apply_to_stretch_roles": "true",
            "management_roles_allowed": "false",
        }
        stretch_job = {"title": "Intermediate Developer", "job_description": "Required: 3 years of experience with Flask."}
        stretch_fit = seniority.score_seniority_fit(stretch_profile, seniority.analyze_job_seniority(stretch_job))
        self.assertTrue(stretch_fit["stretch_role"])
        self.assertGreater(stretch_fit["score"], 45)

        manager_job = {"title": "Developer Manager", "job_description": "Manage staff and run performance reviews."}
        manager_fit = seniority.score_seniority_fit(stretch_profile, seniority.analyze_job_seniority(manager_job))
        self.assertLess(manager_fit["score"], stretch_fit["score"])

    def test_title_description_disagreement_and_experience_false_positives(self):
        job = {
            "title": "Senior Software Developer",
            "job_description": "No experience required. Salary is $80,000 per year. Python 3.12. Team of 6. Founded in 2024.",
        }
        analysis = seniority.analyze_job_seniority(job)
        self.assertTrue(analysis["title_description_disagreement"])
        self.assertEqual(analysis["required_experience"]["minimum"], 0)
        self.assertFalse(any(item["minimum"] == 6 for item in analysis["ambiguous_experience"]))
        self.assertFalse(any(item["minimum"] == 2024 for item in analysis["ambiguous_experience"]))

    def test_direct_skill_match_semantic_transferable_and_unrelated_low(self):
        profile = {
            "key_skills": "Python, Flask, SQL",
            "work_experience": "Built Flask APIs and SQL reports for operations teams.",
            "projects": "Created scripts that automated weekly reporting workflows.",
        }
        job = {
            "title": "Backend Developer",
            "required_skills": "Flask APIs",
            "responsibilities": "Automate reporting workflows; maintain biology lab equipment",
        }
        result = semantic_matching.analyze_resume_match(
            job,
            profile,
            cache_dir=None,
            embedding_backend=FakeEmbeddingBackend(),
        )
        self.assertTrue(any(match["match_type"] == "direct" for match in result["matches"]))
        self.assertTrue(any("Automate reporting" in match["job_requirement"] and match["confidence"] >= 45 for match in result["matches"]))
        unrelated = [match for match in result["matches"] if "biology" in match["job_requirement"].lower()]
        self.assertTrue(unrelated)
        self.assertLess(unrelated[0]["confidence"], 45)

    def test_professional_evidence_weighted_above_project_evidence(self):
        profile = {
            "work_experience": "Built Python automation for production reporting.",
            "projects": "Built Python automation for a hobby dashboard.",
        }
        result = semantic_matching.analyze_resume_match(
            {"required_skills": "Python automation"},
            profile,
            cache_dir=None,
            embedding_backend=FakeEmbeddingBackend(),
        )
        best = result["matches"][0]
        self.assertEqual(best["evidence_source"], "professional experience")

    def test_unsupported_skill_not_added_and_original_bullet_preserved(self):
        profile = {
            "key_skills": "Python, SQL",
            "work_experience": "Built Python reports with SQL for operations teams.",
        }
        job = {"required_skills": "Kubernetes, Python", "job_description": "Use Kubernetes and Python."}
        match = semantic_matching.analyze_resume_match(job, profile, cache_dir=None, embedding_backend=FakeEmbeddingBackend())
        draft = tailoring.generate_tailored_resume_draft(
            job,
            profile,
            match,
            {"explanation": "Junior fit.", "score": 80},
        )
        self.assertIn("Built Python reports", draft["text"])
        self.assertNotIn("Built Kubernetes", draft["text"])
        self.assertTrue(all("original_bullet" in item for item in draft["rewrite_suggestions"]))

    def test_rule_based_rewrite_does_not_invent_content(self):
        original = "made Flask APIs for internal reporting"
        proposed = tailoring.controlled_rewrite(original, "Flask APIs")
        self.assertIn("Flask APIs", proposed)
        self.assertNotIn("Kubernetes", proposed)
        self.assertNotIn("%", proposed)

    def rece_profile(self):
        return {
            "name": "Avery Chen",
            "preferred_roles": "Registered Early Childhood Educator, Early Childhood Educator",
            "key_skills": "child-centred programming, play-based learning, developmental observation, family communication, health and safety compliance, inclusive practice, teamwork",
            "work_experience": "Registered Early Childhood Educator at Little Pines Child Care. Planned child-centred and play-based programming, completed developmental observation and documentation, maintained health and safety compliance, communicated with families, and supported inclusive classroom routines.",
            "education": "Early Childhood Education diploma.",
            "projects": "",
            "resume_text": "",
        }

    def software_profile(self, extra_skills=""):
        skills = "Python, SQL, Flask, cloud, automation"
        if extra_skills:
            skills = f"{skills}, {extra_skills}"
        return {
            "name": "Terry Mak",
            "preferred_roles": "software developer",
            "key_skills": skills,
            "work_experience": "Software developer with Python automation, Flask APIs, SQL reports, and cloud deployment experience.",
            "projects": "Built Flask tools for reporting and automation.",
            "education": "Computer programming diploma.",
            "resume_text": "Built Python, SQL, Flask, cloud, and automation projects.",
        }

    def ece_job(self):
        return {
            "company": "Little Sprouts",
            "title": "Early Childhood Educator",
            "required_skills": "child-centred programming, developmental observation, family communication, health and safety compliance",
            "job_description": "Plan play-based learning, document development, communicate with families, and maintain safe inclusive classrooms.",
        }

    def software_job(self, required="Python, SQL, Flask, cloud, automation"):
        return {
            "company": "BuildCo",
            "title": "Software Developer",
            "required_skills": required,
            "job_description": f"Build software using {required}.",
        }

    def scotiabank_job(self):
        return {
            "company": "Scotiabank",
            "title": "Software Developer",
            "location": "Toronto, ON",
            "required_skills": "Develop and maintain software applications, write well-designed testable REST APIs, software documentation, collaboration with developers, product managers, and designers, version control and Git branching, production troubleshooting, software reliability, sprint objectives and project milestones, backend integrations, database design, cloud platforms",
            "preferred_skills": "Node.js, Express, JavaScript, Agile, Scrum, microservice architecture, unit and integration testing",
            "responsibilities": "Build backend integrations, support production incident troubleshooting, document release steps, and deliver project milestones with cross-functional teams.",
            "job_description": "Develop and maintain software applications and REST APIs, collaborate with developers, product managers, and designers, support production troubleshooting, strengthen software reliability, and contribute to sprint objectives and project milestones. The role also values cloud platform work, database design, technical documentation, and backend integrations.",
        }

    def test_rece_profile_does_not_generate_software_terminology(self):
        output = app_module.generate_application_materials(self.ece_job(), self.rece_profile())
        combined = "\n".join([output["resume_text"], output["cover_letter"], app_module.answers_to_text(output["short_answers"])])
        self.assertIn("Registered Early Childhood Educator", combined)
        self.assertIn("child-centred", combined)
        self.assertNotRegex(combined.lower(), r"\b(aws|python|flask|react|sql|dashboard|api|cloud|software projects)\b")

    def test_software_profile_can_generate_supported_software_terms(self):
        output = app_module.generate_application_materials(self.software_job(), self.software_profile())
        combined = "\n".join([output["resume_text"], output["cover_letter"]])
        self.assertIn("Python", combined)
        self.assertIn("Flask", combined)
        self.assertIn("software developer", combined.lower())

    def test_aws_suggestion_requires_candidate_evidence(self):
        without_aws = app_module.generate_application_materials(self.software_job("AWS, Python"), self.software_profile())
        self.assertIn("aws", [item.lower() for item in without_aws["unsupported_keywords"]])
        self.assertNotIn("AWS", without_aws["skill_guidance"]["supported_relevant"])

        with_aws = app_module.generate_application_materials(self.software_job("AWS, Python"), self.software_profile("AWS"))
        self.assertIn("AWS", with_aws["skill_guidance"]["supported_relevant"])

    def test_corrupted_pdf_text_detected_and_rejected(self):
        noisy_pdf = b"%PDF-1.4\n1 0 obj\n<< /Type /Page /FontName /CIDFont /XObject >>\nstream\n\\x00\\xff\\x10FFFFAAAAFFFFAAAAFFFFAAAAFFFF\nendstream\nendobj\nxref\ntrailer\nstartxref"
        with self.assertRaisesRegex(ValueError, "This PDF could not be read reliably"):
            app_module.extract_pdf_text(noisy_pdf)

    def test_structured_profile_used_when_pdf_extraction_fails(self):
        profile = self.rece_profile()
        profile["resume_text"] = "endobj xref trailer stream glyph CIDFont FFFFEEEEFFFFEEEE"
        output = app_module.generate_application_materials(self.ece_job(), profile)
        self.assertIn("developmental observation", output["resume_text"])
        self.assertNotIn("CIDFont", output["resume_text"])

    def test_cross_domain_generation_blocks_unsupported_terms(self):
        with self.assertRaisesRegex(ValueError, "Unsupported or conflicting phrases"):
            app_module.generate_application_materials(self.software_job(), self.rece_profile())

    def test_cache_invalidation_after_profile_change(self):
        job = self.software_job()
        first = semantic_matching.analyze_resume_match(job, self.software_profile(), cache_dir=None, use_semantic=False)
        second = semantic_matching.analyze_resume_match(job, self.software_profile("AWS"), cache_dir=None, use_semantic=False)
        self.assertNotEqual(first["profile_resume_hash"], second["profile_resume_hash"])

    def test_generated_versions_mark_outdated_after_profile_change(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/jobs",
            json={"company": "OpsTools", "title": "Software Developer", "job_description": "Python SQL Flask"},
        ).get_json()
        generated = self.client.post("/api/generate", json={**created, "id": created["id"]})
        self.assertEqual(generated.status_code, 200)
        self.client.post("/api/profile", json={**self.rece_profile(), "name": "New Profile"})
        versions = self.client.get(f"/api/jobs/{created['id']}/versions").get_json()
        self.assertTrue(any(version["status"] == "Outdated — regenerate using current profile" for version in versions))

    def test_tailored_resume_preserves_master_docx_formatting(self):
        self.seed_profile_for_documents(with_template=True)
        profile = self.client.get("/api/profile").get_json()
        master_path = Path(profile["resume_docx_template_path"])
        master_before = master_path.read_bytes()
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/job/format",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        self.assertEqual(created["resume_generation_status"], "generated")
        resume_path = Path(created["resume_docx_path"])
        self.assertEqual(master_before, master_path.read_bytes())

        master_doc = Document(io.BytesIO(master_path.read_bytes()))
        tailored_doc = Document(io.BytesIO(resume_path.read_bytes()))

        self.assertEqual(len(master_doc.sections), len(tailored_doc.sections))
        self.assertEqual(len(master_doc.paragraphs), len(tailored_doc.paragraphs))
        self.assertEqual(len(master_doc.tables), len(tailored_doc.tables))

        self.assertEqual(master_doc.paragraphs[0].style.name, tailored_doc.paragraphs[0].style.name)
        self.assertEqual(master_doc.paragraphs[0].text, tailored_doc.paragraphs[0].text)

        self.assertEqual(master_doc.paragraphs[2].text, tailored_doc.paragraphs[2].text)
        self.assertEqual(master_doc.paragraphs[2].runs[0].font.name, tailored_doc.paragraphs[2].runs[0].font.name)
        self.assertEqual(master_doc.paragraphs[2].runs[0].font.size.pt, tailored_doc.paragraphs[2].runs[0].font.size.pt)

        self.assertEqual(master_doc.paragraphs[5].style.name, "List Bullet")
        self.assertEqual(tailored_doc.paragraphs[5].style.name, "List Bullet")
        self.assertEqual(master_doc.paragraphs[5].text, tailored_doc.paragraphs[5].text)
        self.assertEqual(master_doc.paragraphs[6].text, tailored_doc.paragraphs[6].text)
        self.assertEqual(master_doc.paragraphs[8].style.name, tailored_doc.paragraphs[8].style.name)
        # Skill reordering is optional and must be rejected when it cannot be
        # proven layout-safe; preserving the source is an approved outcome.
        self.assertEqual(master_doc.paragraphs[8].text, tailored_doc.paragraphs[8].text)

        self.assertEqual(tailored_doc.tables[0].cell(1, 1).text, "Flask, React")
        self.assertEqual(tailored_doc.tables[1].cell(1, 1).text, "Computer programming diploma")
        with zipfile.ZipFile(resume_path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
        self.assertNotIn("Suggested summary", xml)
        self.assertIn("w:hyperlink", xml)

    def test_resume_without_summary_keeps_protected_paragraphs_and_changes_only_bullet_text(self):
        self.seed_profile_for_documents(with_template=True, resume_template="no_summary")
        profile = self.client.get("/api/profile").get_json()
        master_path = Path(profile["resume_docx_template_path"])
        master_before = master_path.read_bytes()
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/job/no-summary",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        self.assertEqual(created["resume_generation_status"], "generated")
        resume_path = Path(created["resume_docx_path"])
        change_log = json.loads((resume_path.parent / "resume_change_log.json").read_text(encoding="utf-8"))
        self.assertTrue(any(item.get("reason") == "Skipped summary tailoring: no existing summary section" for item in change_log["skipped"]))
        self.assertEqual(master_before, master_path.read_bytes())

        master_doc = Document(io.BytesIO(master_path.read_bytes()))
        tailored_doc = Document(io.BytesIO(resume_path.read_bytes()))
        self.assertEqual(len(master_doc.paragraphs), len(tailored_doc.paragraphs))
        self.assertEqual(len(master_doc.tables), len(tailored_doc.tables))
        self.assertEqual(master_doc.paragraphs[0].style.name, tailored_doc.paragraphs[0].style.name)
        self.assertEqual(master_doc.paragraphs[0].text, tailored_doc.paragraphs[0].text)

        protected_texts = [
            "Terry Mak",
            "terry@example.com | LinkedIn | GitHub",
            "Software Developer | Northstar Labs | 2021-2025",
            "George Brown College",
            "Computer programming diploma | 2023",
            "Work Experience",
            "Education",
            "Skills",
            "Portfolio: example.com/portfolio",
        ]
        for text_value in protected_texts:
            self.assertIn(text_value, [paragraph.text for paragraph in master_doc.paragraphs])
            self.assertIn(text_value, [paragraph.text for paragraph in tailored_doc.paragraphs])

        master_bullet = next(paragraph for paragraph in master_doc.paragraphs if paragraph.text.startswith("made Python automation"))
        tailored_bullet = next(paragraph for paragraph in tailored_doc.paragraphs if paragraph.text.startswith("made Python automation"))
        self.assertEqual(master_bullet.style.name, tailored_bullet.style.name)
        self.assertEqual(master_bullet.text, tailored_bullet.text)

        xml = zipfile.ZipFile(resume_path).read("word/document.xml").decode("utf-8", errors="ignore")
        self.assertNotIn("Suggested summary", xml)
        self.assertNotIn("analysis", xml.lower())

    def test_resume_tailoring_makes_grounded_ats_edits_for_python_developer_job(self):
        self.seed_python_developer_profile()
        self.upload_resume_template(self.build_python_developer_resume_fixture())
        profile = self.client.get("/api/profile").get_json()
        master_path = Path(profile["resume_docx_template_path"])
        master_before = master_path.read_bytes()

        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar",
                "title": "Python Developer",
                "location": "Toronto",
                "url": "https://example.com/job/python-dev",
                "job_description": "Python SQL database design AWS Linux testing debugging documentation collaboration",
            },
        ).get_json()

        self.assertEqual(created["resume_generation_status"], "generated")
        resume_path = Path(created["resume_docx_path"])
        self.assertEqual(master_before, master_path.read_bytes())

        change_log = json.loads((resume_path.parent / "resume_change_log.json").read_text(encoding="utf-8"))
        self.assertGreaterEqual(change_log["summary"]["total_editable_paragraphs"], 6)
        self.assertGreaterEqual(change_log["summary"]["total_proposed_changes"], 1)
        self.assertLessEqual(change_log["summary"]["total_applied_changes"], 1)
        self.assertGreaterEqual(change_log["summary"]["total_skipped_changes"], 1)
        self.assertLessEqual(len(change_log["applied"]), 1)
        self.assertTrue(any(item.get("skip_reason") for item in change_log["skipped"]))

        tailored_doc = Document(io.BytesIO(resume_path.read_bytes()))
        tailored_text = "\n".join(paragraph.text for paragraph in tailored_doc.paragraphs)
        self.assertIn("George Brown College", tailored_text)
        self.assertIn("Software Developer | Northstar Labs | 2021-2025", tailored_text)
        self.assertIn("Computer programming diploma | 2023", tailored_text)
        self.assertIn("validation, logging", tailored_text.lower())
        self.assertIn("reduce manual data entry errors by 90%", tailored_text)
        self.assertNotIn("database structures and data models", tailored_text)
        self.assertNotIn("cross-functional data workflows using Retool and Google Docs", tailored_text)

        skills_line = next(paragraph.text for paragraph in tailored_doc.paragraphs if paragraph.text.startswith("Skills:"))
        self.assertEqual(skills_line, "Skills: React, dashboards, AWS, Linux, Python, SQL, Flask, documentation, collaboration, testing")
        self.assertIn("documentation", skills_line.lower())
        self.assertIn("collaboration", skills_line.lower())

        bullet_texts = [paragraph.text for paragraph in tailored_doc.paragraphs if getattr(getattr(paragraph, "style", None), "name", "").startswith("List")]
        self.assertTrue(any("reduce manual data entry errors by 90%" in text for text in bullet_texts))
        self.assertFalse(any("database structures and data models" in text for text in bullet_texts))
        self.assertFalse(any("cross-functional data workflows using Retool and Google Docs" in text for text in bullet_texts))

        opened = Document(io.BytesIO(resume_path.read_bytes()))
        self.assertGreater(len(opened.paragraphs), 0)

    def test_scotiabank_resume_tailoring_makes_meaningful_supported_edits(self):
        self.seed_python_developer_profile()
        self.upload_resume_template(self.build_python_developer_resume_fixture())
        profile = self.client.get("/api/profile").get_json()
        master_path = Path(profile["resume_docx_template_path"])
        master_doc = Document(io.BytesIO(master_path.read_bytes()))
        master_before = master_path.read_bytes()

        created = self.client.post(
            "/api/shortlist",
            json=self.scotiabank_job(),
        ).get_json()

        self.assertEqual(created["resume_generation_status"], "generated")
        resume_path = Path(created["resume_docx_path"])
        self.assertEqual(master_before, master_path.read_bytes())
        change_log = json.loads((resume_path.parent / "resume_change_log.json").read_text(encoding="utf-8"))
        self.assertLessEqual(change_log["summary"]["total_applied_changes"], 3)
        self.assertLessEqual(change_log["summary"]["total_applied_changes"], 6)
        self.assertGreaterEqual(change_log["summary"]["total_skipped_changes"], 1)
        self.assertTrue(
            all(
                re.sub(r"[^a-z0-9]+", "", item["original"].lower())
                != re.sub(r"[^a-z0-9]+", "", item["replacement"].lower())
                for item in change_log["applied"]
            )
        )
        self.assertTrue(all(len(item["replacement"]) <= len(item["original"]) for item in change_log["applied"]))

        tailored_doc = Document(io.BytesIO(resume_path.read_bytes()))
        changed_paragraphs = [
            {
                "index": index,
                "base": base.text,
                "tailored": tailored.text,
                "style": getattr(getattr(base, "style", None), "name", "") or "",
            }
            for index, (base, tailored) in enumerate(zip(master_doc.paragraphs, tailored_doc.paragraphs))
            if base.text != tailored.text
        ]
        substantive_bullet_changes = [
            item for item in changed_paragraphs if item["style"].startswith("List") and not tailoring.is_cosmetic_only_change(item["base"], item["tailored"])
        ]
        self.assertLessEqual(len(substantive_bullet_changes), 3)
        tailored_text = "\n".join(paragraph.text for paragraph in tailored_doc.paragraphs)
        self.assertNotIn("Strengthened software reliability", tailored_text)
        self.assertNotIn("collaborated with cross-functional teammates", tailored_text.lower())
        self.assertNotIn("cross-functional data workflows", tailored_text.lower())
        self.assertNotIn("saved profile", tailored_text.lower())

    def test_scotiabank_requirements_rank_and_map_to_grounded_evidence(self):
        self.seed_python_developer_profile()
        self.upload_resume_template(self.build_python_developer_resume_fixture())
        profile = self.client.get("/api/profile").get_json()
        job = self.scotiabank_job()
        match = semantic_matching.analyze_resume_match(job, profile, cache_dir=None, use_semantic=False)
        plan = tailoring_pipeline.build_tailoring_plan(job, profile, match)

        ranked = plan["ranked_requirements"]
        self.assertGreaterEqual(len(ranked), 8)
        self.assertEqual(ranked[0]["source_field"], "required_skills")
        self.assertTrue(any("database" in item["phrase"].lower() for item in ranked[:5]))
        self.assertTrue(any("testing" in item["phrase"].lower() or "document" in item["phrase"].lower() for item in ranked[:8]))

        inventory = plan["requirement_evidence_inventory"]
        self.assertGreaterEqual(len(inventory), 5)
        top_match = inventory[0]
        self.assertIn("requirement", top_match)
        self.assertIn("matched_source_bullet", top_match)
        self.assertIn("match_reason", top_match)
        self.assertIn("supported_technologies", top_match)
        self.assertIn("supported_outcome", top_match)
        self.assertGreaterEqual(top_match["confidence"], 40)
        self.assertFalse(any("database structures" in item["matched_source_bullet"].lower() for item in inventory))

    def test_controlled_rewrite_rejects_cosmetic_only_changes(self):
        original = "Built supplier specific items and data entry workflows for operations."
        rewritten = tailoring.controlled_rewrite(
            original,
            "database design",
            evidence_text=original,
            section="projects",
        )
        self.assertEqual(rewritten, original)

    def test_minimal_resume_substitution_is_noop_or_explicit_alias_only(self):
        original = "Built NodeJS services for reporting and reduced errors by 90%."
        evidence = {"full_text": original}
        alias_match = {"job_phrase": "Node.js services", "evidence_text": original}
        proposed, reason = tailoring_pipeline.minimal_resume_substitution(original, alias_match)
        validation = tailoring_pipeline.validate_minimal_substitution(original, proposed, alias_match, evidence)
        self.assertEqual(proposed, "Built Node.js services for reporting and reduced errors by 90%.")
        self.assertTrue(validation["grounded"])
        self.assertIn("NodeJS", original)
        self.assertIn("90%", proposed)

        noop, _ = tailoring_pipeline.minimal_resume_substitution(original, {"job_phrase": "Python services"})
        self.assertEqual(noop, original)
        unsupported, _ = tailoring_pipeline.minimal_resume_substitution(original, {"job_phrase": "Kubernetes services"})
        self.assertEqual(unsupported, original)

    def test_layout_validation_rejects_line_or_page_changes(self):
        with patch(
            "resume_layout.measure_docx_layout",
            side_effect=[
                {"available": True, "pages": 1, "paragraphs": [{"index": 0, "lines": 2}]},
                {"available": True, "pages": 2, "paragraphs": [{"index": 0, "lines": 3}]},
            ],
        ):
            result = resume_layout.compare_docx_layout("original.docx", "candidate.docx")
        self.assertFalse(result["verified"])
        self.assertTrue(result["page_mismatch"])
        self.assertTrue(result["line_mismatches"])

    def test_word_layout_measurement_cache_uses_content_fingerprint(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "resume.docx"
            path.write_bytes(f"stable document {time.time_ns()}".encode())
            measurement = {"available": True, "pages": 1, "paragraphs": [{"index": 0, "lines": 1}]}
            with patch.object(resume_layout, "_run_word", return_value=(measurement, None)) as run_word:
                first = resume_layout.measure_docx_layout(path)
                second = resume_layout.measure_docx_layout(path)
            self.assertTrue(first["available"])
            self.assertTrue(second["cached"])
            self.assertEqual(run_word.call_count, 1)
            path.write_bytes(f"changed document {time.time_ns()}".encode())
            with patch.object(resume_layout, "_run_word", return_value=(measurement, None)) as changed_run:
                resume_layout.measure_docx_layout(path)
            self.assertEqual(changed_run.call_count, 1)

    def test_application_workspace_header_keeps_existing_queue_as_single_workflow(self):
        html = self.client.get("/").data.decode("utf-8")
        self.assertIn('data-tab="shortlist" type="button">Application workspace', html)
        self.assertIn("Discover in Pipeline, prepare and tailor here", html)
        self.assertIn('class="application-workflow-rail"', html)

    def test_keyword_coverage_deduplicates_aliases_and_warns_on_unsupported_groups(self):
        job = {"title": "Software Engineer", "required_skills": "NodeJS, Node.js, Kubernetes, database design", "preferred_skills": "Postgres, PostgreSQL"}
        profile = {"work_experience": "Built NodeJS services with SQL reporting.", "key_skills": "NodeJS, SQL"}
        semantic = {
            "matches": [
                {"job_requirement": "NodeJS", "category": "required_skill", "matched_resume_evidence": "Built NodeJS services.", "confidence": 80, "match_type": "direct"},
                {"job_requirement": "Node.js", "category": "required_skill", "matched_resume_evidence": "Built NodeJS services.", "confidence": 80, "match_type": "direct"},
                {"job_requirement": "Kubernetes", "category": "required_skill", "matched_resume_evidence": "", "confidence": 0, "match_type": "missing"},
                {"job_requirement": "database design", "category": "required_skill", "matched_resume_evidence": "", "confidence": 0, "match_type": "missing"},
                {"job_requirement": "Postgres", "category": "preferred_skill", "matched_resume_evidence": "SQL reporting", "confidence": 60, "match_type": "transferable"},
                {"job_requirement": "PostgreSQL", "category": "preferred_skill", "matched_resume_evidence": "SQL reporting", "confidence": 60, "match_type": "transferable"},
            ]
        }
        report = app_module.build_keyword_coverage_report(job, profile, semantic)
        self.assertEqual(report["required"]["total"], 4)
        self.assertEqual(report["preferred"]["total"], 1)
        self.assertTrue(report["review_job_fit"])
        self.assertTrue(any("below 60%" in warning for warning in report["warnings"]))
        self.assertTrue(any(item["requirement"] == "Kubernetes" for item in report["required"]["unsupported"]))

    def test_keywords_score_uses_boundaries_aliases_and_separate_requirement_types(self):
        job = {
            "title": "Software Engineer",
            "job_description": "Requirements: Node.js, Java, C++, C#, backend and REST APIs. Preferred: PostgreSQL and Figma. AI is mentioned in company marketing.",
            "required_skills": "Node.js, Java, C++, C#, back-end, RESTful APIs",
            "preferred_skills": "Postgres, PostgreSQL, Figma",
        }
        profile = {
            "work_experience": "Built NodeJS services with JavaScript and PostgreSQL. Used C++ and C# on Linux.",
            "key_skills": "NodeJS, Java, C++, C#, backend, REST APIs",
        }
        report = grounding.build_keywords_score_report(job, profile)
        normalized = {item["normalized"] for item in report["matched"]}
        self.assertIn("node.js", normalized)
        self.assertIn("rest apis", normalized)
        self.assertNotIn("javascript", {item["normalized"] for item in report["missing"]})
        self.assertEqual(report["preferred"]["status"], "assessed")
        self.assertLess(report["preferred"]["score_percent"], 100)
        self.assertTrue(any(item["normalized"] == "figma" for item in report["missing"]))
        self.assertFalse(grounding.keyword_present("paid training", "ai"))
        self.assertTrue(grounding.keyword_present("AI systems", "ai"))
        self.assertFalse(grounding.keyword_present("JavaScript", "java"))

    def test_keywords_score_keeps_related_evidence_out_of_match_numerator(self):
        job = {
            "title": "Platform Engineer",
            "job_description": "Required: DevOps and PaaS experience. Preferred: Kafka.",
            "required_skills": "DevOps, PaaS",
            "preferred_skills": "Kafka",
        }
        profile = {"work_experience": "Deployed cloud applications to production and built a multi-tenant platform."}
        report = grounding.build_keywords_score_report(job, profile)
        self.assertTrue(report["related"])
        self.assertEqual(report["matched_count"], 0)
        self.assertEqual(report["required"]["matched"], 0)
        self.assertTrue(report["review_before_applying"])

    def test_keywords_score_uses_only_selected_resume_text(self):
        job = {
            "title": "Software Engineer",
            "required_skills": "Python, JavaScript",
            "job_description": "Requirements: Python and JavaScript.",
        }
        profile = {
            "work_experience": "Built Python services for the saved profile.",
            "key_skills": "Python",
            "resume_text": "Selected resume evidence: Built JavaScript interfaces.",
        }
        report = grounding.build_keywords_score_report(job, profile, final_resume_text=profile["resume_text"])
        matched = {item["normalized"] for item in report["matched"]}
        missing = {item["normalized"] for item in report["missing"]}
        self.assertIn("javascript", matched)
        self.assertIn("python", missing)
        self.assertTrue(all("saved profile" not in item["evidence"] for item in report["matched"]))

    def test_keywords_score_returns_bounded_source_evidence_and_actual_alias_spelling(self):
        job = {
            "title": "Software Engineer",
            "job_description": "Requirements: Node.js and DevOps. Preferred: Kafka.",
            "required_skills": "Node.js, DevOps",
            "preferred_skills": "Kafka",
        }
        long_resume = "NodeJS, Express, REST APIs, PostgreSQL. Deployed cloud applications to production. " + ("Additional unrelated résumé detail. " * 30)
        profile = {
            "resume_text": long_resume,
            "_resume_sections": [{"label": "Skills, Technologies", "text": long_resume}],
        }
        report = grounding.build_keywords_score_report(job, profile)
        node = next(item for item in report["matched"] if item["normalized"] == "node.js")
        self.assertEqual(node["match_type"], "equivalent")
        self.assertEqual(node["found_as"], "nodejs")
        self.assertEqual(node["source_section"], "Skills, Technologies")
        self.assertLessEqual(len(node["evidence"]), 160)
        self.assertNotIn("Additional unrelated", node["evidence"])
        related = next(item for item in report["related"] if item["normalized"] == "devops")
        self.assertLessEqual(len(related["evidence"]), 160)
        kafka = next(item for item in report["missing"] if item["normalized"] == "kafka")
        self.assertEqual(kafka["evidence"], "")

    def test_keywords_score_deduplicates_overlapping_aliases(self):
        job = {
            "title": "Web Engineer",
            "job_description": "NodeJS and Node.js, Java Script, back-end and backend, front-end and frontend, RESTful APIs, C++, C#, .NET and SQL.",
        }
        inventory = grounding.extract_job_keywords(job)
        normalized = [item["normalized"] for item in inventory]
        self.assertEqual(len(normalized), len(set(normalized)))
        for expected in {"node.js", "javascript", "backend", "frontend", "rest apis", "c++", "c#", ".net", "sql"}:
            self.assertIn(expected, normalized)
        self.assertNotIn("java", normalized)
        self.assertNotIn("api", normalized)

    def test_keywords_score_requires_explicit_requirement_sections(self):
        job = {
            "title": "Platform Engineer",
            "job_description": "Responsibilities:\nBuild Kafka services.\nPreferred:\nFigma experience.\nRequirements:\nPython experience.",
        }
        inventory = {item["normalized"]: item for item in grounding.extract_job_keywords(job)}
        self.assertEqual(inventory["kafka"]["requirement_type"], "unclassified")
        self.assertEqual(inventory["figma"]["requirement_type"], "preferred")
        self.assertEqual(inventory["python"]["requirement_type"], "required")

    def test_tailoring_studio_builds_full_lakeview_requirement_inventory(self):
        job = {
            "title": "Backend Engineer",
            "required_skills": "3 to 5+ years of production backend experience, advanced distributed systems, Kafka, Kubernetes",
            "preferred_skills": "Python, PostgreSQL, Docker",
            "job_description": (
                "Requirements:\n"
                "3 to 5+ years of production backend experience.\n"
                "Advanced distributed systems expertise.\n"
                "Build Python services and APIs with PostgreSQL.\n"
                "Preferred:\nKafka, Kubernetes, Docker.\n"
                "Collaborate with product teams and participate in code review."
            ),
        }
        resume = "Built Python Flask APIs with PostgreSQL and automated testing. Collaborated with product teams and used Git."
        report = grounding.build_tailoring_coverage_report(job, resume)
        inventory = {item["normalized"]: item for group in ("matched", "related", "missing") for item in report[group]}
        self.assertGreater(report["total"], 1)
        self.assertIn("backend programming languages", inventory)
        self.assertIn("distributed and event-driven systems", inventory)
        self.assertIn("messaging and queue platforms", inventory)
        self.assertIn("production backend experience", inventory)
        self.assertEqual(inventory["backend programming languages"]["match_type"], "exact")
        self.assertEqual(inventory["messaging and queue platforms"]["match_type"], "missing")
        self.assertEqual(inventory["distributed and event-driven systems"]["match_type"], "missing")
        self.assertEqual(inventory["distributed and event-driven systems"]["requirement_type"], "required")
        self.assertTrue(report["review_before_applying"])
        self.assertLess(report["score_percent"], 100)

    def test_complete_lakeview_fixture_groups_alternatives_and_preserves_sections(self):
        description = (Path(__file__).parent / "fixtures" / "lakeview_backend_engineer_job.txt").read_text(encoding="utf-8")
        job = {"title": "Backend Engineer", "job_description": description}
        report = grounding.build_tailoring_coverage_report(
            job,
            "Built Python Flask REST APIs with PostgreSQL, automated testing, AWS and data models.",
        )
        inventory = {item["normalized"]: item for group in ("matched", "related", "missing") for item in report[group]}
        self.assertGreaterEqual(report["total"], 15)
        self.assertIn("backend programming languages", inventory)
        self.assertIn("messaging and queue platforms", inventory)
        self.assertIn("transactional correctness", inventory)
        self.assertIn("llm and ai orchestration", inventory)
        self.assertEqual(inventory["messaging and queue platforms"]["requirement_type"], "preferred")
        self.assertTrue(inventory["api and service design"]["source_evidence"])
        self.assertNotIn("and the ability to build dependable", {item["keyword"] for item in inventory.values()})
        self.assertEqual(inventory["backend programming languages"]["match_type"], "exact")

    def test_gatekeeper_testing_suggestions_match_morphology_and_prefer_quantified_experience(self):
        description = (Path(__file__).parent / "fixtures" / "lakeview_backend_engineer_job.txt").read_text(encoding="utf-8")
        self.assertIn("testing", description.lower())
        bullets = [
            {
                "paragraph_index": 12,
                "text": "Used Codex to accelerate debugging and automated testing while preparing changes for review.",
                "section": "projects",
                "editable": True,
            },
            {
                "paragraph_index": 27,
                "text": "Added conflict protection, validation, migrations, and 25+ automated backend tests across purchasing, inventory, and reporting workflows.",
                "section": "work_experience",
                "editable": True,
            },
            {
                "paragraph_index": 4,
                "text": "Technologies: Python, Flask, PostgreSQL, testing",
                "section": "skills",
                "editable": True,
            },
        ]
        suggestions = grounding.rank_resume_bullets_for_requirement("testing", bullets)
        self.assertEqual(suggestions[0]["paragraph_index"], 27)
        self.assertIn("tests", suggestions[0]["overlap"])
        self.assertEqual(suggestions[0]["evidence_strength"], "quantified substantive evidence")
        self.assertIn("test", grounding.resume_requirement_token_forms("testing") & grounding.resume_requirement_token_forms("tests"))
        self.assertIn("testing", description.lower())
    def test_nodejs_alias_substitution_preserves_word_layout_and_reports_final_docx_coverage(self):
        master = Path(self.tmpdir.name) / "nodejs_alias_master.docx"
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)
        title = document.add_paragraph()
        title.add_run("Alex Morgan").bold = True
        summary = document.add_paragraph("Software developer focused on reliable internal tools.")
        document.add_paragraph("Work Experience", style="Heading 1")
        experience = document.add_paragraph(style="List Bullet")
        experience.add_run("Built NodeJS services and reduced review time by 35%.")
        document.add_paragraph("Skills", style="Heading 1")
        skills = document.add_paragraph("Skills: NodeJS, SQL")
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(10)
        document.save(master)
        master_before = master.read_bytes()

        self.client.post(
            "/api/profile",
            json={
                "name": "Alex Morgan",
                "key_skills": "NodeJS, SQL",
                "work_experience": "Built NodeJS services and reduced review time by 35%.",
                "resume_text": "Built NodeJS services and reduced review time by 35%.",
            },
        )
        self.upload_resume_template(master)
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Example Systems",
                "title": "Software Developer",
                "job_description": "Requirements: Node.js experience. Preferred: SQL.",
                "required_skills": "Node.js",
                "preferred_skills": "SQL",
            },
        ).get_json()
        self.assertEqual(created["resume_generation_status"], "generated", created)
        resume_path = Path(created["resume_docx_path"])
        self.assertEqual(master_before, master.read_bytes())

        comparison = resume_layout.compare_docx_layout(master, resume_path)
        if not comparison["available"]:
            self.skipTest(f"Word layout measurement unavailable: {comparison.get('reason', 'unknown reason')}")
        self.assertTrue(comparison["available"], comparison.get("reason"))
        self.assertTrue(comparison["verified"], comparison)
        self.assertEqual(comparison["original_pages"], 1)
        self.assertEqual(comparison["candidate_pages"], 1)

        tailored = Document(resume_path)
        tailored_text = "\n".join(paragraph.text for paragraph in tailored.paragraphs)
        change_log = json.loads((resume_path.parent / "resume_change_log.json").read_text(encoding="utf-8"))
        self.assertIn("Built Node.js services and reduced review time by 35%.", tailored_text, change_log)
        self.assertNotIn("Built NodeJS services", tailored_text)
        self.assertIn("35%", tailored_text)
        self.assertEqual(document.paragraphs[3].style.name, tailored.paragraphs[3].style.name)
        self.assertEqual(document.paragraphs[3].runs[0].font.size, tailored.paragraphs[3].runs[0].font.size)
        del tailored

        self.assertTrue(any(
            "NodeJS" in item.get("original", "") and "Node.js" in item.get("replacement", "")
            for item in change_log["applied"]
        ))
        preview = self.client.get(f"/api/shortlist/{created['id']}/preview").get_json()
        downloaded = self.client.get(f"/api/shortlist/{created['id']}/resume")
        self.assertEqual(downloaded.status_code, 200)
        downloaded_document = Document(io.BytesIO(downloaded.data))
        downloaded_text = "\n".join(
            paragraph.text for paragraph in app_module.iter_document_paragraphs(downloaded_document) if paragraph.text.strip()
        )
        del downloaded_document
        downloaded.close()
        self.assertEqual(preview["resume_source"], "successful_docx")
        self.assertEqual(preview["resume"], downloaded_text)
        self.assertIn("Node.js", preview["resume"])

        metadata = json.loads((resume_path.parent / "metadata.json").read_text(encoding="utf-8"))
        coverage = metadata["keyword_coverage"]
        self.assertIn("before", coverage)
        self.assertIn("after", coverage)
        self.assertEqual(coverage["after"]["required"]["coverage_percent"], coverage["required"]["coverage_percent"])
        self.assertEqual(coverage["before"]["required"]["coverage_percent"], coverage["after"]["required"]["coverage_percent"])
        self.assertTrue(any("node.js" in item["requirement"].lower() for item in coverage["after"]["required"]["covered"]))
        self.assertTrue(any("sql" in item["requirement"].lower() for item in coverage["after"]["preferred"]["covered"]))
        gc.collect()

    def test_keyword_coverage_ignores_incidental_description_mentions_without_requirement_labels(self):
        job = {
            "title": "Software Developer",
            "job_description": "Our team uses Kubernetes for one internal service. Requirements: Node.js experience. Preferred: SQL.",
        }
        profile = {"work_experience": "Built NodeJS services for reporting.", "key_skills": "NodeJS"}
        report = app_module.build_keyword_coverage_report(job, profile, {"matches": []})
        requirements = [item["requirement"] for item in report["required"]["covered"] + report["required"]["unsupported"]]
        self.assertTrue(any("node.js" in requirement.lower() for requirement in requirements))
        self.assertNotIn("Kubernetes", requirements)

    def test_cover_letter_template_upload_and_placeholder_replacement(self):
        self.seed_profile_for_documents(with_template=True)
        profile = self.client.get("/api/profile").get_json()
        master_path = self.build_cover_letter_template_fixture()
        master_before = master_path.read_bytes()
        payload = self.upload_cover_letter_template(master_path)
        self.assertEqual(payload["cover_letter_template_status"], "DOCX template uploaded")

        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar",
                "title": "Python Developer",
                "location": "Toronto",
                "url": "https://example.com/job/cover-template",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        self.assertEqual(created["cover_letter_generation_status"], "generated")
        cover_path = Path(created["cover_letter_docx_path"])
        self.assertEqual(master_before, master_path.read_bytes())

        master_doc = Document(io.BytesIO(master_path.read_bytes()))
        cover_doc = Document(io.BytesIO(cover_path.read_bytes()))
        master_text = "\n".join(paragraph.text for paragraph in master_doc.paragraphs)
        cover_text = "\n".join(paragraph.text for paragraph in cover_doc.paragraphs)
        self.assertIn("{{DATE}}", master_text)
        self.assertNotIn("{{DATE}}", cover_text)
        self.assertNotIn("{{OPENING_PARAGRAPH}}", cover_text)
        self.assertIn("Northstar Hiring Team", cover_text)
        self.assertIn("Sincerely", cover_text)
        self.assertEqual(master_doc.paragraphs[0].style.name, cover_doc.paragraphs[0].style.name)
        self.assertEqual(master_doc.paragraphs[-1].alignment, cover_doc.paragraphs[-1].alignment)
        self.assertEqual(master_doc.paragraphs[-1].text, cover_doc.paragraphs[-1].text)

    def test_structured_style_template_replaces_role_specific_body_without_plain_fallback(self):
        self.seed_profile_for_documents(with_template=True)
        template = Path(self.tmpdir.name) / "pinterest_style_fixture.docx"
        document = Document()
        document.add_paragraph("Terry Mak")
        document.add_paragraph("Toronto, ON")
        document.add_paragraph("(647) 449-9669")
        document.add_paragraph("terry.hcmak@gmail.com")
        document.add_paragraph("")
        document.add_paragraph("October 1, 2026")
        document.add_paragraph("")
        document.add_paragraph("Hiring Manager")
        document.add_paragraph("Pinterest")
        document.add_paragraph("Toronto, ON")
        document.add_paragraph("Re: Software Engineer II, Web")
        document.add_paragraph("Dear Hiring Manager,")
        document.add_paragraph("I am writing to apply for a software role at Pinterest.")
        document.add_paragraph("I designed a restaurant operations platform using supported technologies.")
        document.add_paragraph("These experiences strengthened my ability to build maintainable web functionality.")
        document.add_paragraph("I am particularly interested in Pinterest's engineering work.")
        document.add_paragraph("")
        document.add_paragraph("Sincerely,")
        document.add_paragraph("Terry Mak")
        document.save(template)
        self.upload_cover_letter_template(template)

        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Gatekeeper Systems Inc",
                "title": "Software Engineer",
                "location": "Abbotsford, BC",
                "job_description": "Build and support Python and SQL software applications.",
                "required_skills": "Python, SQL",
            },
        ).get_json()
        cover_path = Path(created["cover_letter_docx_path"])
        cover = Document(cover_path)
        paragraphs = [paragraph.text for paragraph in cover.paragraphs if paragraph.text.strip()]
        joined = "\n".join(paragraphs)
        self.assertEqual(created["cover_letter_generation_status"], "generated")
        self.assertEqual(created["generation_debug"]["cover_result"]["output_path"], str(cover_path))
        self.assertIn("Gatekeeper Systems Inc", joined)
        self.assertIn("Dear Gatekeeper Systems Inc Hiring Team,", joined)
        self.assertNotIn("Pinterest", joined)
        self.assertIn("restaurant operations", joined.lower())
        self.assertNotRegex(joined, r"[-–—]")
        self.assertEqual(paragraphs[-2:], ["Sincerely,", "Terry Mak"])

    def test_structured_template_preserves_factual_body_paragraphs_and_logs_adaptation(self):
        template = Path(self.tmpdir.name) / "preserving_cover_letter_fixture.docx"
        document = Document()
        document.add_paragraph("Terry Mak")
        document.add_paragraph("Toronto, ON")
        document.add_paragraph("terry@example.com")
        document.add_paragraph("October 1, 2026")
        document.add_paragraph("Hiring Manager")
        document.add_paragraph("Pinterest")
        document.add_paragraph("Toronto, ON")
        document.add_paragraph("Re: Software Engineer II, Web")
        document.add_paragraph("Dear Hiring Manager,")
        opening = document.add_paragraph(
            "I am writing to apply for the Software Engineer II, Web position at Pinterest. "
            "With a Bachelor of Science in Computer Science, I bring relevant experience."
        )
        flowtally = document.add_paragraph(
            "Most recently, I designed and built Flowtally, an operations platform. "
            "I also built 25+ automated backend tests."
        )
        employers = document.add_paragraph(
            "At Oasis Elea Inc., I built internal web tooling and helped launch a storefront with "
            "40 products and 450+ variants. At Harmony 88 Co., I reduced data errors by 75 percent."
        )
        motivation = document.add_paragraph(
            "I am particularly interested in Pinterest's focus on web platform development and testing."
        )
        document.add_paragraph("Sincerely,")
        document.add_paragraph("Terry Mak")
        document.save(template)
        original_bytes = template.read_bytes()
        profile = {
            "name": "Terry Mak",
            "work_experience": "Built internal web tools with Python and SQL.",
            "projects": "Built Flowtally and automated backend tests.",
            "resume_text": "Terry Mak. Built internal web tools with Python and SQL. Built Flowtally and automated backend tests.",
            "cover_letter_docx_template_path": str(template),
            "cover_letter_no_dashes": "true",
        }
        job = {
            "company": "Gatekeeper Systems Inc",
            "title": "Software Engineer",
            "location": "Abbotsford, BC",
            "job_description": "Gatekeeper Systems Inc provides intelligent video and data solutions for safer transportation.",
        }
        output = app_module.generate_application_materials(job, profile, raise_on_validation=False)
        output_path = Path(self.tmpdir.name) / "preserving_cover_letter_output.docx"
        artifact = app_module.build_format_preserving_cover_letter(job, profile, output, output_path)
        generated = Document(output_path)
        generated_text = "\n".join(paragraph.text for paragraph in generated.paragraphs)
        self.assertIn("Software Engineer position at Gatekeeper Systems Inc", generated_text)
        self.assertIn("25+ automated backend tests", generated_text)
        self.assertIn("40 products and 450+ variants", generated_text)
        self.assertIn("reduced data errors by 75 percent", generated_text)
        self.assertIn("safer transportation", generated_text)
        self.assertNotIn("Pinterest", generated_text)
        self.assertNotRegex(generated_text, r"[-–—]")
        self.assertEqual(original_bytes, template.read_bytes())

        self.assertEqual(artifact["status"], "generated")
        self.assertGreaterEqual(len(artifact["change_log"]["preserved"]), 2)
        self.assertTrue(any("Flowtally" in item["original"] for item in artifact["change_log"]["preserved"]))
        self.assertTrue(any("Pinterest" in item["original"] and "Gatekeeper" in item["replacement"] for item in artifact["change_log"]["applied"]))

    def test_structured_template_replaces_selected_identity_email_hyperlink(self):
        template = Path(self.tmpdir.name) / "structured_hyperlinked_identity.docx"
        document = Document()
        document.add_paragraph("Template Applicant")
        contact = document.add_paragraph()
        self.add_hyperlink(contact, "mailto:template@example.com", "template@example.com")
        document.add_paragraph("October 1, 2026")
        document.add_paragraph("Hiring Manager")
        document.add_paragraph("Old Company")
        document.add_paragraph("Toronto, ON")
        document.add_paragraph("Re: Software Engineer")
        document.add_paragraph("Dear Hiring Manager,")
        document.add_paragraph("I built supported software systems and automated tests.")
        document.add_paragraph("I would welcome the opportunity to contribute.")
        document.add_paragraph("Sincerely,")
        document.add_paragraph("Template Applicant")
        document.save(template)
        profile = {
            "name": "Terry Mak",
            "email": "terry.hcmak@gmail.com",
            "work_experience": "Built supported software systems and automated tests.",
            "resume_text": "Terry Mak. Built supported software systems and automated tests.",
            "cover_letter_docx_template_path": str(template),
            "cover_letter_no_dashes": "true",
        }
        job = {"company": "Lakeview Loan Servicing", "title": "Backend Engineer", "location": "Toronto, Ontario", "job_description": "Build Python APIs and automated tests."}
        output = app_module.generate_application_materials(job, profile, raise_on_validation=False)
        output_path = Path(self.tmpdir.name) / "structured_hyperlinked_identity_output.docx"
        artifact = app_module.build_format_preserving_cover_letter(job, profile, output, output_path)
        self.assertEqual(artifact["status"], "generated", artifact)
        generated = Document(output_path)
        generated_text = "\n".join(paragraph.text for paragraph in app_module.iter_document_paragraphs(generated))
        self.assertEqual(generated_text.count("terry.hcmak@gmail.com"), 1)
        self.assertNotIn("template@example.com", generated_text)
        self.assertIn("Terry Mak", generated_text)
        self.assertEqual(
            [rel.target_ref for rel in generated.paragraphs[1].part.rels.values() if rel.reltype.endswith("/hyperlink")],
            ["mailto:terry.hcmak@gmail.com"],
        )

    def test_selected_cover_template_missing_fails_without_generic_fallback(self):
        self.seed_profile_for_documents(with_template=True)
        missing_template = Path(self.tmpdir.name) / "missing-cover-letter.docx"
        profile = {
            "name": "Terry Mak",
            "key_skills": "Python",
            "work_experience": "Built Python software.",
            "resume_text": "Terry Mak. Built Python software.",
            "cover_letter_docx_template_path": str(missing_template),
            "cover_letter_no_dashes": "true",
        }
        job = {
            "company": "Gatekeeper Systems Inc",
            "title": "Software Engineer",
            "location": "Abbotsford, BC",
            "job_description": "Build and support Python software.",
        }
        output = {"cover_letter": "A generated letter must never replace the selected template."}
        output_path = Path(self.tmpdir.name) / "missing-cover-letter-output.docx"
        artifact = app_module.build_format_preserving_cover_letter(job, profile, output, output_path)
        self.assertEqual(artifact["status"], "failed")
        self.assertEqual(artifact["change_log"]["template_mode"], "template_unavailable")
        self.assertFalse(output_path.exists())

    def test_shortlist_selected_cover_template_failure_preserves_previous_document(self):
        self.seed_profile_for_documents(with_template=True)
        profile = self.client.post("/api/candidate-profiles", json={"name": "Selected Profile", "description": "Built Python software.", "identity_mode": "shared"}).get_json()
        profile_id = profile["id"]
        resume_document = Document()
        resume_document.add_paragraph("Selected Profile")
        resume_document.add_paragraph("Summary")
        resume_document.add_paragraph("Software developer with Python experience.")
        resume_document.add_paragraph("Experience")
        resume_document.add_paragraph("Built Python software.")
        resume_document.add_paragraph("Skills")
        resume_document.add_paragraph("Python")
        resume_bytes = io.BytesIO()
        resume_document.save(resume_bytes)
        resume_upload = self.client.post(
            f"/api/candidate-profiles/{profile_id}/resumes/upload",
            data={"file": (io.BytesIO(resume_bytes.getvalue()), "base-resume.docx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(resume_upload.status_code, 201)
        resume = resume_upload.get_json()
        cover = self.client.post(
            f"/api/candidate-profiles/{profile_id}/cover-letter-templates",
            json={"name": "Missing template", "content": "Hiring Manager\\nExample Corp\\nRe: Developer\\nDear Hiring Manager,\\n{{OPENING_PARAGRAPH}}\\nSincerely,\\nTerry Mak"},
        ).get_json()
        with app_module.app.app_context():
            app_module.get_db().execute("UPDATE cover_letter_templates SET docx_path = ? WHERE id = ?", [str(Path(self.tmpdir.name) / "missing.docx"), cover["id"]])
            app_module.get_db().commit()
        fake_output = {
            "resume_text": "Selected Profile\nBuilt Python software.",
            "cover_letter": "A generated letter",
            "resume_validation": {"ok": True, "conflicting_phrases": []},
            "matched_keywords": [],
            "keyword_coverage": {},
            "job_hash": "",
            "candidate_evidence_hash": "",
            "tailoring_plan": {},
        }
        with patch.object(app_module, "generate_application_materials", return_value=fake_output), patch.object(app_module, "profile_has_generation_source", return_value=True):
            job = self.client.post(
                "/api/shortlist",
                json={
                    "company": "Gatekeeper Systems Inc",
                    "title": "Software Engineer",
                    "job_description": "Build and support Python software.",
                    "candidate_profile_id": profile_id,
                    "resume_id": resume["id"],
                    "cover_letter_template_id": cover["id"],
                },
            ).get_json()
        self.assertEqual(job.get("cover_letter_generation_status"), "failed", job)
        self.assertIn("template", job["cover_letter_generation_error"].lower())
        self.assertEqual(job["cover_letter_docx_path"], "")

    def test_cover_letter_motivation_and_evidence_follow_each_job(self):
        profile = {
            "name": "Alex Rivera",
            "key_skills": "Python, SQL, testing, documentation",
            "work_experience": "Built internal reporting workflows with Python and SQL for operations teams.",
            "projects": "Flowtally project: Added validation, migrations, and 25+ automated backend tests across reporting workflows.",
            "resume_text": "Alex Rivera. Built internal reporting workflows with Python and SQL for operations teams. Flowtally project: Added validation, migrations, and 25+ automated backend tests across reporting workflows.",
        }
        gatekeeper = {
            "company": "Gatekeeper Systems Inc",
            "title": "Software Engineer",
            "job_description": "Gatekeeper Systems Inc is a provider of intelligent video and data solutions for safer transportation. Build and test Python software.",
            "required_skills": "Python, testing",
        }
        finance = {
            "company": "LedgerWorks",
            "title": "Software Engineer",
            "job_description": "LedgerWorks is a financial software company building reporting and reconciliation tools. Develop Python services and document reliable workflows.",
            "required_skills": "Python, reporting, documentation",
        }
        gatekeeper_map = app_module.build_cover_letter_section_map(profile, gatekeeper, [], grounding.build_candidate_evidence(profile))
        finance_map = app_module.build_cover_letter_section_map(profile, finance, [], grounding.build_candidate_evidence(profile))
        self.assertIn("safer transportation", gatekeeper_map["motivation"].lower())
        self.assertNotIn("transportation", finance_map["motivation"].lower())
        self.assertIn("financial software", finance_map["motivation"].lower())
        self.assertIn("25+", gatekeeper_map["experience"] + gatekeeper_map["experience_secondary"])
        self.assertNotIn("role's focus on", gatekeeper_map["opening"].lower())
        self.assertNotIn("role's focus on", finance_map["opening"].lower())

        lakeview = {
            "company": "Lakeview Loan Servicing",
            "title": "Backend Engineer",
            "job_description": "Build Python and Flask APIs with PostgreSQL, authentication, permissions, migrations, validation, and automated testing.",
            "required_skills": "Python, Flask, PostgreSQL, APIs, testing",
        }
        lakeview_map = app_module.build_cover_letter_section_map(profile, lakeview, [], grounding.build_candidate_evidence(profile))
        self.assertIn("Python", lakeview_map["motivation"])
        self.assertIn("SQL", lakeview_map["motivation"])
        self.assertEqual(lakeview_map["motivation"].lower().count("python"), 1)
        self.assertNotIn("Kafka", lakeview_map["motivation"])
        self.assertNotIn("Kubernetes", lakeview_map["motivation"])

    def test_cover_letter_validation_uses_body_boundaries_for_contact_details(self):
        evidence = grounding.build_candidate_evidence(self.software_profile())
        header_safe = (
            "Terry Mak\n"
            "terry@example.com | Toronto, ON | linkedin.com/in/terrymak\n\n"
            "June 17, 2026\n\n"
            "Scotiabank - Toronto, ON\n\n"
            "Dear Scotiabank Hiring Team,\n\n"
            "I built Python Flask applications backed by SQL and documented workflows.\n\n"
            "Sincerely,\n"
            "Terry Mak"
        )
        header_validation = grounding.validate_cover_letter_output(header_safe, evidence, self.software_job())
        self.assertTrue(header_validation["ok"])

        body_leak = (
            "Terry Mak\n"
            "terry@example.com | Toronto, ON | linkedin.com/in/terrymak\n\n"
            "June 17, 2026\n\n"
            "Scotiabank - Toronto, ON\n\n"
            "Dear Scotiabank Hiring Team,\n\n"
            "I built Python Flask applications backed by SQL. Contact me at terry@example.com for details.\n\n"
            "Sincerely,\n"
            "Terry Mak"
        )
        body_validation = grounding.validate_cover_letter_output(body_leak, evidence, self.software_job())
        self.assertFalse(body_validation["ok"])
        self.assertTrue(any("Contact details detected in body paragraph" in item for item in body_validation["conflicting_phrases"]))

    def test_cover_letter_email_hyperlink_is_replaced_atomically(self):
        document = Document()
        contact = document.add_paragraph()
        contact.add_run("Email: ")
        self.add_hyperlink(contact, "mailto:old@example.com", "old@example.com")
        split_contact = document.add_paragraph()
        split_contact.add_run("old")
        split_contact.add_run("@example")
        split_contact.add_run(".com")

        self.assertTrue(app_module.replace_cover_letter_email(contact, "new@example.com"))
        self.assertTrue(app_module.replace_cover_letter_email(split_contact, "new@example.com"))
        self.assertEqual(contact.text, "Email: new@example.com")
        self.assertEqual(split_contact.text, "new@example.com")
        self.assertEqual(
            [rel.target_ref for rel in contact.part.rels.values() if rel.reltype.endswith("/hyperlink")],
            ["mailto:new@example.com"],
        )
        self.assertEqual(len(app_module.EMAIL_PATTERN.findall(contact.text)), 1)
        self.assertEqual(len(app_module.EMAIL_PATTERN.findall(split_contact.text)), 1)

    def test_format_preserving_cover_letter_updates_hyperlinked_identity_and_keeps_range_natural(self):
        source = Path(self.tmpdir.name) / "hyperlinked_cover_letter.docx"
        document = Document()
        heading = document.add_paragraph("Terry Mak")
        contact = document.add_paragraph()
        self.add_hyperlink(contact, "mailto:ter23rymhc@gmail.com", "ter23rymhc@gmail.com")
        document.add_paragraph("May 1, 2024")
        document.add_paragraph("Old Company")
        document.add_paragraph("Toronto, Ontario")
        document.add_paragraph("Dear Old Company Hiring Team,")
        document.add_paragraph("Re: Application for Software Engineer")
        document.add_paragraph("I am pleased to apply for the Software Engineer role at Old Company.")
        document.add_paragraph("I tracked 50–250 SKUs and built Python and Flask tools.")
        document.add_paragraph("My experience would support this team.")
        document.add_paragraph("Thank you for your consideration.")
        document.add_paragraph("Sincerely,")
        document.add_paragraph("Terry Mak")
        document.save(source)

        output_path = Path(self.tmpdir.name) / "generated_cover_letter.docx"
        profile = {"name": "Terry Mak", "email": "terry.hcmak@gmail.com", "cover_letter_docx_template_path": str(source), "cover_letter_no_dashes": "true"}
        job = {"company": "Lakeview Loan Servicing", "title": "Backend Engineer", "location": "Toronto, Ontario", "job_description": "Build Python and Flask APIs with PostgreSQL and automated testing."}
        result = app_module.build_format_preserving_cover_letter(
            job,
            profile,
            {"cover_letter": "", "candidate_evidence": grounding.build_candidate_evidence({"work_experience": "Built Python and Flask APIs with PostgreSQL and automated testing for internal operations workflows."})},
            output_path,
        )
        self.assertEqual(result["status"], "generated", result)
        generated = Document(output_path)
        generated_text = "\n".join(paragraph.text for paragraph in app_module.iter_document_paragraphs(generated))
        self.assertEqual(generated_text.count("terry.hcmak@gmail.com"), 1)
        self.assertNotIn("ter23rymhc@gmail.com", generated_text)
        self.assertIn("50 to 250 SKUs", generated_text)
        self.assertNotIn("50–250", generated_text)
        self.assertEqual(
            [rel.target_ref for rel in generated.paragraphs[1].part.rels.values() if rel.reltype.endswith("/hyperlink")],
            ["mailto:terry.hcmak@gmail.com"],
        )

    def test_resume_identity_uses_the_same_selected_email_as_cover_letter(self):
        document = Document()
        document.add_paragraph("Old Applicant")
        contact = document.add_paragraph()
        self.add_hyperlink(contact, "mailto:ter23rymhc@gmail.com", "ter23rymhc@gmail.com")
        self.assertTrue(app_module.replace_resume_identity(document, "Terry Mak", "terry.hcmak@gmail.com"))
        self.assertEqual(document.paragraphs[0].text, "Terry Mak")
        self.assertEqual(document.paragraphs[1].text, "terry.hcmak@gmail.com")
        self.assertNotIn("ter23rymhc@gmail.com", document.paragraphs[1].text)
        self.assertEqual(
            [rel.target_ref for rel in contact.part.rels.values() if rel.reltype.endswith("/hyperlink")],
            ["mailto:terry.hcmak@gmail.com"],
        )

    def test_no_dash_phrase_cleanup_does_not_duplicate_experience(self):
        self.assertEqual(
            tailoring_pipeline.remove_cover_letter_dashes("I have hands-on experience and tracked 50–250 SKUs."),
            "I have practical experience and tracked 50 to 250 SKUs.",
        )

    def test_scotiabank_cover_letter_template_preserves_structure_and_replaces_stale_content(self):
        self.seed_profile_for_documents(with_template=True)
        self.upload_resume_template(self.build_python_developer_resume_fixture())
        master_path = self.build_software_cover_letter_fixture()
        master_before = master_path.read_bytes()
        payload = self.upload_cover_letter_template(master_path)
        self.assertEqual(payload["cover_letter_template_status"], "DOCX template uploaded")

        created = self.client.post(
            "/api/shortlist",
            json=self.scotiabank_job(),
        ).get_json()
        self.assertEqual(created["cover_letter_generation_status"], "failed")
        self.assertIn("incompatible", created["cover_letter_generation_error"].lower())
        self.assertEqual(created["cover_letter_docx_path"], "")
        self.assertEqual(created["tailoring_changes"]["cover_letter"]["template_mode"], "template_rejected")
        self.assertEqual(master_before, master_path.read_bytes())

    def test_scotiabank_cover_letter_rewrites_all_visible_template_content(self):
        self.seed_profile_for_documents(with_template=True)
        self.upload_resume_template(self.build_python_developer_resume_fixture())
        master_path = self.build_centralsquare_cover_letter_fixture()
        self.upload_cover_letter_template(master_path)

        created = self.client.post(
            "/api/shortlist",
            json=self.scotiabank_job(),
        ).get_json()
        self.assertEqual(created["cover_letter_generation_status"], "failed")
        self.assertIn("incompatible", created["cover_letter_generation_error"].lower())
        self.assertEqual(created["cover_letter_docx_path"], "")
        self.assertEqual(created["tailoring_changes"]["cover_letter"]["template_mode"], "template_rejected")

    def test_cover_letter_validation_reports_hidden_metadata_location(self):
        source = self.build_cover_letter_template_fixture()
        hidden = Path(self.tmpdir.name) / "cover_letter_template_hidden_metadata.docx"
        with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(hidden, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "docProps/core.xml":
                    data = data.replace(b"python-docx", b"python-docx CentralSquare")
                zout.writestr(item, data)

        doc = Document(io.BytesIO(hidden.read_bytes()))
        warning = app_module.validate_cover_letter_document(doc, self.software_job(), {"name": "Terry Mak"})
        self.assertIsNotNone(warning)
        self.assertIn("metadata part docProps/core.xml", warning)

    def test_cover_letter_no_template_fallback_and_validation_guards(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar",
                "title": "Python Developer",
                "location": "Toronto",
                "url": "https://example.com/job/cover-fallback",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        self.assertEqual(created["cover_letter_generation_status"], "generated")
        cover_path = Path(created["cover_letter_docx_path"])
        cover_doc = Document(io.BytesIO(cover_path.read_bytes()))
        cover_paragraphs = [paragraph.text.strip() for paragraph in cover_doc.paragraphs if paragraph.text.strip()]
        cover_text = "\n".join(cover_paragraphs)
        cover_body_start = next(index for index, text in enumerate(cover_paragraphs) if text.startswith("I am pleased"))
        cover_body_end = next(index for index, text in enumerate(cover_paragraphs) if text.startswith("Thank you"))
        cover_body = cover_paragraphs[cover_body_start:cover_body_end]
        self.assertIn("Dear Northstar Hiring Team", cover_text)
        self.assertIn("Sincerely", cover_text)
        self.assertNotIn("EDUCATION", cover_text)
        self.assertNotIn("SKILLS", cover_text)
        self.assertNotIn("EXPERIENCE", cover_text)
        self.assertNotIn("PROJECTS", cover_text)
        self.assertEqual(len(cover_body), 3)
        self.assertGreaterEqual(len(" ".join(cover_body).split()), 50)
        self.assertLessEqual(len(" ".join(cover_body).split()), 140)
        self.assertLessEqual(max(len(paragraph.text) for paragraph in cover_doc.paragraphs), 400)
        self.assertNotIn("Terry Mak. Developer Intern at Northstar", cover_text)

        bad_validation = grounding.validate_cover_letter_output(
            "EDUCATION\nSKILLS\nEXPERIENCE\nPROJECTS",
            grounding.build_candidate_evidence(self.rece_profile()),
            self.ece_job(),
        )
        self.assertFalse(bad_validation["ok"])
        self.assertTrue(any("Resume section dump detected" in item for item in bad_validation["conflicting_phrases"]))

        prose_validation = grounding.validate_cover_letter_output(
            "I built Python and Flask applications backed by SQL and used validation, logging, and data-cleaning logic to improve reliability.",
            grounding.build_candidate_evidence(self.software_profile()),
            self.software_job(),
        )
        self.assertTrue(prose_validation["ok"])

        keyword_validation = grounding.validate_cover_letter_output(
            "Python, Flask, SQL, AWS, Linux, React, JavaScript, APIs, cloud, dashboards, automation.",
            grounding.build_candidate_evidence(self.software_profile("AWS")),
            self.software_job("AWS, Python"),
        )
        self.assertFalse(keyword_validation["ok"])
        self.assertTrue(any("Raw keyword list detected" in item for item in keyword_validation["conflicting_phrases"]))

        section_dump_validation = grounding.validate_cover_letter_output(
            "Skills\nPython Flask SQL\nExperience\nBuilt reports.",
            grounding.build_candidate_evidence(self.software_profile()),
            self.software_job(),
        )
        self.assertFalse(section_dump_validation["ok"])
        self.assertTrue(any("Resume section dump detected" in item for item in section_dump_validation["conflicting_phrases"]))

        long_paragraph = "Python " * 90
        length_validation = grounding.validate_cover_letter_output(
            long_paragraph,
            grounding.build_candidate_evidence(self.rece_profile()),
            self.ece_job(),
        )
        self.assertFalse(length_validation["ok"])

    def test_cover_letter_template_rejects_unrelated_retail_content_without_fallback(self):
        self.seed_profile_for_documents(with_template=True)
        retail_template = self.build_retail_cover_letter_fixture()
        payload = self.upload_cover_letter_template(retail_template)
        self.assertEqual(payload["cover_letter_template_status"], "DOCX template uploaded")

        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar",
                "title": "Python Developer",
                "location": "Toronto",
                "url": "https://example.com/job/cover-retail",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        self.assertEqual(created["cover_letter_generation_status"], "failed")
        self.assertIn("incompatible", created["cover_letter_generation_error"].lower())
        self.assertEqual(created["cover_letter_docx_path"], "")
        self.assertEqual(created["tailoring_changes"]["cover_letter"]["template_mode"], "template_rejected")

    def test_shortlist_restore_reuses_existing_row_and_regenerates_documents(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "RestorableCo",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/job/restorable",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        job_id = created["id"]

        removed = self.client.delete(f"/api/shortlist/{job_id}").get_json()
        self.assertTrue(removed["restored_to_saved"])
        self.assertEqual(self.client.get("/api/shortlist").get_json(), [])

        restored = self.client.post(
            "/api/shortlist",
            json={
                "company": "RestorableCo",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/job/restorable?utm_source=search",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        )
        self.assertEqual(restored.status_code, 200)
        restored_job = restored.get_json()
        self.assertTrue(restored_job["restored"])
        self.assertEqual(restored_job["id"], job_id)
        self.assertEqual(len(self.client.get("/api/shortlist").get_json()), 1)
        self.assertEqual(restored_job["resume_generation_status"], "generated")
        self.assertEqual(restored_job["cover_letter_generation_status"], "generated")

    def test_cover_failure_does_not_block_resume_download(self):
        self.seed_profile_for_documents(with_template=True)
        with patch.object(app_module, "build_format_preserving_cover_letter", side_effect=RuntimeError("cover exploded")):
            created = self.client.post(
                "/api/shortlist",
                json={
                    "company": "SplitCo",
                    "title": "Software Developer",
                    "location": "Toronto",
                    "url": "https://example.com/job/split-cover",
                    "job_description": "Python Flask SQL React dashboards cloud automation",
                },
            ).get_json()
        self.assertEqual(created["resume_generation_status"], "generated")
        self.assertEqual(created["cover_letter_generation_status"], "failed")
        self.assertIn("cover exploded", created["cover_letter_generation_error"])
        self.assertTrue(Path(created["resume_docx_path"]).exists())
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/resume").status_code, 200)
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/cover-letter").status_code, 404)

    def test_resume_failure_does_not_block_cover_download(self):
        self.seed_profile_for_documents(with_template=True)
        with patch.object(app_module, "build_format_preserving_resume", side_effect=RuntimeError("resume exploded")):
            created = self.client.post(
                "/api/shortlist",
                json={
                    "company": "SplitCo",
                    "title": "Software Developer",
                    "location": "Toronto",
                    "url": "https://example.com/job/split-resume",
                    "job_description": "Python Flask SQL React dashboards cloud automation",
                },
            ).get_json()
        self.assertEqual(created["resume_generation_status"], "failed")
        self.assertEqual(created["cover_letter_generation_status"], "generated")
        self.assertIn("resume exploded", created["resume_generation_error"])
        self.assertTrue(Path(created["cover_letter_docx_path"]).exists())
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/resume").status_code, 404)
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/cover-letter").status_code, 200)

    def test_failed_regeneration_preserves_previous_successful_documents(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "RecoveryCo",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/job/recovery",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        previous_resume = Path(created["resume_docx_path"])
        previous_cover = Path(created["cover_letter_docx_path"])
        with patch.object(app_module, "build_format_preserving_resume", side_effect=RuntimeError("resume retry failed")), patch.object(
            app_module, "build_format_preserving_cover_letter", side_effect=RuntimeError("cover retry failed")
        ):
            regenerated = self.client.post(f"/api/shortlist/{created['id']}/regenerate", json={}).get_json()

        self.assertEqual(regenerated["resume_generation_status"], "failed")
        self.assertEqual(regenerated["cover_letter_generation_status"], "failed")
        self.assertEqual(Path(regenerated["resume_docx_path"]), previous_resume)
        self.assertEqual(Path(regenerated["cover_letter_docx_path"]), previous_cover)
        self.assertTrue(previous_resume.is_file())
        self.assertTrue(previous_cover.is_file())
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/resume").status_code, 200)
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/cover-letter").status_code, 200)

    def test_pdf_only_profile_cannot_request_format_preserving_export(self):
        self.client.post(
            "/api/profile",
            json={
                "name": "Sparse Profile",
                "resume_text": "Built Python reports with SQL.",
                "work_experience": "Supported families, planned play-based learning, and completed developmental observation and documentation in a classroom setting.",
                "preferred_roles": "software developer",
                "key_skills": "Python, SQL",
            },
        )
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "SparseCo",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/sparse",
                "job_description": "Python SQL Flask",
            },
        ).get_json()
        self.assertEqual(created["resume_generation_status"], "failed")
        self.assertIn("Upload a DOCX resume", created["generation_error"])

    def test_output_docx_is_readable(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/job/readable",
                "job_description": "Python Flask SQL React dashboards cloud automation",
            },
        ).get_json()
        resume_path = Path(created["resume_docx_path"])
        opened = Document(io.BytesIO(resume_path.read_bytes()))
        self.assertGreater(len(opened.paragraphs), 0)
        self.assertIn("Terry Mak", opened.paragraphs[0].text)

    def test_shortlist_can_generate_from_selected_candidate_profile_documents(self):
        self.seed_profile_for_documents(with_template=True)
        current_profile = self.client.get("/api/profile").get_json()
        created_profile = self.client.post(
            "/api/candidate-profiles",
            json={"name": "SWE Resume", "description": "Python developer with documented API and SQL project experience.", "identity_mode": "distinct", "applicant_name": "Alex Rivera"},
        )
        self.assertEqual(created_profile.status_code, 201)
        profile_id = created_profile.get_json()["id"]
        resume_response = self.client.post(
            f"/api/candidate-profiles/{profile_id}/resumes",
            json={
                "name": "Alternate base resume",
                "content": "Python developer with documented API and SQL project experience and reliable testing practices.",
                "docx_path": current_profile["resume_docx_template_path"],
                "is_default": True,
            },
        )
        self.assertEqual(resume_response.status_code, 201)
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "ProfileCo",
                "title": "Python Developer",
                "location": "Toronto",
                "url": "https://example.com/job/profile-selection",
                "job_description": "Python SQL API testing",
                "candidate_profile_id": profile_id,
            },
        )
        self.assertEqual(created.status_code, 201)
        job = created.get_json()
        self.assertEqual(job["resume_generation_status"], "generated")
        self.assertEqual(job["cover_letter_generation_status"], "generated")
        self.assertIn("Alex_Rivera", Path(job["resume_docx_path"]).name)
        cover_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(Path(job["cover_letter_docx_path"]).read_bytes())).paragraphs)
        self.assertIn("Alex Rivera", cover_text)
        self.assertNotIn("SWE Resume", cover_text)

    def test_legacy_generate_without_resume_evidence_returns_helpful_400(self):
        response = self.client.post(
            "/api/generate",
            json={"company": "Validation Co", "title": "Developer", "job_description": "Python"},
        )
        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.get_json()["validation_error"])
        self.assertIn("Complete the Profile page", response.get_json()["error"])

    def test_shortlist_keyword_score_uses_selected_resume_without_generating_documents(self):
        self.seed_profile_for_documents()
        profile_response = self.client.post(
            "/api/candidate-profiles",
            json={"name": "Selected QA Profile", "identity_mode": "distinct", "applicant_name": "Alex Rivera"},
        )
        self.assertEqual(profile_response.status_code, 201)
        profile_id = profile_response.get_json()["id"]
        resume_path = self.build_formatted_resume_fixture()
        with resume_path.open("rb") as handle:
            uploaded = self.client.post(
                f"/api/candidate-profiles/{profile_id}/resumes/upload",
                data={"file": (handle, "selected.docx"), "name": "Selected résumé", "is_default": "true"},
                content_type="multipart/form-data",
            )
        self.assertEqual(uploaded.status_code, 201)
        resume_id = uploaded.get_json()["id"]
        job = self.client.post(
            "/api/jobs",
            json={"company": "Score Co", "title": "Python Developer", "status": "shortlisted", "job_description": "Python SQL APIs"},
        ).get_json()

        scored = self.client.post(
            f"/api/shortlist/{job['id']}/keyword-score",
            json={"candidate_profile_id": profile_id, "resume_id": resume_id},
        )
        self.assertEqual(scored.status_code, 200)
        body = scored.get_json()
        self.assertFalse(body["generated"])
        self.assertEqual(body["candidate_profile_id"], str(profile_id))
        self.assertGreaterEqual(body["keywords_score"]["total"], 1)
        saved = next(item for item in self.client.get("/api/jobs").get_json() if item["id"] == job["id"])
        self.assertFalse(saved.get("resume_docx_path"))
        self.assertFalse(saved.get("cover_letter_docx_path"))

    def test_shortlist_keyword_score_changes_when_selected_resume_changes(self):
        self.seed_profile_for_documents()
        profile_id = self.client.post(
            "/api/candidate-profiles",
            json={"name": "Resume Choice Profile", "identity_mode": "distinct", "applicant_name": "Alex Rivera"},
        ).get_json()["id"]
        first_path = self.build_formatted_resume_fixture()
        second_path = Path(self.tmpdir.name) / "unrelated_resume.docx"
        second_doc = Document()
        second_doc.add_paragraph("Alex Rivera")
        second_doc.add_paragraph("Customer support specialist with scheduling and communications experience.")
        second_doc.save(second_path)
        first = self.client.post(
            f"/api/candidate-profiles/{profile_id}/resumes",
            json={"name": "Technical résumé", "content": "Python SQL APIs", "docx_path": str(first_path), "is_default": True},
        ).get_json()
        second = self.client.post(
            f"/api/candidate-profiles/{profile_id}/resumes",
            json={"name": "Other résumé", "content": "Customer support", "docx_path": str(second_path), "is_default": False},
        ).get_json()
        job = self.client.post(
            "/api/jobs",
            json={"company": "Score Co", "title": "Python Developer", "status": "shortlisted", "job_description": "Python SQL APIs"},
        ).get_json()
        first_score = self.client.post(
            f"/api/shortlist/{job['id']}/keyword-score",
            json={"candidate_profile_id": profile_id, "resume_id": first["id"]},
        ).get_json()["keywords_score"]
        second_response = self.client.post(
            f"/api/shortlist/{job['id']}/keyword-score",
            json={"candidate_profile_id": profile_id, "resume_id": second["id"]},
        )
        self.assertEqual(second_response.status_code, 200, second_response.get_data(as_text=True))
        second_score = second_response.get_json()["keywords_score"]
        self.assertGreater(first_score["matched_count"], second_score["matched_count"])

    def test_shortlist_uses_keywords_score_as_primary_display_metric(self):
        source = (Path(__file__).parents[1] / "static" / "app.js").read_text(encoding="utf-8")
        shortlist = source.split("function renderShortlistCard", 1)[1].split("async function regenerateShortlist", 1)[0]
        discovery = source.split("function renderDiscoveryRow", 1)[1].split("function renderDiscoveryCard", 1)[0]
        self.assertIn("Keywords Score", shortlist)
        self.assertNotIn("Overall Fit</strong>", shortlist)
        self.assertNotIn("miniBreakdown(job)", shortlist)
        self.assertIn("Job Priority", discovery)
        self.assertNotIn("Overall score", discovery)

    def test_selected_candidate_profile_isolated_from_legacy_and_other_profiles(self):
        self.seed_profile_for_documents(with_template=True)
        legacy = self.client.post(
            "/api/profile",
            json={
                "name": "Legacy Candidate",
                "work_experience": "LegacyCorp maintained legacy systems and legacy customer workflows.",
                "resume_text": "LegacyCorp maintained legacy systems and legacy customer workflows.",
            },
        )
        self.assertEqual(legacy.status_code, 200)
        profiles = []
        for name, evidence in [
            ("Profile Alpha", "Alpha Systems built Python APIs and automated reporting for finance teams with careful validation."),
            ("Profile Beta", "Beta Labs managed unrelated Java services and mobile release operations."),
        ]:
            created_profile = self.client.post("/api/candidate-profiles", json={"name": name, "description": evidence})
            self.assertEqual(created_profile.status_code, 201)
            profile_id = created_profile.get_json()["id"]
            resume_response = self.client.post(
                f"/api/candidate-profiles/{profile_id}/resumes",
                json={
                    "name": f"{name} resume",
                    "content": evidence + " This work included documented tests and reliable handoffs for maintainers.",
                    "docx_path": str(self.build_formatted_resume_fixture()),
                    "is_default": True,
                },
            )
            self.assertEqual(resume_response.status_code, 201)
            profiles.append(profile_id)

        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Alpha Hiring",
                "title": "Python Developer",
                "location": "Toronto",
                "url": "https://example.com/job/profile-isolation",
                "job_description": "Python APIs validation reporting testing",
                "candidate_profile_id": profiles[0],
            },
        )
        self.assertEqual(created.status_code, 201)
        job = created.get_json()
        resume_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(Path(job["resume_docx_path"]).read_bytes())).paragraphs)
        cover_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(Path(job["cover_letter_docx_path"]).read_bytes())).paragraphs)
        self.assertIn("Northstar Labs", resume_text + cover_text)
        self.assertNotIn("Alpha Systems", resume_text + cover_text)
        self.assertNotIn("LegacyCorp", resume_text + cover_text)
        self.assertNotIn("Beta Labs", resume_text + cover_text)

    def test_structured_resume_evidence_produces_grounded_professional_documents(self):
        self.seed_profile_for_documents(with_template=True)
        resume_path = self.build_formatted_resume_fixture()
        sections = app_module.extract_docx_profile_sections(str(resume_path))
        self.assertIn("Northstar Labs", sections["work_experience"])
        self.assertNotIn("Skills:", sections["work_experience"])
        self.assertNotIn("Category:", sections["work_experience"])
        self.assertIn("Python", sections["key_skills"])

        current_profile = self.client.get("/api/profile").get_json()
        cover_template = self.upload_cover_letter_template(self.build_cover_letter_template_fixture())
        created_profile = self.client.post(
            "/api/candidate-profiles",
            json={
                "name": "Software Engineering",
                "description": "A focused profile for Python service work.",
                "identity_mode": "distinct",
                "applicant_name": "Alex Rivera",
                "applicant_email": "alex.rivera@example.test",
            },
        ).get_json()
        profile_id = created_profile["id"]
        self.client.post(
            f"/api/candidate-profiles/{profile_id}/resumes",
            json={
                "name": "Alex resume",
                "content": "Raw registration text should not replace structured DOCX evidence.",
                "docx_path": str(resume_path),
                "is_default": True,
            },
        )
        self.client.post(
            f"/api/candidate-profiles/{profile_id}/cover-letter-templates",
            json={
                "name": "Shared cover letter",
                "content": "{{date}}\n{{company}}\n{{location}}\nDear Hiring Manager,\n{{opening_paragraph}}\n{{experience_paragraph}}\n{{fit_paragraph}}\n{{closing_paragraph}}\nSincerely,\n{{applicant_name}}",
                "docx_path": cover_template["cover_letter_docx_template_path"],
                "is_default": True,
            },
        )
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Northstar Labs",
                "title": "Python Developer",
                "location": "Toronto",
                "url": "https://example.com/job/grounded-documents",
                "job_description": "Python and PostgreSQL service development with validation and documentation.",
                "candidate_profile_id": profile_id,
            },
        ).get_json()
        self.assertEqual(created["resume_generation_status"], "generated")
        self.assertEqual(created["cover_letter_generation_status"], "generated")
        resume_doc = Document(io.BytesIO(Path(created["resume_docx_path"]).read_bytes()))
        cover_doc = Document(io.BytesIO(Path(created["cover_letter_docx_path"]).read_bytes()))
        resume_text = "\n".join(paragraph.text for paragraph in resume_doc.paragraphs)
        cover_text = "\n".join(paragraph.text for paragraph in cover_doc.paragraphs)
        self.assertIn("Alex Rivera", resume_text + cover_text)
        self.assertIn("Software developer with Python automation", resume_text)
        self.assertNotIn("Candidate with supported experience", resume_text)
        self.assertNotIn("documented profile experience", resume_text.lower())
        self.assertNotIn("Skills Area Tools and practices Delivery", cover_text)
        self.assertNotIn("Python, Flask, SQL, and automation claims", cover_text)
        self.assertNotIn("Skills:", cover_text)
        self.assertNotRegex(cover_text, r"\s[-–—]\s")
        self.assertEqual(created["generation"]["generation_debug"]["resume_result"]["status"], "generated")

    def test_selected_document_validation_and_preview_are_actionable(self):
        self.seed_profile_for_documents(with_template=True)
        created_profile = self.client.post(
            "/api/candidate-profiles",
            json={"name": "Incomplete profile", "description": "A profile with enough documented evidence for a useful generation error."},
        )
        profile_id = created_profile.get_json()["id"]
        invalid = self.client.post(
            "/api/shortlist",
            json={
                "company": "Invalid Selection Co",
                "title": "Developer",
                "location": "Toronto",
                "job_description": "Python testing",
                "candidate_profile_id": profile_id,
                "resume_id": 999999,
            },
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("resume_id", invalid.get_json()["error"])

        valid = self.client.post(
            "/api/shortlist",
            json={
                "company": "Preview Co",
                "title": "Python Developer",
                "location": "Toronto",
                "job_description": "Python Flask SQL testing",
            },
        )
        self.assertEqual(valid.status_code, 201)
        job = valid.get_json()
        preview = self.client.get(f"/api/shortlist/{job['id']}/preview")
        self.assertEqual(preview.status_code, 200)
        self.assertTrue(preview.get_json()["resume"])
        self.assertTrue(preview.get_json()["cover_letter"])

    def test_generated_cover_letter_uses_current_date_and_default_no_dash_preference(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Current Date Co",
                "title": "Python Developer",
                "location": "Toronto",
                "job_description": "Python Flask SQL testing",
            },
        )
        self.assertEqual(created.status_code, 201)
        cover_doc = Document(io.BytesIO(Path(created.get_json()["cover_letter_docx_path"]).read_bytes()))
        cover_text = "\n".join(paragraph.text for paragraph in app_module.iter_document_paragraphs(cover_doc))
        self.assertRegex(cover_text, rf"{date.today().strftime('%B')}\s+0?{date.today().day},\s+{date.today().year}")
        self.assertNotRegex(cover_text, r"[-–—]")

    def test_profile_label_does_not_replace_shared_or_distinct_applicant_identity(self):
        self.seed_profile_for_documents(with_template=True)
        shared_profile = self.client.post(
            "/api/candidate-profiles",
            json={"name": "Business Applications", "description": "Python application delivery and testing.", "identity_mode": "shared"},
        ).get_json()
        shared_resume = self.client.post(
            f"/api/candidate-profiles/{shared_profile['id']}/resumes",
            json={"name": "Shared resume", "content": "Python application delivery and testing with documented handoffs.", "docx_path": self.client.get("/api/profile").get_json()["resume_docx_template_path"], "is_default": True},
        )
        self.assertEqual(shared_resume.status_code, 201)
        shared_job = self.client.post(
            "/api/shortlist",
            json={"company": "Shared Co", "title": "Developer", "job_description": "Python testing", "candidate_profile_id": shared_profile["id"]},
        ).get_json()
        shared_cover = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(Path(shared_job["cover_letter_docx_path"]).read_bytes())).paragraphs)
        self.assertIn("Terry Mak", shared_cover)
        self.assertNotIn("Business Applications", shared_cover)

        distinct = self.client.post(
            "/api/candidate-profiles",
            json={"name": "SWE Resume", "description": "Python application delivery and testing.", "identity_mode": "distinct", "applicant_name": "Jordan Lee", "applicant_email": "jordan@example.com"},
        ).get_json()
        distinct_resume = self.client.post(
            f"/api/candidate-profiles/{distinct['id']}/resumes",
            json={"name": "Jordan resume", "content": "Python application delivery and testing with documented handoffs.", "docx_path": self.client.get("/api/profile").get_json()["resume_docx_template_path"], "is_default": True},
        )
        self.assertEqual(distinct_resume.status_code, 201)
        distinct_job = self.client.post(
            "/api/shortlist",
            json={"company": "Distinct Co", "title": "Developer", "job_description": "Python testing", "candidate_profile_id": distinct["id"]},
        ).get_json()
        distinct_cover = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(Path(distinct_job["cover_letter_docx_path"]).read_bytes())).paragraphs)
        distinct_resume_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(Path(distinct_job["resume_docx_path"]).read_bytes())).paragraphs)
        self.assertIn("Jordan Lee", distinct_cover + distinct_resume_text)
        self.assertNotIn("SWE Resume", distinct_cover + distinct_resume_text)

    def test_keyword_only_evidence_does_not_create_fabricated_achievements(self):
        self.client.post(
            "/api/profile",
            json={"name": "Evidence Candidate", "resume_text": "Monitoring validation with PostgreSQL and routine reporting.", "work_experience": "Monitoring validation with PostgreSQL and routine reporting."},
        )
        self.upload_master_resume_template()
        created = self.client.post(
            "/api/shortlist",
            json={"company": "Evidence Co", "title": "Developer", "job_description": "monitoring validation PostgreSQL"},
        ).get_json()
        cover = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(Path(created["cover_letter_docx_path"]).read_bytes())).paragraphs)
        self.assertNotIn("300,000", cover)
        self.assertNotIn("90%", cover)
        self.assertNotIn("Harmony 88", cover)
        self.assertIn("monitor validation with postgresql", cover.lower())

    def test_cover_letter_does_not_attribute_prior_work_to_hiring_company(self):
        profile = {
            "name": "Alex Rivera",
            "email": "alex.rivera@example.test",
            "location": "Toronto, ON",
            "work_experience": "Northstar Labs — Business Analyst. Led requirements workshops and translated stakeholder needs into SQL reporting workflows.",
            "resume_text": "Alex Rivera\nNorthstar Labs — Business Analyst\nLed requirements workshops and translated stakeholder needs into SQL reporting workflows.",
            "key_skills": "SQL, requirements workshops, documentation",
            "cover_letter_no_dashes": "true",
        }
        job = {
            "company": "Cedar Bridge Consulting",
            "title": "Business Analyst",
            "location": "Toronto, ON",
            "job_description": "Lead requirements workshops, create SQL reports, document processes, and coordinate validation.",
            "required_skills": "SQL, requirements workshops",
        }
        output = app_module.generate_application_materials(job, profile)
        letter = output["cover_letter"]
        self.assertIn("Cedar Bridge Consulting", letter)
        self.assertIn("At Northstar Labs", letter)
        self.assertNotIn("At Cedar Bridge Consulting,", letter)
        self.assertEqual(letter.lower().count("role's focus on"), 1)
        self.assertNotRegex(letter, r"[-–—]")

    def test_cedar_bridge_docx_preview_and_metadata_keep_prior_employer(self):
        self.seed_profile_for_documents(with_template=True)
        self.upload_cover_letter_template(self.build_cover_letter_template_fixture())
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Cedar Bridge Consulting",
                "title": "Business Analyst",
                "location": "Toronto, ON",
                "url": "https://example.com/job/cedar-bridge-employer-check",
                "job_description": "Lead requirements workshops, create SQL reports, document processes, and coordinate validation.",
            },
        ).get_json()
        preview = self.client.get(f"/api/shortlist/{created['id']}/preview").get_json()
        downloaded = self.client.get(f"/api/shortlist/{created['id']}/cover-letter")
        document = Document(io.BytesIO(downloaded.data))
        docx_text = "\n".join(paragraph.text for paragraph in app_module.iter_document_paragraphs(document) if paragraph.text.strip())
        del document
        self.assertEqual(preview["cover_letter_source"], "successful_docx")
        self.assertEqual(preview["cover_letter"], docx_text)
        self.assertIn("At Northstar", docx_text)
        self.assertNotIn("At Cedar Bridge Consulting,", docx_text)
        self.assertEqual(created["generation_debug"]["final_cover_letter_text"], preview["cover_letter"])
        self.assertNotRegex(docx_text, r"[-–—]")
        gc.collect()

    def test_preview_matches_successful_docx_and_survives_failed_retry(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post(
            "/api/shortlist",
            json={"company": "Preview Co", "title": "Developer", "location": "Toronto", "job_description": "Python testing"},
        ).get_json()
        preview_before = self.client.get(f"/api/shortlist/{created['id']}/preview").get_json()
        downloaded_resume = self.client.get(f"/api/shortlist/{created['id']}/resume")
        downloaded_cover = self.client.get(f"/api/shortlist/{created['id']}/cover-letter")
        def docx_text(payload: bytes) -> str:
            document = Document(io.BytesIO(payload))
            value = "\n".join(paragraph.text for paragraph in app_module.iter_document_paragraphs(document) if paragraph.text.strip())
            del document
            return value
        self.assertEqual(preview_before["resume_source"], "successful_docx")
        self.assertEqual(preview_before["cover_letter_source"], "successful_docx")
        self.assertEqual(preview_before["resume"], docx_text(downloaded_resume.data))
        self.assertEqual(preview_before["cover_letter"], docx_text(downloaded_cover.data))
        self.assertNotRegex(docx_text(downloaded_cover.data), r"[-–—]")
        with patch.object(app_module, "build_format_preserving_resume", side_effect=RuntimeError("resume retry failed")), patch.object(app_module, "build_format_preserving_cover_letter", side_effect=RuntimeError("cover retry failed")):
            failed = self.client.post(f"/api/shortlist/{created['id']}/regenerate", json={}).get_json()
        self.assertEqual(failed["resume_generation_status"], "failed")
        self.assertEqual(failed["cover_letter_generation_status"], "failed")
        preview_after = self.client.get(f"/api/shortlist/{created['id']}/preview").get_json()
        self.assertEqual(preview_after["resume"], preview_before["resume"])
        self.assertEqual(preview_after["cover_letter"], preview_before["cover_letter"])
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/resume").status_code, 200)
        self.assertEqual(self.client.get(f"/api/shortlist/{created['id']}/cover-letter").status_code, 200)
        gc.collect()

    def test_empty_projects_or_education_do_not_fabricate(self):
        profile = self.rece_profile()
        profile["projects"] = ""
        profile["education"] = ""
        output = app_module.generate_application_materials(self.ece_job(), profile)
        self.assertNotIn("project", output["summary"].lower())
        self.assertNotIn("education", output["summary"].lower())

    def test_cache_reuse_and_sentence_transformer_fallback(self):
        cache_dir = Path(self.tmpdir.name) / "semantic_cache"
        profile = {"work_experience": "Built Flask APIs."}
        job = {"required_skills": "Flask APIs"}
        first = semantic_matching.analyze_resume_match(job, profile, cache_dir=cache_dir, embedding_backend=FakeEmbeddingBackend())
        second = semantic_matching.analyze_resume_match(job, profile, cache_dir=cache_dir, embedding_backend=FakeEmbeddingBackend())
        self.assertFalse(first["cache_hit"])
        self.assertTrue(second["cache_hit"])

        fallback = semantic_matching.analyze_resume_match(job, profile, cache_dir=None, use_semantic=False)
        self.assertEqual(fallback["model"], "keyword-fallback")
        self.assertFalse(fallback["model_available"])

    def test_profile_changes_cause_different_score(self):
        job = {
            "title": "Junior Developer",
            "location": "Toronto",
            "job_description": "Required 2 years of experience. Python SQL APIs",
            "salary_source": "missing",
        }
        junior = {"target_seniority_levels": "junior", "maximum_required_years": "2", "key_skills": "Python, SQL"}
        senior_only = {"target_seniority_levels": "director", "maximum_required_years": "20", "key_skills": "Python, SQL"}
        self.assertNotEqual(
            app_module.score_job(job, junior)["seniority_score"],
            app_module.score_job(job, senior_only)["seniority_score"],
        )

    def seed_profile_for_documents(self, with_template: bool = False, resume_template: str = "formatted"):
        self.client.post(
            "/api/profile",
            json={
                "name": "Terry Mak",
                "preferred_roles": "software developer",
                "preferred_locations": "Toronto, remote",
                "key_skills": "Python, Flask, SQL, React, dashboards, cloud",
                "projects": "Built a Flask job tracker with SQLite, scoring, dashboards, and document generation.",
                "education": "Computer programming diploma, Toronto.",
                "work_experience": "Developer Intern at Northstar from 2024 to 2025. Built Python automation, cloud deployment notes, and SQL reports for operations teams.",
                "resume_text": "Terry Mak. Developer Intern at Northstar from 2024 to 2025. made Python automation, Flask tools, SQL reports, React interfaces, and dashboards. Education: Computer programming diploma.",
            },
        )
        if with_template:
            if resume_template == "no_summary":
                self.upload_resume_template(self.build_resume_without_summary_fixture())
            else:
                self.upload_master_resume_template()

    def seed_python_developer_profile(self):
        self.client.post(
            "/api/profile",
            json={
                "name": "Terry Mak",
                "preferred_roles": "software developer",
                "preferred_locations": "Toronto, remote",
                "key_skills": "Python, SQL, database design, AWS, Linux, testing, debugging, documentation, collaboration, Flask, React, cloud",
                "projects": "Built a Python inventory tool with database design, AWS deployment scripts, Linux server checks, testing, and documentation for a small operations team.",
                "education": "Computer programming diploma, George Brown College.",
                "work_experience": "Software Developer at Northstar Labs from 2021 to 2025. Implemented robust validation, logging, and data cleaning logic to detect malformed inputs and reduce manual data entry errors by 90%. Designed data models and application logic for products, supplier specific items, and pricing records. Built backend integrations with real-time data from Retool and Google Docs for operations reporting and handoffs. Worked with cross-functional teammates to debug issues and document release steps.",
                "resume_text": "Terry Mak. Software Developer at Northstar Labs from 2021 to 2025. Implemented robust validation, logging, and data cleaning logic to detect malformed inputs and reduce manual data entry errors by 90%. Designed data models and application logic for products, supplier specific items, and pricing records. Built backend integrations with real-time data from Retool and Google Docs for operations reporting and handoffs. Worked with cross-functional teammates to debug issues and document release steps. Education: Computer programming diploma, George Brown College.",
            },
        )

    def upload_master_resume_template(self):
        self.upload_resume_template(self.build_formatted_resume_fixture())

    def upload_resume_template(self, master: Path):
        with open(master, "rb") as handle:
            response = self.client.post(
                "/api/profile/upload",
                data={"file": (handle, "master_resume.docx")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("resume_docx_template_path", response.get_json())

    def upload_cover_letter_template(self, master: Path):
        with open(master, "rb") as handle:
            response = self.client.post(
                "/api/profile/upload",
                data={"file": (handle, "cover_letter_template.docx"), "template_kind": "cover_letter"},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertIn("cover_letter_docx_template_path", payload)
        return payload

    def build_formatted_resume_fixture(self):
        path = Path(self.tmpdir.name) / "master_resume_fixture.docx"
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Terry Mak")
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(18)

        summary_heading = doc.add_paragraph()
        summary_heading.style = "Heading 1"
        summary_heading.add_run("Professional Summary")
        summary = doc.add_paragraph()
        summary_run = summary.add_run("Software developer with Python automation and SQL reporting experience.")
        summary_run.font.name = "Georgia"
        summary_run.font.size = Pt(11)

        experience_heading = doc.add_paragraph()
        experience_heading.style = "Heading 1"
        experience_heading.add_run("Work Experience")
        experience = doc.add_paragraph()
        experience.add_run("Software Developer").bold = True
        experience.add_run(" | Northstar Labs | 2021-2025")

        bullet1 = doc.add_paragraph(style="List Bullet")
        bullet1.add_run("made Python automation, Flask tools, SQL reports, React interfaces, and dashboards.")
        bullet2 = doc.add_paragraph(style="List Bullet")
        bullet2_run1 = bullet2.add_run("Maintained team documentation standards")
        bullet2_run1.italic = True
        bullet2.add_run(" across internal projects.")

        skills_heading = doc.add_paragraph()
        skills_heading.style = "Heading 1"
        skills_heading.add_run("Skills")
        skills = doc.add_paragraph()
        skills.add_run("Skills: ").bold = True
        skills.add_run("React, dashboards, Python, SQL, Flask, cloud")

        categories = doc.add_table(rows=2, cols=3)
        categories.style = "Table Grid"
        categories.cell(0, 0).text = "Languages"
        categories.cell(0, 1).text = "Technologies"
        categories.cell(0, 2).text = "Tools & Platforms"
        categories.cell(1, 0).text = "Python, SQL"
        categories.cell(1, 1).text = "Flask, React"
        categories.cell(1, 2).text = "cloud, dashboards"

        contact = doc.add_paragraph()
        contact.add_run("Portfolio: ")
        self.add_hyperlink(contact, "https://example.com/portfolio", "example.com/portfolio")

        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Category"
        table.cell(0, 1).text = "Detail"
        table.cell(1, 0).text = "Certification"
        table.cell(1, 1).text = "Computer programming diploma"

        doc.save(path)
        return path

    def build_resume_without_summary_fixture(self):
        path = Path(self.tmpdir.name) / "resume_without_summary_fixture.docx"
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Terry Mak")
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(18)

        contact = doc.add_paragraph()
        contact.add_run("terry@example.com | LinkedIn | GitHub")

        experience_heading = doc.add_paragraph()
        experience_heading.style = "Heading 1"
        experience_heading.add_run("Work Experience")
        experience = doc.add_paragraph()
        experience.add_run("Software Developer").bold = True
        experience.add_run(" | Northstar Labs | 2021-2025")

        bullet1 = doc.add_paragraph(style="List Bullet")
        bullet1.add_run("made Python automation, Flask tools, SQL reports, React interfaces, and dashboards.")
        bullet2 = doc.add_paragraph(style="List Bullet")
        bullet2.add_run("Maintained team documentation standards across internal projects.")

        education_heading = doc.add_paragraph()
        education_heading.style = "Heading 1"
        education_heading.add_run("Education")
        college = doc.add_paragraph()
        college.add_run("George Brown College")
        diploma = doc.add_paragraph()
        diploma.add_run("Computer programming diploma | 2023")

        skills_heading = doc.add_paragraph()
        skills_heading.style = "Heading 1"
        skills_heading.add_run("Skills")
        skills = doc.add_paragraph()
        skills.add_run("Python, Flask, SQL, React")

        contact_two = doc.add_paragraph()
        contact_two.add_run("Portfolio: example.com/portfolio")

        doc.save(path)
        return path

    def build_python_developer_resume_fixture(self):
        path = Path(self.tmpdir.name) / "python_developer_resume_fixture.docx"
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Terry Mak")
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(18)

        contact = doc.add_paragraph()
        contact.add_run("terry@example.com | LinkedIn | GitHub")

        summary_heading = doc.add_paragraph()
        summary_heading.style = "Heading 1"
        summary_heading.add_run("Professional Summary")
        summary = doc.add_paragraph()
        summary.add_run("Software developer with Python, SQL, AWS, Linux, and documentation experience.")

        experience_heading = doc.add_paragraph()
        experience_heading.style = "Heading 1"
        experience_heading.add_run("Work Experience")
        experience = doc.add_paragraph()
        experience.add_run("Software Developer").bold = True
        experience.add_run(" | Northstar Labs | 2021-2025")

        bullet1 = doc.add_paragraph(style="List Bullet")
        bullet1.add_run("Implemented robust validation, logging, and data cleaning logic to detect malformed inputs and reduce manual data entry errors by 90%.")
        bullet2 = doc.add_paragraph(style="List Bullet")
        bullet2.add_run("Worked with cross-functional teammates to debug issues and document release steps across Python and SQL tools.")

        projects_heading = doc.add_paragraph()
        projects_heading.style = "Heading 1"
        projects_heading.add_run("Projects")
        project_title = doc.add_paragraph()
        project_title.add_run("Inventory Reconciler").bold = True
        project_title.add_run(" | Personal Project | 2024")
        project_bullet1 = doc.add_paragraph(style="List Bullet")
        project_bullet1.add_run("Designed data models and application logic for products, supplier specific items, and pricing records in an internal tool.")
        project_bullet2 = doc.add_paragraph(style="List Bullet")
        project_bullet2.add_run("Built backend integrations with real-time data from Retool and Google Docs for reporting and handoffs.")

        skills_heading = doc.add_paragraph()
        skills_heading.style = "Heading 1"
        skills_heading.add_run("Skills")
        skills = doc.add_paragraph()
        skills.add_run("Skills: ").bold = True
        skills.add_run("React, dashboards, AWS, Linux, Python, SQL, Flask, documentation, collaboration, testing")

        education_heading = doc.add_paragraph()
        education_heading.style = "Heading 1"
        education_heading.add_run("Education")
        college = doc.add_paragraph()
        college.add_run("George Brown College")
        diploma = doc.add_paragraph()
        diploma.add_run("Computer programming diploma | 2023")

        contact_two = doc.add_paragraph()
        contact_two.add_run("Portfolio: example.com/portfolio")

        doc.save(path)
        return path

    def build_asin_resume_fixture(self):
        path = Path(self.tmpdir.name) / "asin_resume_fixture.docx"
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("Terry Mak")
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(18)

        summary_heading = doc.add_paragraph()
        summary_heading.style = "Heading 1"
        summary_heading.add_run("Professional Summary")
        summary = doc.add_paragraph()
        summary.add_run("Backend developer with PostgreSQL, Python, automation, and reliability experience.")

        experience_heading = doc.add_paragraph()
        experience_heading.style = "Heading 1"
        experience_heading.add_run("Work Experience")
        experience = doc.add_paragraph()
        experience.add_run("Harmony 88 Co.").bold = True
        experience.add_run(" | Backend Developer | 2023-2025")

        bullet1 = doc.add_paragraph(style="List Bullet")
        bullet1.add_run("Built an automated ASIN monitoring system that tracked pricing, rank, and rating trends across 50-250 SKUs per session, delivering hourly insights over custom time frames and storing 300,000+ records in PostgreSQL.")
        bullet2 = doc.add_paragraph(style="List Bullet")
        bullet2.add_run("Implemented robust validation, logging, and data cleaning logic to detect malformed inputs and reduce manual data entry errors by 90%.")
        bullet3 = doc.add_paragraph(style="List Bullet")
        bullet3.add_run("Built backend integrations with real-time data from Retool and Google Docs for operations reporting and handoffs.")

        skills_heading = doc.add_paragraph()
        skills_heading.style = "Heading 1"
        skills_heading.add_run("Skills")
        skills = doc.add_paragraph()
        skills.add_run("Skills: ").bold = True
        skills.add_run("Python, PostgreSQL, Flask, AWS, dashboards, automation, APIs")

        doc.save(path)
        return path

    def build_cover_letter_template_fixture(self):
        path = Path(self.tmpdir.name) / "cover_letter_template_fixture.docx"
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_run = heading.add_run("Terry Mak")
        heading_run.bold = True
        heading_run.font.name = "Aptos"
        heading_run.font.size = Pt(18)

        date_para = doc.add_paragraph()
        date_run = date_para.add_run("{{DATE}}")
        date_run.font.name = "Georgia"
        date_run.font.size = Pt(11)

        greeting = doc.add_paragraph()
        greeting.add_run("Dear {{HIRING_MANAGER}},")

        opening = doc.add_paragraph()
        opening.add_run("{{OPENING_PARAGRAPH}}")

        experience = doc.add_paragraph()
        experience.add_run("{{EXPERIENCE_PARAGRAPH}}")

        fit = doc.add_paragraph()
        fit.add_run("{{FIT_PARAGRAPH}}")

        closing = doc.add_paragraph()
        closing.add_run("{{CLOSING_PARAGRAPH}}")

        signoff = doc.add_paragraph()
        signoff.add_run("Sincerely,")

        signature = doc.add_paragraph()
        signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        signature_run = signature.add_run("Terry Mak")
        signature_run.bold = True

        doc.save(path)
        return path

    def build_retail_cover_letter_fixture(self):
        path = Path(self.tmpdir.name) / "retail_cover_letter_template_fixture.docx"
        doc = Document()
        heading = doc.add_paragraph()
        heading.add_run("Terry Mak")
        doc.add_paragraph("Retail customer service associate with supermarket and cashier experience.")
        doc.add_paragraph("Managed customer service duties, inventory, and merchandising for a busy store.")
        doc.add_paragraph("Sincerely,")
        doc.add_paragraph("Terry Mak")
        doc.save(path)
        return path

    def build_software_cover_letter_fixture(self):
        path = Path(self.tmpdir.name) / "software_cover_letter_template_fixture.docx"
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run("Terry Mak")

        contact = doc.add_paragraph()
        contact.add_run("terry@example.com | Toronto, ON | linkedin.com/in/terrymak")

        date_para = doc.add_paragraph()
        date_para.add_run("May 1, 2024")

        company_para = doc.add_paragraph()
        company_para.add_run("CentralSquare")

        location_para = doc.add_paragraph()
        location_para.add_run("Toronto, Ontario")

        greeting = doc.add_paragraph()
        greeting.add_run("Dear CentralSquare Hiring Team,")

        opening = doc.add_paragraph()
        opening.add_run("I am excited to apply for the Software Developer co-op role at CentralSquare for an 8/12-month placement.")

        experience = doc.add_paragraph()
        experience.add_run("I supported profile content and ATS-oriented language for the role.")

        fit = doc.add_paragraph()
        fit.add_run("My saved profile matches the posting and I would bring grounded wording.")

        closing = doc.add_paragraph()
        closing.add_run("Thank you for your consideration. I look forward to hearing from you.")

        signoff = doc.add_paragraph()
        signoff.add_run("Sincerely,")

        signature = doc.add_paragraph()
        signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        signature.add_run("CentralSquare Applicant")

        doc.save(path)
        return path

    def build_centralsquare_cover_letter_fixture(self):
        path = Path(self.tmpdir.name) / "centralsquare_cover_letter_template_fixture.docx"
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        header = section.header.paragraphs[0]
        header.text = "hochitterry.mak@georgebrown.ca | CentralSquare Technologies"

        footer = section.footer.paragraphs[0]
        footer.text = "CentralSquare Technologies | Remote - Ontario"

        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run("Terry Mak")

        contact = doc.add_paragraph()
        contact.add_run("hochitterry.mak@georgebrown.ca | Toronto, Ontario")

        date_para = doc.add_paragraph()
        date_para.add_run("May 1, 2024")

        company_para = doc.add_paragraph()
        company_para.add_run("CentralSquare")

        location_para = doc.add_paragraph()
        location_para.add_run("Remote - Ontario")

        greeting = doc.add_paragraph()
        greeting.add_run("Dear Hiring Manager,")

        subject = doc.add_paragraph()
        subject.add_run("Re: Software Developer Co-op – Remote, Ontario")

        opening = doc.add_paragraph()
        opening.add_run("I am excited to apply for the Software Developer Co-op role at CentralSquare for an 8- or 12-month term.")

        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "CentralSquare Technologies"
        table.cell(0, 1).text = "CentralSquare's technology initiatives"

        experience = doc.add_paragraph()
        experience.add_run("I supported profile content and ATS-oriented language for the role.")

        fit = doc.add_paragraph()
        fit.add_run("My saved profile matches the posting and I would bring grounded wording.")

        closing = doc.add_paragraph()
        closing.add_run("Thank you for your consideration. I look forward to hearing from you.")

        signoff = doc.add_paragraph()
        signoff.add_run("Sincerely,")

        signature = doc.add_paragraph()
        signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        signature.add_run("CentralSquare Applicant")

        doc.save(path)
        return path

    def add_hyperlink(self, paragraph, url, text):
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0000FF")
        rPr.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rPr.append(underline)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def test_shortlist_persistence_generation_safe_names_and_downloads(self):
        self.seed_profile_for_documents(with_template=True)
        response = self.client.post(
            "/api/shortlist",
            json={
                "company": "Acme/Tools",
                "title": "Junior Software Developer:*?",
                "location": "Toronto",
                "url": "https://example.com/job/abc",
                "application_url": "https://example.com/apply/abc",
                "source": "LinkedIn",
                "salary": "CAD $70,000-$85,000 per year",
                "job_description": "Python Flask SQL React dashboards APIs",
            },
        )
        self.assertEqual(response.status_code, 201)
        job = response.get_json()
        self.assertEqual(job["resume_generation_status"], "generated")
        self.assertEqual(job["cover_letter_generation_status"], "generated")
        resume_path = Path(job["resume_docx_path"])
        cover_path = Path(job["cover_letter_docx_path"])
        self.assertTrue(resume_path.exists())
        self.assertTrue(cover_path.exists())
        self.assertNotRegex(resume_path.name, r'[<>:"/\\|?*]')

        resume = self.client.get(f"/api/shortlist/{job['id']}/resume")
        cover = self.client.get(f"/api/shortlist/{job['id']}/cover-letter")
        self.assertEqual(resume.status_code, 200)
        self.assertEqual(cover.status_code, 200)
        self.assertTrue(resume.data.startswith(b"PK"))

        listed = self.client.get("/api/shortlist").get_json()
        self.assertEqual(len(listed), 1)

    def test_regenerate_both_serves_fresh_files_and_rewrites_keyword_lists(self):
        self.client.post(
            "/api/profile",
            json={
                "name": "Terry Mak",
                "preferred_roles": "software developer, backend developer",
                "preferred_locations": "Toronto",
                "key_skills": "Python, PostgreSQL, Flask, AWS, automation, dashboards, validation, logging",
                "projects": "Harmony 88 Co. built automated ASIN monitoring for pricing and rating trends with PostgreSQL storage.",
                "work_experience": (
                    "Harmony 88 Co. built an automated ASIN monitoring system that tracked pricing, rank, and rating trends across 50-250 SKUs per session, "
                    "delivering hourly insights over custom time frames and storing 300,000+ records in PostgreSQL. "
                    "Implemented robust validation, logging, and data cleaning logic to detect malformed inputs and reduce manual data entry errors by 90%. "
                    "Built backend integrations with real-time data from Retool and Google Docs for operations reporting and handoffs."
                ),
                "resume_text": (
                    "Terry Mak. Harmony 88 Co. built an automated ASIN monitoring system that tracked pricing, rank, and rating trends across 50-250 SKUs per session, "
                    "delivering hourly insights over custom time frames and storing 300,000+ records in PostgreSQL. "
                    "Implemented robust validation, logging, and data cleaning logic to detect malformed inputs and reduce manual data entry errors by 90%. "
                    "Built backend integrations with real-time data from Retool and Google Docs for operations reporting and handoffs."
                ),
            },
        )
        self.upload_resume_template(self.build_asin_resume_fixture())
        self.upload_cover_letter_template(self.build_cover_letter_template_fixture())

        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "Scotiabank",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/job/scotiabank",
                "job_description": "Python PostgreSQL Flask AWS automation dashboards APIs validation logging documentation cross-functional collaboration",
            },
        ).get_json()
        job_id = created["id"]

        regenerated = self.client.post(f"/api/shortlist/{job_id}/regenerate", json={})
        self.assertEqual(regenerated.status_code, 200)
        job = regenerated.get_json()
        self.assertEqual(job["resume_generation_status"], "generated")
        self.assertEqual(job["cover_letter_generation_status"], "generated")
        self.assertIn("generation_debug", job)
        debug = job["generation_debug"]
        self.assertTrue(debug.get("generated_at"))
        self.assertGreater(debug.get("approved_evidence_count", 0), 0)
        self.assertTrue(debug.get("approved_evidence_ids"))
        self.assertTrue(debug.get("selected_approved_evidence_ids"))
        self.assertTrue(set(debug.get("selected_approved_evidence_ids", [])).issubset(set(debug.get("approved_evidence_ids", []))))
        self.assertTrue(debug.get("matched_evidence_ids"))
        self.assertIn("ollama", debug)
        self.assertIn("called", debug["ollama"])

        resume_path = Path(job["resume_docx_path"])
        cover_path = Path(job["cover_letter_docx_path"])
        self.assertTrue(resume_path.exists())
        self.assertTrue(cover_path.exists())
        generation_debug_path = resume_path.parent / "generation_debug.json"
        self.assertTrue(generation_debug_path.exists())
        self.assertEqual(json.loads(generation_debug_path.read_text(encoding="utf-8"))["generated_at"], debug["generated_at"])

        resume_doc = Document(io.BytesIO(resume_path.read_bytes()))
        cover_doc = Document(io.BytesIO(cover_path.read_bytes()))
        resume_text = "\n".join(paragraph.text for paragraph in resume_doc.paragraphs)
        cover_text = "\n".join(paragraph.text for paragraph in cover_doc.paragraphs)
        cover_paragraphs = [paragraph.text.strip() for paragraph in cover_doc.paragraphs if paragraph.text.strip()]
        cover_body_start = next(index for index, text in enumerate(cover_paragraphs) if text.startswith("I am pleased"))
        cover_body_end = next(index for index, text in enumerate(cover_paragraphs) if text.startswith("Thank you"))
        cover_body = cover_paragraphs[cover_body_start:cover_body_end]

        original_resume_1 = "Built an automated ASIN monitoring system that tracked pricing, rank, and rating trends across 50-250 SKUs per session, delivering hourly insights over custom time frames and storing 300,000+ records in PostgreSQL."
        original_resume_2 = "Implemented robust validation, logging, and data cleaning logic to detect malformed inputs and reduce manual data entry errors by 90%."
        rewritten_resume_1 = "Built a backend ASIN monitoring system that automated pricing, rank, and rating tracking across 50-250 SKUs per session, delivered hourly insights over custom time frames, and stored 300,000+ records reliably in PostgreSQL."
        rewritten_resume_2 = "Strengthened software reliability by implementing validation, logging, and data-cleaning logic that detected malformed inputs and reduced manual data-entry errors by 90%."
        bad_cover_paragraph = "Python, JavaScript, Flask, React, SQL, cloud, AWS, Azure, dashboards, automation, APIs, reporting. This is the kind of work I would bring to Scotiabank."

        self.assertIn(original_resume_1, resume_text)
        self.assertIn(original_resume_2, resume_text)
        self.assertNotIn(rewritten_resume_1, resume_text)
        self.assertNotIn(rewritten_resume_2, resume_text)
        self.assertNotIn("testing", resume_text.lower())
        self.assertNotIn(bad_cover_paragraph, cover_text)
        self.assertEqual(len(cover_body), 3)
        self.assertGreaterEqual(len(" ".join(cover_body).split()), 50)
        self.assertLessEqual(len(" ".join(cover_body).split()), 260)
        self.assertIn("automated ASIN monitoring system", cover_text)
        self.assertIn("PostgreSQL", cover_text)
        self.assertIn("300,000+ records", cover_text)
        self.assertIn("validation, logging, and data cleaning", cover_text)
        self.assertIsNone(app_module.validate_cover_letter_document(cover_doc, self.scotiabank_job(), {"name": "Terry Mak"}))

        download_resume = self.client.get(f"/api/shortlist/{job_id}/resume")
        download_cover = self.client.get(f"/api/shortlist/{job_id}/cover-letter")
        self.assertEqual(download_resume.status_code, 200)
        self.assertEqual(download_cover.status_code, 200)
        self.assertEqual(download_resume.data, resume_path.read_bytes())
        self.assertEqual(download_cover.data, cover_path.read_bytes())

    def test_duplicate_shortlist_prevention_and_generation_failure_preserves_job(self):
        missing_profile = self.client.post(
            "/api/shortlist",
            json={
                "company": "SparseCo",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/sparse",
                "job_description": "Python SQL",
            },
        )
        self.assertEqual(missing_profile.status_code, 201)
        first = missing_profile.get_json()
        self.assertEqual(first["resume_generation_status"], "failed")
        self.assertIn("Complete the Profile page", first["generation_error"])

        duplicate = self.client.post(
            "/api/shortlist",
            json={
                "company": "SparseCo",
                "title": "Software Developer",
                "location": "Toronto",
                "url": "https://example.com/sparse?utm_source=x",
                "job_description": "Python SQL",
            },
        )
        self.assertEqual(duplicate.status_code, 200)
        self.assertTrue(duplicate.get_json()["duplicate"])
        self.assertEqual(len(self.client.get("/api/shortlist").get_json()), 1)

    def test_candidate_profile_uploads_are_safe_and_defaults_are_selectable(self):
        created = self.client.post(
            "/api/candidate-profiles",
            json={"name": "Managed Profile", "identity_mode": "distinct", "applicant_name": "QA Applicant"},
        )
        self.assertEqual(created.status_code, 201)
        profile_id = created.get_json()["id"]
        document = Document()
        document.add_paragraph("QA Applicant")
        document.add_paragraph("Documented Python and SQL experience.")
        buffer = io.BytesIO()
        document.save(buffer)
        first = self.client.post(
            f"/api/candidate-profiles/{profile_id}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "first resume.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        )
        second = self.client.post(
            f"/api/candidate-profiles/{profile_id}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "second resume.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        )
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 201)
        documents = self.client.get(f"/api/candidate-profiles/{profile_id}/documents").get_json()
        self.assertEqual(sum(item["is_default"] for item in documents["resumes"]), 1)
        self.assertTrue(Path(first.get_json()["path"]).exists())
        self.assertTrue(Path(second.get_json()["path"]).exists())
        selected = next(item for item in documents["resumes"] if item["is_default"])
        self.assertEqual(selected["id"], second.get_json()["id"])
        response = self.client.post(f"/api/candidate-profiles/{profile_id}/resumes/{first.get_json()['id']}/default")
        self.assertEqual(response.status_code, 200)
        documents = self.client.get(f"/api/candidate-profiles/{profile_id}/documents").get_json()
        self.assertEqual([item["id"] for item in documents["resumes"] if item["is_default"]], [first.get_json()["id"]])

    def test_generation_response_exposes_actual_status_for_frontend_feedback(self):
        self.seed_profile_for_documents(with_template=True)
        created = self.client.post("/api/shortlist", json={"company": "Feedback Co", "title": "Developer", "job_description": "Python SQL"}).get_json()
        self.assertEqual(created["generation"]["resume"], "generated")
        self.assertEqual(created["generation"]["cover_letter"], "generated")
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn("function showGenerationOutcome", frontend)
        self.assertIn("previous successful document remains available", frontend)

    def test_regeneration_after_profile_update(self):
        created = self.client.post(
            "/api/shortlist",
            json={
                "company": "LaterCo",
                "title": "Cloud Support Analyst",
                "location": "Remote Canada",
                "url": "https://example.com/later",
                "job_description": "Azure SQL APIs support automation dashboards",
            },
        ).get_json()
        self.assertEqual(created["resume_generation_status"], "failed")
        self.seed_profile_for_documents(with_template=True)
        regenerated = self.client.post(f"/api/shortlist/{created['id']}/regenerate", json={})
        self.assertEqual(regenerated.status_code, 200)
        job = regenerated.get_json()
        self.assertEqual(job["resume_generation_status"], "generated")
        self.assertTrue(Path(job["resume_docx_path"]).exists())

    def test_application_preparation_queue_is_persistent_and_does_not_change_job_status(self):
        profile = self.client.post(
            "/api/candidate-profiles",
            json={"name": "Phase B QA", "identity_mode": "distinct", "applicant_name": "QA Applicant"},
        ).get_json()
        document = Document()
        document.add_paragraph("QA Applicant")
        document.add_paragraph("Python and SQL experience")
        buffer = io.BytesIO()
        document.save(buffer)
        uploaded = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "qa.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        )
        self.assertEqual(uploaded.status_code, 201)
        cover_document = Document()
        cover_document.add_paragraph("QA Applicant")
        cover_document.add_paragraph("I am applying to QueueCo for the Developer role.")
        cover_buffer = io.BytesIO()
        cover_document.save(cover_buffer)
        cover_uploaded = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/cover-letter-templates/upload",
            data={"file": (io.BytesIO(cover_buffer.getvalue()), "qa-cover.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        )
        self.assertEqual(cover_uploaded.status_code, 201)
        job = self.client.post("/api/jobs", json={"company": "QueueCo", "title": "Developer", "location": "Toronto", "overall_score": 88, "remote_eligibility": "canadian_eligible", "job_description": "Python SQL"}).get_json()
        queued = self.client.post("/api/preparation/queue", json={"job_id": job["id"], "candidate_profile_id": profile["id"]})
        self.assertEqual(queued.status_code, 201, queued.get_json())
        item = queued.get_json()
        duplicate = self.client.post("/api/preparation/queue", json={"job_id": job["id"], "candidate_profile_id": profile["id"]})
        self.assertEqual(duplicate.status_code, 200)
        self.assertTrue(duplicate.get_json()["duplicate"])
        self.assertEqual(len(self.client.get("/api/preparation/queue").get_json()), 1)
        self.assertEqual(self.client.get("/api/jobs").get_json()[0]["status"], "Saved")
        generated_folder = app_module.generated_applications_dir() / str(job["id"])
        generated_folder.mkdir(parents=True, exist_ok=True)
        generated_resume = generated_folder / "queue-resume.docx"
        generated_cover = generated_folder / "queue-cover.docx"
        generated_resume.write_bytes(buffer.getvalue())
        generated_cover.write_bytes(cover_buffer.getvalue())

        def fake_prepare(job_id, kinds, profile):
            with app_module.app.app_context():
                app_module.get_db().execute(
                    "UPDATE jobs SET resume_generation_status = 'generated', cover_letter_generation_status = 'generated', resume_docx_path = ?, cover_letter_docx_path = ? WHERE id = ?",
                    [str(generated_resume), str(generated_cover), job_id],
                )
                app_module.get_db().commit()
            return {"resume": "generated", "cover_letter": "generated", "error": ""}

        with patch.object(app_module, "generate_and_store_shortlist_documents", side_effect=fake_prepare):
            prepared = self.client.post(f"/api/preparation/{item['id']}/prepare")
        self.assertEqual(prepared.status_code, 200, prepared.get_json())
        self.assertEqual(prepared.get_json()["status"], "ready_for_review")
        approved = self.client.post(f"/api/preparation/{item['id']}/review", json={"action": "approve"})
        self.assertEqual(approved.status_code, 200)
        self.assertEqual(approved.get_json()["status"], "approved")
        self.assertEqual(self.client.get(f"/api/preparation/{item['id']}/preview").status_code, 200)
        self.assertEqual(self.client.get(f"/api/preparation/{item['id']}/download/resume").status_code, 200)
        self.assertEqual(self.client.get(f"/api/preparation/{item['id']}/download/cover-letter").status_code, 200)
        self.assertEqual(self.client.post(f"/api/preparation/{item['id']}/regenerate").status_code, 409)
        self.assertEqual(self.client.get("/api/jobs").get_json()[0]["status"], "Saved")

    def test_application_preparation_auto_eligibility_requires_priority_and_confirmed_canada(self):
        search = {"auto_prepare_enabled": True, "auto_prepare_min_priority": 80, "auto_prepare_require_canadian": True}
        eligible, reason, _ = app_module.preparation_eligibility({"overall_score": 90, "remote_eligibility": "canadian_eligible"}, search)
        self.assertTrue(eligible)
        self.assertEqual(reason, "eligible")
        eligible, reason, _ = app_module.preparation_eligibility({"overall_score": 90, "remote_eligibility": "unknown"}, search)
        self.assertFalse(eligible)
        self.assertEqual(reason, "manual_eligibility_review")
        eligible, reason, _ = app_module.preparation_eligibility({"overall_score": 90, "remote_eligibility": "restricted"}, search)
        self.assertFalse(eligible)
        self.assertEqual(reason, "explicitly_restricted")
        eligible, reason, _ = app_module.preparation_eligibility({"overall_score": 70, "remote_eligibility": "canadian_eligible"}, search)
        self.assertFalse(eligible)
        self.assertEqual(reason, "priority_below_threshold")
        for require_canadian in (True, False):
            eligible, reason, _ = app_module.preparation_eligibility(
                {"overall_score": 90, "remote_eligibility": "unknown"},
                {"auto_prepare_min_priority": 80, "auto_prepare_require_canadian": require_canadian},
            )
            self.assertFalse(eligible)
            self.assertEqual(reason, "manual_eligibility_review")

    def test_automatic_preparation_never_chooses_first_profile_implicitly(self):
        with self.assertRaisesRegex(ValueError, "automatic preparation cannot choose one implicitly"):
            app_module.queue_profile_payload(search={"auto_prepare_enabled": True})

    def test_preparation_approval_requires_ready_generated_package(self):
        job = self.client.post("/api/jobs", json={"company": "Approval Co", "title": "Developer"}).get_json()
        with app_module.app.app_context():
            app_module.get_db().execute("INSERT INTO application_preparation_queue (job_id, status) VALUES (?, 'queued')", [job["id"]])
            app_module.get_db().commit()
            item_id = app_module.get_db().execute("SELECT last_insert_rowid() AS id").fetchone()["id"]
        response = self.client.post(f"/api/preparation/{item_id}/review", json={"action": "approve"})
        self.assertEqual(response.status_code, 409)
        self.assertIn("ready_for_review", response.get_json()["error"])

    def test_preparation_preview_uses_only_documents_belonging_to_package(self):
        job = self.client.post("/api/jobs", json={"company": "Preview Co", "title": "Developer", "status": "shortlisted"}).get_json()
        package_dir = app_module.generated_applications_dir() / str(job["id"])
        package_dir.mkdir(parents=True, exist_ok=True)

        def docx_bytes(text):
            document = Document()
            document.add_paragraph(text)
            output = io.BytesIO()
            document.save(output)
            return output.getvalue()

        current_resume = package_dir / "current-resume.docx"
        current_cover = package_dir / "current-cover.docx"
        package_resume = package_dir / "package-resume.docx"
        package_cover = package_dir / "package-cover.docx"
        current_resume.write_bytes(docx_bytes("CURRENT JOB DOCUMENT"))
        current_cover.write_bytes(docx_bytes("CURRENT JOB COVER LETTER"))
        package_resume.write_bytes(docx_bytes("PACKAGE RESUME"))
        package_cover.write_bytes(docx_bytes("PACKAGE COVER LETTER"))
        with app_module.app.app_context():
            app_module.get_db().execute(
                """INSERT INTO application_preparation_queue
                   (job_id, status, resume_status, cover_letter_status, prepared_resume_path, prepared_cover_letter_path)
                   VALUES (?, 'ready_for_review', 'generated', 'generated', ?, ?)""",
                [job["id"], str(package_resume), str(package_cover)],
            )
            item_id = app_module.get_db().execute("SELECT last_insert_rowid() AS id").fetchone()["id"]
            app_module.get_db().execute(
                "UPDATE jobs SET resume_docx_path = ?, cover_letter_docx_path = ?, status = 'shortlisted' WHERE id = ?",
                [str(current_resume), str(current_cover), job["id"]],
            )
            app_module.get_db().commit()

        response = self.client.get(f"/api/preparation/{item_id}/preview")
        self.assertEqual(response.status_code, 200)
        preview = response.get_json()
        self.assertEqual(preview["resume"], "PACKAGE RESUME")
        self.assertEqual(preview["cover_letter"], "PACKAGE COVER LETTER")
        self.assertEqual(preview["errors"], {})
        self.assertEqual(preview["company"], "Preview Co")
        self.assertEqual(preview["title"], "Developer")
        self.assertNotIn("CURRENT JOB", preview["resume"] + preview["cover_letter"])

    def test_preparation_preview_reports_missing_package_artifacts_explicitly(self):
        job = self.client.post("/api/jobs", json={"company": "Missing Preview Co", "title": "Developer"}).get_json()
        with app_module.app.app_context():
            item_id = app_module.get_db().execute(
                "INSERT INTO application_preparation_queue (job_id, status, prepared_resume_path, prepared_cover_letter_path) VALUES (?, 'ready_for_review', ?, ?)",
                [job["id"], str(app_module.generated_applications_dir() / "missing-resume.docx"), ""],
            ).lastrowid
            app_module.get_db().commit()
        response = self.client.get(f"/api/preparation/{item_id}/preview")
        self.assertEqual(response.status_code, 200)
        errors = response.get_json()["errors"]
        self.assertIn("missing from this application environment", errors["resume"])
        self.assertIn("not recorded", errors["cover_letter"])

    def test_preparation_preview_exports_exact_package_docs_and_caches_by_content_hash(self):
        job = self.client.post("/api/jobs", json={"company": "PDF Preview Co", "title": "Developer"}).get_json()
        package_dir = app_module.generated_applications_dir() / str(job["id"])
        package_dir.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_paragraph("Package layout preview")
        content = io.BytesIO()
        document.save(content)
        resume = package_dir / "resume.docx"
        cover = package_dir / "cover.docx"
        resume.write_bytes(content.getvalue())
        cover.write_bytes(content.getvalue())
        with app_module.app.app_context():
            item_id = app_module.get_db().execute(
                "INSERT INTO application_preparation_queue (job_id, status, prepared_resume_path, prepared_cover_letter_path) VALUES (?, 'ready_for_review', ?, ?)",
                [job["id"], str(resume), str(cover)],
            ).lastrowid
            app_module.get_db().commit()

        def fake_export(source, output):
            Path(output).parent.mkdir(parents=True, exist_ok=True)
            Path(output).write_bytes(b"%PDF-1.4\nmock preview")
            return {"available": True, "path": str(output)}

        with patch.object(app_module, "export_docx_to_pdf", side_effect=fake_export) as export:
            first = self.client.get(f"/api/preparation/{item_id}/preview")
            self.assertEqual(first.status_code, 200)
            body = first.get_json()
            self.assertTrue(body["documents"]["resume"]["available"])
            self.assertTrue(body["documents"]["cover_letter"]["available"])
            resume_pdf = self.client.get(body["documents"]["resume"]["pdf_url"])
            self.assertEqual(resume_pdf.status_code, 200)
            self.assertEqual(resume_pdf.mimetype, "application/pdf")
            second = self.client.get(f"/api/preparation/{item_id}/preview")
            self.assertEqual(second.status_code, 200)
            self.assertEqual(export.call_count, 2)
            resume.write_bytes(content.getvalue() + b"changed")
            changed = self.client.get(f"/api/preparation/{item_id}/preview").get_json()
            self.assertNotEqual(changed["documents"]["resume"]["content_hash"], body["documents"]["resume"]["content_hash"])
            self.assertEqual(export.call_count, 3)
            stale = self.client.get(body["documents"]["resume"]["pdf_url"])
            self.assertEqual(stale.status_code, 409)

    def test_preparation_preview_is_inline_and_tailoring_approval_uses_compact_diff(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        template = Path(app_module.BASE_DIR, "templates", "index.html").read_text(encoding="utf-8")
        self.assertIn("preparationPreviews", frontend)
        self.assertIn("preparation-preview-inline", frontend)
        self.assertNotIn("preparation-document-preview", template)
        self.assertIn("preparation-preview-pdf", frontend)
        self.assertIn("Open PDF separately", frontend)
        self.assertIn("tailoringDiffMarkup", frontend)
        self.assertIn("tailoring-compact-diff", frontend)
        self.assertNotIn('insertAdjacentHTML("afterbegin", `<details class="preparation-preview"', frontend)
        self.assertIn('id="tailoring-diff-heading"', template)

    def test_preparation_preview_retry_discards_stale_responses_and_live_diff_remains_for_new_edits(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn("async function previewPreparation(id, retry = false)", frontend)
        self.assertIn("previewPreparation(${item.id}, true)", frontend)
        self.assertIn("preparationPreviewEpoch", frontend)
        self.assertIn("preparationPreviewRequests", frontend)
        self.assertIn("state.preparationPreviewEpoch !== previewEpoch", frontend)
        self.assertIn("state.preparationPreviewRequests[id] !== requestId", frontend)
        self.assertIn("const hideStandaloneDiff = Boolean(changes.length && hasAppliedChange && !hasMeaningfulChange)", frontend)

    def test_manual_discovery_preparation_requires_explicit_profile_and_docx_pair(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        template = Path(app_module.BASE_DIR, "templates", "index.html").read_text(encoding="utf-8")
        self.assertIn("Choose the candidate profile for this package by number", frontend)
        self.assertIn("must have both a DOCX résumé and a DOCX cover letter template", frontend)
        self.assertIn("Unknown Canadian eligibility always requires manual confirmation", template)
        self.assertNotIn("saved-search-auto-require-canadian", template)

    def test_pipeline_regeneration_feedback_is_visible_and_deduplicated(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn("Preparing documents", frontend)
        self.assertIn("Last successful generation", frontend)
        self.assertIn("New document version ready", frontend)
        self.assertIn("data-view-application-package", frontend)
        self.assertIn("prepared_at", frontend)

    def test_preparation_queue_regeneration_has_persistent_busy_success_and_error_states(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn("preparationRegenerations", frontend)
        self.assertIn("/api/preparation/${id}/regenerate", frontend)
        self.assertIn("Regenerating...", frontend)
        self.assertIn("Generating résumé and cover letter. Please wait.", frontend)
        self.assertIn("New document version ready", frontend)
        self.assertIn("Retry when the issue is resolved.", frontend)
        self.assertIn("preparedAt", frontend)

    def test_regeneration_does_not_reset_an_active_preparation_lease(self):
        job = self.client.post("/api/jobs", json={"company": "Lease Co", "title": "Developer"}).get_json()
        with app_module.app.app_context():
            app_module.get_db().execute(
                "INSERT INTO application_preparation_queue (job_id, status, lease_owner, lease_expires_at) VALUES (?, 'preparing', 'other-worker', ?)",
                [job["id"], (datetime.now() + timedelta(minutes=10)).isoformat(timespec="seconds")],
            )
            app_module.get_db().commit()
            item_id = app_module.get_db().execute("SELECT last_insert_rowid() AS id").fetchone()["id"]
        response = self.client.post(f"/api/preparation/{item_id}/regenerate")
        self.assertEqual(response.status_code, 409)
        with app_module.app.app_context():
            row = app_module.query_db("SELECT status, lease_owner FROM application_preparation_queue WHERE id = ?", [item_id], one=True)
        self.assertEqual(row["status"], "preparing")
        self.assertEqual(row["lease_owner"], "other-worker")

    def test_unknown_eligibility_requires_explicit_review_before_preparation(self):
        profile = self.client.post("/api/candidate-profiles", json={"name": "Eligibility QA", "identity_mode": "distinct", "applicant_name": "QA Applicant"}).get_json()
        resume = Document()
        resume.add_paragraph("QA Applicant")
        resume_bytes = io.BytesIO()
        resume.save(resume_bytes)
        self.client.post(f"/api/candidate-profiles/{profile['id']}/resumes/upload", data={"file": (io.BytesIO(resume_bytes.getvalue()), "resume.docx")}, content_type="multipart/form-data")
        cover = Document()
        cover.add_paragraph("QA Applicant")
        cover_bytes = io.BytesIO()
        cover.save(cover_bytes)
        self.client.post(f"/api/candidate-profiles/{profile['id']}/cover-letter-templates/upload", data={"file": (io.BytesIO(cover_bytes.getvalue()), "cover.docx")}, content_type="multipart/form-data")
        job = self.client.post("/api/jobs", json={"company": "Unknown Co", "title": "Developer", "overall_score": 90, "remote_eligibility": "unknown"}).get_json()
        queued = self.client.post("/api/preparation/queue", json={"job_id": job["id"], "candidate_profile_id": profile["id"]})
        self.assertEqual(queued.status_code, 201, queued.get_json())
        item = queued.get_json()
        self.assertEqual(item["review_status"], "manual_eligibility")
        blocked = self.client.post(f"/api/preparation/{item['id']}/prepare")
        self.assertEqual(blocked.status_code, 409)
        reviewed = self.client.post(f"/api/preparation/{item['id']}/eligibility-review", json={"confirmed": True})
        self.assertEqual(reviewed.status_code, 200)
        self.assertEqual(reviewed.get_json()["eligibility_review_status"], "confirmed")

    def test_discovery_search_config_keeps_total_cap_authoritative(self):
        config = app_module.discovery_search_config({
            "search_depth": "standard",
            "enabled_sources": ["linkedin", "indeed"],
            "per_source_cap": 100,
            "total_cap": 20,
        })
        self.assertEqual(config["total_cap"], 20)
        self.assertEqual(config["per_source_cap"], 100)
        self.assertEqual(config["request_per_source_cap"], 10)
        self.assertEqual(config["aggregate_request_budget"], 20)
        quick = app_module.discovery_search_config({"search_depth": "quick", "enabled_sources": ["linkedin", "indeed"]})
        self.assertEqual(quick["total_cap"], 50)
        self.assertEqual(quick["request_per_source_cap"], 25)
        three = app_module.discovery_search_config({
            "enabled_sources": ["linkedin", "indeed", "google"],
            "per_source_cap": 100,
            "total_cap": 20,
        })
        self.assertEqual(three["request_per_source_cap"], 6)
        self.assertLessEqual(three["aggregate_request_budget"], 20)
        minimum = app_module.discovery_search_config({
            "enabled_sources": ["linkedin", "indeed", "google"],
            "per_source_cap": 100,
            "total_cap": 3,
        })
        self.assertEqual(minimum["request_per_source_cap"], 1)
        with self.assertRaisesRegex(ValueError, "at least 3"):
            app_module.discovery_search_config({
                "enabled_sources": ["linkedin", "indeed", "google"],
                "total_cap": 2,
            })

    def test_manual_search_rejects_aggregate_limit_below_selected_source_count(self):
        response = self.client.post("/api/discovery/jobspy", json={
            "keywords": "backend", "enabled_sources": ["linkedin", "indeed", "google"], "total_cap": 2,
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn("at least 3", response.get_json()["error"])

    def test_saved_search_preserves_current_limits_and_is_editable_without_running(self):
        created = self.client.post("/api/discovery/searches", json={
            "name": "Saved current search",
            "keywords": "backend developer",
            "location": "Toronto",
            "remote_preference": "hybrid",
            "distance": 25,
            "enabled_sources": ["linkedin", "indeed", "google"],
            "date_window": "7d",
            "search_depth": "thorough",
            "per_source_cap": 100,
            "total_cap": 20,
        })
        self.assertEqual(created.status_code, 201)
        search = created.get_json()
        self.assertEqual(search["distance"], 25)
        self.assertEqual(search["total_cap"], 20)
        self.assertEqual(search["per_source_cap"], 100)
        self.assertEqual(search["search_depth"], "thorough")
        self.assertFalse(search["enabled"])
        updated = self.client.put(f"/api/discovery/searches/{search['id']}", json={**search, "name": "Edited current search", "total_cap": 17, "distance": 40})
        self.assertEqual(updated.status_code, 200)
        refreshed = self.client.get("/api/discovery/searches").get_json()
        saved = next(item for item in refreshed if item["id"] == search["id"])
        self.assertEqual(saved["name"], "Edited current search")
        self.assertEqual(saved["total_cap"], 17)
        self.assertEqual(saved["distance"], 40)

    def test_saved_search_create_and_edit_reject_invalid_aggregate_budget(self):
        invalid = self.client.post("/api/discovery/searches", json={
            "name": "Invalid budget", "keywords": "backend", "enabled_sources": ["linkedin", "indeed", "google"], "total_cap": 2,
        })
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("at least 3", invalid.get_json()["error"])
        valid = self.client.post("/api/discovery/searches", json={
            "name": "Editable budget", "keywords": "backend", "enabled_sources": ["linkedin", "indeed", "google"], "total_cap": 3,
        }).get_json()
        invalid_update = self.client.put(f"/api/discovery/searches/{valid['id']}", json={**valid, "total_cap": 2})
        self.assertEqual(invalid_update.status_code, 400)
        self.assertIn("at least 3", invalid_update.get_json()["error"])
        preserved = self.client.get("/api/discovery/searches").get_json()
        stored = next(item for item in preserved if item["id"] == valid["id"])
        self.assertEqual(stored["total_cap"], 3)

    def test_saved_search_run_uses_persisted_fetch_limits_and_reports_snapshot_separately(self):
        search = self.client.post("/api/discovery/searches", json={
            "name": "Limited saved run", "keywords": "backend", "enabled_sources": ["linkedin", "indeed", "google"],
            "total_cap": 20, "per_source_cap": 100,
        }).get_json()
        fake = {"jobs": [], "errors": {}, "source_counts": {}, "source_diagnostics": {}, "source_snapshots": {}}
        location = {"original_input": "", "normalized_input": "", "distance_origin": "", "discovery_location": "", "resolution_method": "test", "lookup": {}}
        with patch.object(app_module, "get_geocoding_provider", return_value=FakeGeocodingProvider({})), patch.object(app_module, "resolve_discovery_location_context", return_value=location), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake) as mocked_fetch:
            response = self.client.post(f"/api/discovery/searches/{search['id']}/run")
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(mocked_fetch.call_args.kwargs["results_wanted"], 6)
        body = response.get_json()
        self.assertEqual(body["search"]["total_cap"], 20)
        self.assertIn("persisted", body)
        self.assertIn("visible_total", body)

    def test_saved_search_run_rejects_invalid_aggregate_budget_before_provider_call(self):
        self.client.get("/api/discovery/searches")
        with app_module.app.app_context():
            cursor = app_module.get_db().execute(
                "INSERT INTO saved_search_profiles (name, keywords, enabled_sources_json, total_cap) VALUES (?, ?, ?, ?)",
                ["Invalid legacy budget", "backend", '["linkedin", "indeed", "google"]', 2],
            )
            app_module.get_db().commit()
            search_id = cursor.lastrowid
        response = self.client.post(f"/api/discovery/searches/{search_id}/run")
        self.assertEqual(response.status_code, 400)
        self.assertIn("at least 3", response.get_json()["error"])

    def test_worker_records_invalid_budget_without_requesting_providers(self):
        from scripts.discovery_worker import run_due_searches

        search = self.client.post("/api/discovery/searches", json={
            "name": "Worker invalid budget", "keywords": "backend", "enabled_sources": ["linkedin", "indeed", "google"], "total_cap": 3, "enabled": True,
        }).get_json()
        with app_module.app.app_context():
            app_module.get_db().execute("UPDATE saved_search_profiles SET total_cap = 2, next_run_at = '' WHERE id = ?", [search["id"]])
            app_module.get_db().execute("INSERT INTO app_settings (key, value) VALUES ('discovery_scheduler_enabled', '1') ON CONFLICT(key) DO UPDATE SET value = '1'")
            app_module.get_db().commit()
            with patch.object(app_module, "fetch_jobspy_jobs") as mocked_fetch:
                result = run_due_searches()
        self.assertEqual(result["runs"], 0)
        self.assertEqual(result["errors"], 1)
        mocked_fetch.assert_not_called()
        run = self.client.get("/api/discovery/runs").get_json()[0]
        self.assertEqual(run["status"], "failed")
        self.assertIn("at least 3", run["error"])

    def test_legacy_saved_search_columns_migrate_without_losing_existing_search(self):
        self.client.get("/api/discovery/searches")
        with app_module.app.app_context():
            cursor = app_module.get_db().execute(
                "INSERT INTO saved_search_profiles (name, keywords) VALUES (?, ?)",
                ["Legacy preserved search", "software engineer"],
            )
            app_module.get_db().commit()
            legacy_id = cursor.lastrowid
        response = self.client.get("/api/discovery/searches")
        self.assertEqual(response.status_code, 200)
        legacy = next(item for item in response.get_json() if item["id"] == legacy_id)
        self.assertEqual(legacy["name"], "Legacy preserved search")
        self.assertEqual(legacy["total_cap"], 500)
        self.assertEqual(legacy["per_source_cap"], 100)
        self.assertEqual(legacy["search_depth"], "standard")

    def test_discovery_total_cap_bounds_two_source_fetch_and_reports_separate_counts(self):
        fake_jobs = [
            {
                "source_job_id": f"job-{index}",
                "title": f"Software Engineer {index}",
                "company": f"Company {index}",
                "location": "Toronto, ON",
                "source": "linkedin" if index % 2 else "indeed",
                "job_url": f"https://example.com/jobs/{index}",
                "description": "Python SQL backend",
            }
            for index in range(40)
        ]
        fake = {"jobs": fake_jobs, "errors": {}, "source_counts": {"linkedin": 20, "indeed": 20}, "source_diagnostics": {}}
        location = {"original_input": "Toronto", "normalized_input": "Toronto", "distance_origin": "Toronto", "discovery_location": "Toronto, ON", "resolution_method": "test", "lookup": {}}
        with patch.object(app_module, "get_geocoding_provider", return_value=FakeGeocodingProvider({})), patch.object(app_module, "resolve_discovery_location_context", return_value=location), patch.object(app_module, "fetch_jobspy_jobs", return_value=fake) as mocked_fetch:
            response = self.client.post("/api/discovery/jobspy", json={
                "keywords": "software engineer", "location": "Toronto", "enabled_sources": ["linkedin", "indeed"],
                "per_source_cap": 100, "total_cap": 20, "bypass_cache": True,
            })
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(mocked_fetch.call_args.kwargs["results_wanted"], 10)
        self.assertEqual(body["search"]["total_cap"], 20)
        self.assertEqual(body["total"], 40)
        self.assertEqual(body["visible_total"], 20)
        self.assertEqual(body["diagnostics"]["duplicates_in_search"], 0)
        self.assertEqual(body["diagnostics"]["already_saved_matches"], 0)
        self.assertEqual(body["diagnostics"]["excluded_count"], 0)
        self.assertEqual(body["diagnostics"]["capped_count"], 20)

    def test_discovery_views_keep_search_results_and_saved_searches_separate(self):
        template = Path(app_module.BASE_DIR, "templates", "index.html").read_text(encoding="utf-8")
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn('data-discovery-view="pipeline"', template)
        self.assertIn('data-discovery-view="find"', template)
        self.assertIn('data-discovery-view="saved"', template)
        self.assertIn('class="panel discovery-results-panel discovery-view-find"', template)
        self.assertIn('class="panel discovery-saved-searches discovery-view-saved"', template)
        self.assertIn("setDiscoverySection", frontend)
        self.assertIn('id="save-current-discovery-search"', template)
        self.assertIn("saveCurrentDiscoverySearch", frontend)
        self.assertIn("Search snapshot:", frontend)
        self.assertIn("discovery-counts", frontend)
        self.assertIn("configured home location", frontend)


    def _make_ai_fixture(self):
        self.client.post("/api/profile", json={"name": "Terry Mak", "email": "terry@example.com", "resume_text": "Flowtally"})
        profile = self.client.post("/api/candidate-profiles", json={"name": "Software Engineering", "identity_mode": "shared"}).get_json()
        document = Document()
        original = "Built 25+ tests and reduced errors by 75% across 3 services."
        document.add_paragraph(original, style="List Bullet")
        buffer = io.BytesIO()
        document.save(buffer)
        resume = self.client.post(
            f"/api/candidate-profiles/{profile['id']}/resumes/upload",
            data={"file": (io.BytesIO(buffer.getvalue()), "ai-fixture.docx"), "is_default": "true"},
            content_type="multipart/form-data",
        ).get_json()
        job = self.client.post("/api/jobs", json={"company": "Example", "title": "Developer", "job_description": "Python and REST APIs"}).get_json()
        workspace = self.client.get(f"/api/tailoring/{job['id']}/workspace?candidate_profile_id={profile['id']}&resume_id={resume['id']}").get_json()
        config = app_module.OllamaConfig(True, "http://ollama", "qwen3:4b", 10, 100, 1000)
        return workspace, original, config

    def test_tailoring_ai_timeout_is_unavailable_without_automatic_retry(self):
        workspace, original, config = self._make_ai_fixture()
        with patch.object(app_module, "request_json_completion", return_value={"ok": False, "error": "timed out", "data": {}, "attempts": 1}) as request_mock:
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                started = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original})
                result = self.wait_for_tailoring_ai(started)
        self.assertEqual(result.status_code, 503)
        self.assertEqual(result.get_json()["status"], "unavailable")
        self.assertEqual(result.get_json()["proposed_text"], original)
        self.assertEqual(request_mock.call_count, 1)
        self.assertEqual(request_mock.call_args.kwargs["max_attempts"], 1)
        self.assertEqual(request_mock.call_args.kwargs["keep_alive"], 0)

    def test_tailoring_ai_cancellation_preserves_original(self):
        workspace, original, config = self._make_ai_fixture()

        def slow_completion(*args, **kwargs):
            event = kwargs["cancel_event"]
            while not event.is_set():
                time.sleep(0.01)
            return {"ok": False, "cancelled": True, "error": "AI improvement was cancelled.", "data": {}, "attempts": 1}

        with patch.object(app_module, "request_json_completion", side_effect=slow_completion):
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                started = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original})
                job_id = started.get_json()["job_id"]
                cancel = self.client.post(f"/api/tailoring/improve-jobs/{job_id}/cancel")
                result = self.client.get(f"/api/tailoring/improve-jobs/{job_id}")
                for _ in range(100):
                    if result.status_code != 202: break
                    time.sleep(0.01)
                    result = self.client.get(f"/api/tailoring/improve-jobs/{job_id}")
        self.assertEqual(cancel.status_code, 202)
        self.assertEqual(result.status_code, 503)
        self.assertEqual(result.get_json()["status"], "cancelled")
        self.assertEqual(result.get_json()["proposed_text"], original)

    def test_tailoring_ai_rejects_unsafe_retry_and_preserves_original(self):
        workspace, original, config = self._make_ai_fixture()
        unsafe = "Built 20+ tests and reduced errors by 75% across 3 services."
        with patch.object(
            app_module,
            "request_json_completion",
            side_effect=[
                {"ok": True, "data": {"proposed_bullet": unsafe}, "attempts": 1},
                {"ok": True, "data": {"proposed_bullet": unsafe}, "attempts": 1},
            ],
        ) as request_mock:
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                started = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original})
                result = self.wait_for_tailoring_ai(started)
        self.assertEqual(result.status_code, 422)
        self.assertEqual(result.get_json()["status"], "rejected")
        self.assertEqual(result.get_json()["proposed_text"], original)
        self.assertEqual(request_mock.call_count, 2)
        self.assertTrue(result.get_json()["retry_attempted"])

    def test_tailoring_ai_rejects_concurrent_request_and_recovers_after_cancellation(self):
        workspace, original, config = self._make_ai_fixture()

        def slow_completion(*args, **kwargs):
            while not kwargs["cancel_event"].is_set():
                time.sleep(0.01)
            return {"ok": False, "cancelled": True, "error": "AI improvement was cancelled.", "data": {}, "attempts": 1}

        with patch.object(app_module, "request_json_completion", side_effect=slow_completion):
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                first = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original})
                second = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original})
                job_id = first.get_json()["job_id"]
                self.client.post(f"/api/tailoring/improve-jobs/{job_id}/cancel")
                cancelled = self.client.get(f"/api/tailoring/improve-jobs/{job_id}")
                for _ in range(100):
                    if cancelled.status_code != 202: break
                    time.sleep(0.01)
                    cancelled = self.client.get(f"/api/tailoring/improve-jobs/{job_id}")
        self.assertEqual(second.status_code, 429)
        self.assertEqual(second.get_json()["status"], "busy")
        self.assertEqual(cancelled.get_json()["status"], "cancelled")

        with patch.object(app_module, "request_json_completion", return_value={"ok": True, "data": {"proposed_bullet": "Built and maintained 25+ automated tests, reducing errors by 75% across 3 services."}, "attempts": 1}):
            with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                recovered = self.wait_for_tailoring_ai(self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original}))
        self.assertEqual(recovered.status_code, 200)
        self.assertEqual(recovered.get_json()["status"], "completed")

    def test_tailoring_ai_completed_jobs_expire_from_memory(self):
        workspace, original, config = self._make_ai_fixture()
        old_retention = app_module.TAILORING_AI_JOB_RETENTION_SECONDS
        app_module.TAILORING_AI_JOB_RETENTION_SECONDS = 600
        try:
            with patch.object(app_module, "request_json_completion", return_value={"ok": True, "data": {"proposed_bullet": original}, "attempts": 1}):
                with patch.object(app_module.OllamaConfig, "from_mapping", return_value=config):
                    started = self.client.post(f"/api/tailoring/drafts/{workspace['draft']['id']}/improve", json={"original_text": original})
                    job_id = started.get_json()["job_id"]
                    result = self.wait_for_tailoring_ai(started)
            app_module.TAILORING_AI_JOB_RETENTION_SECONDS = 0
            app_module.cleanup_tailoring_ai_jobs()
            expired = self.client.get(f"/api/tailoring/improve-jobs/{job_id}")
            self.assertEqual(result.get_json()["status"], "unchanged")
            self.assertEqual(expired.status_code, 404)
        finally:
            app_module.TAILORING_AI_JOB_RETENTION_SECONDS = old_retention
    def test_ollama_deadline_is_checked_before_network_request(self):
        config = ollama_client.OllamaConfig(True, "http://127.0.0.1:11434", "qwen3:4b", 30, 128, 2000)
        with patch.object(ollama_client.urllib.request, "urlopen") as urlopen:
            result = ollama_client.request_json_completion(config, "Return JSON.", max_attempts=2, deadline=time.monotonic() - 1, keep_alive=0)
        self.assertTrue(result.get("deadline_exceeded"))
        urlopen.assert_not_called()

    def test_tailoring_ai_frontend_consumes_structured_terminal_statuses(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn("getOutcome", frontend)
        self.assertIn('result.status === "rejected"', frontend)
        self.assertIn('result.status === "unavailable"', frontend)
        self.assertIn("Original wording was preserved.", frontend)
        self.assertIn('result.status === "cancelled"', frontend)

    def test_tailoring_frontend_preserves_backend_evidence_ranking(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn("const existing = group.suggestions.get(bullet.paragraph_index)", frontend)
        self.assertIn("Number(b.relevance_score ?? 0) - Number(a.relevance_score ?? 0)", frontend)
        self.assertIn("Number(item.relevance_score ?? Number.NEGATIVE_INFINITY) > selectedRelevance", frontend)
        self.assertNotIn("const evidenceScore =", frontend)

    def test_application_workspace_updates_status_only_after_explicit_manual_action(self):
        job = self.client.post("/api/jobs", json={"company": "Workspace Co", "title": "Developer", "status": "Saved"}).get_json()
        with app_module.app.app_context():
            cursor = app_module.get_db().execute(
                "INSERT INTO application_preparation_queue (job_id, status, review_status) VALUES (?, 'approved', 'approved')",
                [job["id"]],
            )
            app_module.get_db().commit()
            item_id = cursor.lastrowid
        response = self.client.post(f"/api/preparation/{item_id}/application-status", json={"status": "Applied"})
        self.assertEqual(response.status_code, 200, response.get_json())
        refreshed = self.client.get("/api/jobs").get_json()[0]
        self.assertEqual(refreshed["status"], "Applied")
        self.assertTrue(refreshed["date_applied"])
        self.assertEqual(self.client.post(f"/api/preparation/{item_id}/application-status", json={"status": "Nope"}).status_code, 400)

    def test_application_workspace_has_master_detail_contextual_actions(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn("application-workspace-grid", frontend)
        self.assertIn("selectPreparationPackage", frontend)
        self.assertIn("application_id", frontend)
        self.assertIn("persistSelectedPreparation", frontend)
        self.assertIn("preparationNextAction", frontend)
        self.assertIn("/api/preparation/${id}/application-status", frontend)
        self.assertIn("Open employer application page", frontend)

    def test_application_workspace_next_action_uses_underlying_job_status(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        for status in ("applied", "interview", "offer", "rejected", "archived"):
            self.assertIn(f'applicationStatus === "{status}"', frontend)
        self.assertIn("Application recorded as submitted", frontend)
        self.assertIn("Interview stage recorded", frontend)
        self.assertIn("Offer stage recorded", frontend)

    def test_approved_application_keeps_final_preview_and_download_actions(self):
        frontend = Path(app_module.BASE_DIR, "static", "app.js").read_text(encoding="utf-8")
        self.assertIn('item.status === "approved"', frontend)
        self.assertIn('onclick="previewPreparation(${item.id})"', frontend)
        self.assertIn('/api/preparation/${item.id}/download/resume', frontend)
        self.assertIn('/api/preparation/${item.id}/download/cover-letter', frontend)

    def test_word_preview_cache_requires_matching_docx_fingerprints_and_invalidates_on_edit(self):
        original = Path(self.tmpdir.name) / "original.docx"
        candidate = Path(self.tmpdir.name) / "candidate.docx"
        output = Path(self.tmpdir.name) / "preview.pdf"
        original.write_bytes(b"original-docx")
        candidate.write_bytes(b"candidate-v1")
        output.write_bytes(b"stale-pdf")
        output.with_name(f"{output.name}.fingerprints.json").write_text(json.dumps({"original": "old", "candidate": "old"}), encoding="utf-8")
        calls = []

        def fake_word(_command, environment):
            calls.append(environment["JOB_TRACKER_PDF_PATH"])
            Path(environment["JOB_TRACKER_PDF_PATH"]).write_bytes(b"fresh-pdf")
            measurement = {"available": True, "pages": 1, "paragraphs": [{"index": 0, "lines": 1}]}
            return {"available": True, "original": measurement, "candidate": measurement, "timings_ms": {"word_layout_validation": 8.0, "pdf_generation": 12.0, "total": 20.0}}, None

        with patch.object(resume_layout, "_run_word", side_effect=fake_word):
            first = resume_layout.compare_docx_layout_and_export(original, candidate, output)
            second = resume_layout.compare_docx_layout_and_export(original, candidate, output)
            candidate.write_bytes(b"candidate-v2")
            third = resume_layout.compare_docx_layout_and_export(original, candidate, output)
        self.assertIn("available", first, first)
        self.assertFalse(first.get("cached", False), first)
        self.assertTrue(second.get("cached"), second)
        self.assertFalse(third.get("cached", False), third)
        self.assertEqual(len(calls), 2)
        self.assertEqual(output.read_bytes(), b"fresh-pdf")
        self.assertEqual(third["timings_ms"]["word_layout_validation"], 8.0)
        self.assertEqual(third["timings_ms"]["pdf_generation"], 12.0)
if __name__ == "__main__":
    unittest.main()
